"""Load the contingency equipment table (04_设备信息表.xlsx) into Dolt.

Dolt is the single source of truth for EquipmentConfig (binding `reader: dolt_rows`):
the ontology reads equipment lifecycle facts — GA/TR5/EOM/EOS/ESS dates — straight from
the Dolt `equipment_config` table, with no committed CSV/JSON in the read path. This script
is the human-facing ingestion step: Excel is edited by hand, then loaded here.

Mechanism (schema-driven, mirrors dolt_mirror._sync_object_type):
  * table name + DDL come from the ObjectType schema (object_type_table_name / build_create_table_sql);
  * Chinese Excel headers map to apiName via each property's `sourceColumnName`;
  * required system columns absent from the sheet (equipmentId / equipmentName / status /
    projectKey / sourceFile) are synthesised;
  * load is DELETE + INSERT (idempotent), then DOLT_COMMIT.

Run (Dolt sql-server must be up; DOLT_DATABASE_URL from ontology/.env):
    cd ontology && python load_contingency_excel_to_dolt.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ONTOLOGY_DIR = Path(__file__).resolve().parent
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))

try:  # load DOLT_DATABASE_URL et al. from ontology/.env when run standalone
    from dotenv import load_dotenv

    load_dotenv(ONTOLOGY_DIR / ".env")
except Exception:  # pragma: no cover - dotenv optional
    pass

import openpyxl  # noqa: E402

from dolt_schema_sync import (  # noqa: E402
    _resolve_engine,
    _sql_text,
    build_add_column_statements,
    build_create_table_sql,
    object_type_table_name,
    quote_ident,
    validate_sql_identifier,
)

OBJECT_TYPE = "EquipmentConfig"
# 04_设备信息表 is the JD 三期 (京东) project's equipment table.
PROJECT_KEY = "京东"
EXCEL_PATH = ONTOLOGY_DIR / "contingency-data" / "预案 excel数据" / "04_设备信息表.xlsx"
OBJECT_TYPES_PATH = ONTOLOGY_DIR / "schema" / "object-types.json"


def _load_object_schema() -> dict:
    schemas = json.loads(OBJECT_TYPES_PATH.read_text(encoding="utf-8"))
    schema = schemas.get(OBJECT_TYPE)
    if not isinstance(schema, dict):
        raise SystemExit(f"ObjectType {OBJECT_TYPE} not found in {OBJECT_TYPES_PATH}")
    return schema


def _schema_columns(schema: dict) -> list[str]:
    columns: list[str] = []
    for fallback, prop in (schema.get("properties") or {}).items():
        if not isinstance(prop, dict):
            continue
        column = str(prop.get("apiName") or fallback).strip()
        if column:
            validate_sql_identifier(column, kind="column")
            columns.append(column)
    return columns


def _source_column_to_api(schema: dict) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for fallback, prop in (schema.get("properties") or {}).items():
        if not isinstance(prop, dict):
            continue
        src = str(prop.get("sourceColumnName") or "").strip()
        if src:
            mapping[src] = str(prop.get("apiName") or fallback).strip()
    return mapping


def _read_excel_rows(schema: dict) -> list[dict]:
    header_to_api = _source_column_to_api(schema)
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["设备信息表"] if "设备信息表" in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in raw_rows[0]]

    rows: list[dict] = []
    for i, raw in enumerate(raw_rows[1:], start=1):
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        record: dict = {}
        for header, value in zip(headers, raw):
            api = header_to_api.get(header)
            if api and value is not None and str(value).strip() != "":
                record[api] = value
        # Synthesise required system columns the sheet does not carry.
        code = str(record.get("productCode") or "").strip()
        version = str(record.get("version") or "").strip()
        record["equipmentId"] = f"EQ-{code}-{version}".strip("-") or f"EQ-{i:03d}"
        record.setdefault("equipmentName", record.get("model") or record["equipmentId"])
        record.setdefault("status", "ACTIVE")
        record["projectKey"] = PROJECT_KEY
        record.setdefault("sourceFile", EXCEL_PATH.name)
        rows.append(record)
    return rows


def _sql_value(value):
    # Empty -> SQL NULL: date/datetime columns reject '' under Dolt strict mode; nullable
    # optional columns take NULL cleanly (required columns always carry a value upstream).
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    return value


def _fetch_column_names(conn, table_name: str) -> set[str]:
    result = conn.execute(
        _sql_text(
            "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = :t"
        ),
        {"t": table_name},
    )
    return {str(row[0]) for row in result.fetchall() if row[0] is not None}


def main() -> int:
    schema = _load_object_schema()
    table_name = object_type_table_name(OBJECT_TYPE, schema)
    columns = _schema_columns(schema)
    rows = _read_excel_rows(schema)
    if not rows:
        raise SystemExit(f"No rows read from {EXCEL_PATH}")

    engine = _resolve_engine(None)
    if engine is None:
        raise SystemExit(
            "Dolt is not reachable. Set DOLT_DATABASE_URL (ontology/.env) and start the dolt sql-server."
        )

    with engine.connect() as conn:
        conn.execute(_sql_text(build_create_table_sql(OBJECT_TYPE, schema)))
        existing = _fetch_column_names(conn, table_name)
        add_statements, _added = build_add_column_statements(table_name, schema, existing)
        for statement in add_statements:
            conn.execute(_sql_text(statement))
        conn.commit()

        conn.execute(_sql_text(f"DELETE FROM {quote_ident(table_name)}"))
        column_sql = ", ".join(quote_ident(c) for c in columns)
        values_sql = ", ".join(f":{c}" for c in columns)
        insert_sql = f"INSERT INTO {quote_ident(table_name)} ({column_sql}) VALUES ({values_sql})"
        for row in rows:
            conn.execute(_sql_text(insert_sql), {c: _sql_value(row.get(c)) for c in columns})
        conn.commit()

        try:
            conn.execute(
                _sql_text("CALL dolt_commit('-A', '-m', :m)"),
                {"m": f"Load {table_name} from {EXCEL_PATH.name} ({len(rows)} rows)"},
            )
            conn.commit()
        except Exception as exc:  # nothing-to-commit is fine
            if "nothing to commit" not in str(exc).lower():
                raise

    print(f"Loaded {len(rows)} rows into Dolt `{table_name}`:")
    for row in rows:
        print(
            f"  {row['equipmentId']:28} {row.get('lifecycleStatus',''):4} "
            f"GA={row.get('gaDate','')} EOM={row.get('eomDate','')} TR5={row.get('tr5Date','')} "
            f"EOS={row.get('eosDate','')} ESS={row.get('essDate','')}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
