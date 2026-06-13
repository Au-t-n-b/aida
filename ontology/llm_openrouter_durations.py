#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
结合计划基线 xlsx 与 POD 规模说明，经 OpenRouter 大模型为交付计划 CSV
填充或刷新「标准工期」「极限工期」列。

自包含复用与 llm_openrouter_backschedule 一致的 OpenRouter 调用方式；基线
读取逻辑与 `xlsx_rows` 类似，并支持选表、截断行数。
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------------------
# 与 schedule_plan_activities 一致的规范化（避免 import 循环，保持与 xlsx 行一致）
# ---------------------------------------------------------------------------


def normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value or "")
    normalized = normalized.replace("\u3000", " ")
    normalized = re.sub(r"\s+", " ", normalized)
    return normalized.strip()


DEFAULT_POD_SPEC = (
    "1个POD 灵衢2688根，综合布线460根。"
    "1个POD384卡，12个计算柜，4个总线柜。"
)

COL_STANDARD = "标准工期"
COL_LIMIT = "极限工期"
COL_SERIAL = "序列号"
COL_NAME = "活动名称"
COL_UNIT = "管理单元"
COL_SLA = "SLA"
COL_DEP = "依赖活动"
COL_ID = "活动ID"

SYSTEM_PROMPT = """\
你是算力/数据中心交付计划分析助手。用户会提供：
(1) 项目 POD 与布线等规模说明；
(2) 历史「计划基线」表中的片段（xlsx 转存为 JSON 行，列名不固定，请自行读列含义）；
(3) 当前需填写「标准工期」与「极限工期」的交付计划行，每行有唯一 row_id。

【工期格式】
- 填写「N天」形式，N 为正整数，如「3天」「10天」。
- 若与基线/常识确实无法给数值，可填空字符串，但应极少出现。
- 若基线中的SLA为空或者无计算逻辑，可填空字符串。

【row_id 规则】
- 以 s: 开头表示有「序列号」的行的序列号全串（小写/原样以用户表为准，须与输入完全一致）。
- 以 r: 开头表示无序列号时按数据行 0 起的行号，如 r:4。

【重要】
- 必须只输出**一个** JSON 对象，不要 markdown 代码围栏，不要除 JSON 外的任何文字。
- 无计算逻辑的直接按「计划基线」中的标准工期、极限工期直接填写
- 有计算逻辑的标准工期和极限工期，严格结合基线、POD 说明与行内活动名/依赖/SLA/设备列表进行计算，不足1天按1天算。

JSON 结构要求：
{
  "batch_reason": "本批不引用具体行号的简短推理",
  "updates": [
    {"row_id": "s:... 或 r:...","标准工期": "x天","极限工期": "x天"}
  ]
}
"""


# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------


def get_openrouter_config(
    args: argparse.Namespace,
) -> tuple[str, str, str]:
    api_key = args.openrouter_api_key or os.getenv("OPENROUTER_API_KEY", "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")
    base_url = args.openrouter_base_url or os.getenv(
        "OPENROUTER_BASE_URL", "https://openrouter.ai/api/v1"
    )
    model = args.openrouter_model or os.getenv("OPENROUTER_MODEL", "GPT-5.4")
    return api_key, base_url.rstrip("/"), model


# ---------------------------------------------------------------------------
# 基线与 CSV
# ---------------------------------------------------------------------------


def _xlsx_data_rows(
    path: Path,
    sheet: str | None,
    max_rows: int,
) -> tuple[list[dict[str, str]], str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            f"无法读取 `{path.name}`：请先安装 openpyxl（pip install openpyxl）。"
        ) from exc

    wb = load_workbook(filename=path, read_only=True, data_only=True)
    names = wb.sheetnames
    if not names:
        return [], ""

    if sheet and sheet in names:
        use_sheet = sheet
    elif sheet and sheet not in names:
        raise ValueError(
            f"工作表「{sheet}」不存在。可用表名：{', '.join(names)}"
        )
    else:
        use_sheet = names[0]

    ws = wb[use_sheet]
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return [], use_sheet

    header = [normalize_text(str(c)) if c is not None else "" for c in all_rows[0]]
    out: list[dict[str, str]] = []
    for values in all_rows[1:]:
        if max_rows > 0 and len(out) >= max_rows:
            break
        row: dict[str, str] = {}
        for i, col in enumerate(header):
            if not col:
                continue
            v = values[i] if i < len(values) else ""
            row[col] = "" if v is None else str(v)
        if any(normalize_text(x) for x in row.values()):
            out.append(row)
    return out, use_sheet


def baseline_to_prompt_chunk(rows: list[dict[str, str]], max_json_chars: int) -> str:
    """将基线行序列化为有限长度的 JSON 字符串，用于提示词。"""
    raw = json.dumps(rows, ensure_ascii=False, indent=0)
    if len(raw) <= max_json_chars:
        return raw
    return raw[: max_json_chars - 3] + "..."


def read_target_csv(
    path: Path,
) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = [dict(r) for r in reader]
    for required in (COL_NAME,):
        if required not in fieldnames:
            raise RuntimeError(f"目标 CSV 缺少列「{required}」。")
    for c in (COL_STANDARD, COL_LIMIT, COL_SERIAL):
        if c not in fieldnames:
            fieldnames.append(c)
    return fieldnames, rows


def make_row_id(data_index: int, row: dict[str, str]) -> str:
    s = normalize_text(row.get(COL_SERIAL, ""))
    if s:
        return f"s:{s}"
    return f"r:{data_index}"


def is_row_relevant(row: dict[str, str]) -> bool:
    """有活动名则视为可参与工期估算的业务行；全空可跳过。"""
    return bool(normalize_text(row.get(COL_NAME, "")))


def is_both_durations_empty(row: dict[str, str]) -> bool:
    a = normalize_text(row.get(COL_STANDARD, ""))
    b = normalize_text(row.get(COL_LIMIT, ""))
    return not a and not b


# ---------------------------------------------------------------------------
# OpenRouter
# ---------------------------------------------------------------------------


def call_openrouter(
    api_key: str,
    base_url: str,
    model: str,
    user_content: str,
    temperature: float,
    timeout: int,
) -> tuple[dict[str, Any] | None, str | None]:
    payload: dict[str, Any] = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_content},
        ],
        "temperature": temperature,
    }
    req = urllib.request.Request(
        url=f"{base_url}/chat/completions",
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            body = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as ex:
        err = ex.read().decode("utf-8", errors="ignore")
        return None, f"HTTP {ex.code}：{err[:500]}"
    except Exception as ex:  # pragma: no cover
        return None, str(ex)

    choices = body.get("choices") or []
    if not choices:
        return None, "无 choices 字段"
    content = (choices[0].get("message") or {}).get("content") or ""
    if not (content and content.strip()):
        return None, "模型返回空 content"
    parsed = try_parse_model_json(content.strip())
    if parsed is None:
        return None, f"无法解析 JSON。原始前 500 字：{content[:500]}"
    return parsed, None


def try_parse_model_json(text: str) -> dict[str, Any] | None:
    text = text.strip()
    m = re.search(
        r"```(?:json)?\s*(\{.*\})\s*```", text, flags=re.DOTALL | re.IGNORECASE
    )
    if m:
        text = m.group(1).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    o = text.find("{")
    c = text.rfind("}")
    if o >= 0 and c > o:
        try:
            return json.loads(text[o : c + 1])
        except json.JSONDecodeError:
            return None
    return None


# ---------------------------------------------------------------------------
# 合批、合并
# ---------------------------------------------------------------------------


@dataclass
class PlanRow:
    data_index: int
    row_id: str
    row: dict[str, str]


def build_user_payload(
    pod_spec: str,
    baseline_json: str,
    batch: list[PlanRow],
) -> str:
    lines = [f"【POD 与规模】\n{pod_spec}\n", "【基线表 JSON 片段】\n" + baseline_json, "\n【本批交付计划行】\n"]
    for pr in batch:
        r = pr.row
        rec = {
            "row_id": pr.row_id,
            COL_ID: r.get(COL_ID, ""),
            COL_UNIT: r.get(COL_UNIT, ""),
            COL_NAME: r.get(COL_NAME, ""),
            COL_SLA: r.get(COL_SLA, ""),
            COL_DEP: r.get(COL_DEP, ""),
            "批次": r.get("批次", ""),
            "设备": r.get("设备型号&数量的列表", r.get("设备", ""))[:2000],
            COL_STANDARD: r.get(COL_STANDARD, ""),
            COL_LIMIT: r.get(COL_LIMIT, ""),
        }
        lines.append(json.dumps(rec, ensure_ascii=False))
    return "\n".join(lines)


def index_rows_by_id(rows: list[dict[str, str]]) -> dict[str, int]:
    m: dict[str, int] = {}
    for i, row in enumerate(rows):
        m[make_row_id(i, row)] = i
    return m


def apply_updates(
    rows: list[dict[str, str]],
    updates: list[dict[str, Any]],
    warnings: list[str],
) -> int:
    id2ix = index_rows_by_id(rows)
    n = 0
    for u in updates:
        if not isinstance(u, dict):
            continue
        rid = normalize_text(str(u.get("row_id", "")))
        st = u.get(COL_STANDARD)
        lm = u.get(COL_LIMIT)
        if st is not None and not isinstance(st, str):
            st = str(st)
        if lm is not None and not isinstance(lm, str):
            lm = str(lm)
        if st is not None:
            st = st.strip() if st else ""
        if lm is not None:
            lm = lm.strip() if lm else ""
        if not rid or rid not in id2ix:
            warnings.append(f"未知 row_id，已跳过：{rid!r}")
            continue
        ix = id2ix[rid]
        if st is not None:
            rows[ix][COL_STANDARD] = st
        if lm is not None:
            rows[ix][COL_LIMIT] = lm
        n += 1
    return n


def default_output_path(source: Path) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return source.with_name(f"{source.stem}_刷新工期_{ts}.csv")


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="根据计划基线 xlsx 与 OpenRouter 刷新交付计划「标准工期/极限工期」。"
    )
    p.add_argument(
        "--baseline", required=True, type=Path, help="计划基线 .xlsx 路径（如 计划基线0402.xlsx）"
    )
    p.add_argument(
        "--file", required=True, type=Path, help="待刷新的交付计划 .csv 路径"
    )
    p.add_argument("--baseline-sheet", help="基线表工作表名；缺省为第一张表")
    p.add_argument(
        "--baseline-max-rows", type=int, default=200, help="基线中最多参与提示的数据行数"
    )
    p.add_argument(
        "--baseline-max-json-chars", type=int, default=80_000, help="基线 JSON 在提示中最大字符数"
    )
    p.add_argument(
        "--pod-spec", type=Path, help="POD 规模说明文件（UTF-8）；缺省使用内置文案"
    )
    p.add_argument(
        "--batch-size", type=int, default=20, help="每批送给模型的行数"
    )
    p.add_argument(
        "--temperature", type=float, default=0.15, help="补全用的采样温度"
    )
    p.add_argument(
        "--request-timeout", type=int, default=120, help="单请求超时秒"
    )
    p.add_argument(
        "--only-empty", action=argparse.BooleanOptionalAction, default=True,
        help="仅当标准工期与极限工期都为空时更新（默认开）"
    )
    p.add_argument(
        "--all-rows", action="store_true",
        help="对表中应处理的全部相关行都请求模型（与 --no-only-empty 同义，覆盖两列）"
    )
    p.add_argument(
        "--max-candidate-rows", type=int, default=0,
        help="最多参与模型补全的候选行数，0 表示不限制；大表试跑时可设小值以缩短耗时",
    )
    p.add_argument(
        "--dry-run", action="store_true", help="调用模型但不写回 CSV，仅打印/日志"
    )
    p.add_argument(
        "--output", type=Path, help="输出 CSV 路径；默认 原名_刷新工期_时间戳.csv"
    )
    p.add_argument(
        "--in-place", action="store_true", help="写回与 --file 同一路径（请谨慎）"
    )
    p.add_argument("--openrouter-api-key", help="或环境变量 OPENROUTER_API_KEY")
    p.add_argument(
        "--openrouter-base-url", help="或环境变量 OPENROUTER_BASE_URL"
    )
    p.add_argument(
        "--openrouter-model", help="或环境变量 OPENROUTER_MODEL"
    )
    p.add_argument(
        "--log-json", type=Path, help="将每批模型解析后的 JSON 另存为 NDJSON 日志"
    )
    return p


def run() -> int:
    ap = build_arg_parser()
    args = ap.parse_args()
    if args.all_rows:
        # 覆盖为「非 only_empty」——即对候选行不检查两列都空
        only_empty = False
    else:
        only_empty = bool(args.only_empty)

    api_key, base_url, model = get_openrouter_config(args)
    if not api_key or not model:
        print("需要 OPENROUTER_API_KEY 与 OPENROUTER_MODEL（或 CLI 参数）。", file=sys.stderr)
        return 1

    baseline_path = args.baseline.expanduser().resolve()
    target_path = args.file.expanduser().resolve()
    for pth in (baseline_path, target_path):
        if not pth.exists():
            print(f"文件不存在：{pth}", file=sys.stderr)
            return 1

    pod_text = DEFAULT_POD_SPEC
    if args.pod_spec:
        sp = args.pod_spec.expanduser().resolve()
        if not sp.exists():
            print(f"POD 文件不存在：{sp}", file=sys.stderr)
            return 1
        pod_text = sp.read_text(encoding="utf-8")

    try:
        b_rows, used_sheet = _xlsx_data_rows(
            baseline_path, args.baseline_sheet, args.baseline_max_rows
        )
    except (RuntimeError, ValueError) as ex:
        print(str(ex), file=sys.stderr)
        return 1

    baseline_json = baseline_to_prompt_chunk(
        b_rows, max_json_chars=args.baseline_max_json_chars
    )
    try:
        fieldnames, t_rows = read_target_csv(target_path)
    except RuntimeError as ex:
        print(str(ex), file=sys.stderr)
        return 1

    candidates: list[PlanRow] = []
    for i, row in enumerate(t_rows):
        if not is_row_relevant(row):
            continue
        if only_empty and not is_both_durations_empty(row):
            continue
        candidates.append(PlanRow(data_index=i, row_id=make_row_id(i, row), row=row))

    mcr = int(args.max_candidate_rows)
    if mcr > 0:
        candidates = candidates[:mcr]

    if not candidates:
        print("没有需处理的行（可检查 --only-empty 或表内容）。", file=sys.stderr)
        return 0

    warnings: list[str] = []
    all_parsed: list[dict[str, Any]] = []
    bs = max(1, int(args.batch_size))
    for start in range(0, len(candidates), bs):
        batch = candidates[start : start + bs]
        user_pl = build_user_payload(pod_text, baseline_json, batch)
        data, err = call_openrouter(
            api_key,
            base_url,
            model,
            user_pl,
            float(args.temperature),
            int(args.request_timeout),
        )
        if err and data is None:
            # 重试一次
            warnings.append(f"批次 {start//bs+1} 首失败：{err}，重试中…")
            data, err = call_openrouter(
                api_key,
                base_url,
                model,
                user_pl,
                float(args.temperature),
                int(args.request_timeout),
            )
        if err and data is None:
            warnings.append(f"批次 {start//bs+1} 放弃：{err}")
            if args.dry_run:
                print("--- 失败批次 user 摘要 ---\n" + user_pl[:800] + "…")
            continue
        if data is not None:
            if args.dry_run:
                print(json.dumps(data, ensure_ascii=False, indent=2))
            all_parsed.append(data)
            if args.log_json:
                with args.log_json.open("a", encoding="utf-8") as lj:
                    lj.write(json.dumps(data, ensure_ascii=False) + "\n")
            ulist = data.get("updates")
            if isinstance(ulist, list):
                apply_updates(t_rows, ulist, warnings)  # type: ignore[arg-type]
            else:
                warnings.append(f"批次 {start//bs+1} 无 updates 列表")

    for w in warnings:
        print(f"警告: {w}", file=sys.stderr)

    out_path: Path | None = None
    if not args.dry_run:
        if args.in_place:
            out_path = target_path
        else:
            out_path = (args.output or default_output_path(target_path)).resolve()
        with out_path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            w.writeheader()
            w.writerows(t_rows)
        print(f"已写出：{out_path}")
    else:
        print("dry-run：未修改 CSV 文件。")
    print(f"基线工作表：{used_sheet}，行数(截断后)：{len(b_rows)}")

    return 0


if __name__ == "__main__":
    sys.exit(run())
