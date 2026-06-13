"""Load the 机房机柜信息表 (PoD × 机柜归位) into Dolt as PodCabinetLayout.

Dolt is the single source of truth for PodCabinetLayout (binding `reader: dolt_rows`,
table `pod_cabinet_layout`): the contingency ontology / 数字孪生 read the real
PoD→机房→机柜归位 facts straight from Dolt. This script is the human-facing ingestion
step — 机房机柜信息表.xlsx is maintained by hand, then loaded here.

The current source of truth is temp/京东三期/早期介入/交付预案/输出结果/机房机柜信息表.xlsx
Sheet1. Only the new header dialect is accepted. `预案版本号` is a load filter only:
blank and `草稿` rows are skipped, and the version value is not stored in the ontology row.

Rows are aggregated by podName after filtering: machineRoomName = first non-empty; each
cabinet column = order-preserving, de-duplicated union of the non-empty values across the
PoD's rows.

Required system columns the sheet lacks (podCabinetLayoutId / status / sourceFile) are
synthesised. Load is DELETE + INSERT (idempotent), then DOLT_COMMIT — mirroring the sibling
contingency loaders.

Run (Dolt sql-server must be up; DOLT_DATABASE_URL from ontology/.env):
    cd ontology && python load_pod_cabinet_layout_to_dolt.py
    # override the source file:
    POD_CABINET_XLSX="D:\\path\\to\\机房机柜信息表.xlsx" python load_pod_cabinet_layout_to_dolt.py
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

ONTOLOGY_DIR = Path(__file__).resolve().parent
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))

try:  # load DOLT_DATABASE_URL et al. from ontology/.env when run standalone
    from dotenv import load_dotenv

    load_dotenv(ONTOLOGY_DIR / ".env")
except Exception:  # pragma: no cover - dotenv optional
    pass

import openpyxl  # noqa: E402

from dolt_schema_sync import (  # noqa: E402
    _resolve_engine,
    _sql_text,
    build_add_column_statements,
    build_create_table_sql,
    object_type_table_name,
    quote_ident,
    validate_sql_identifier,
)

OBJECT_TYPE = "PodCabinetLayout"
OBJECT_TYPES_PATH = ONTOLOGY_DIR / "schema" / "object-types.json"

# Authoritative current source. Override with $POD_CABINET_XLSX.
DEFAULT_EXCEL_PATH = Path(
    r"D:\Code\aida\temp\京东三期\早期介入\交付预案\输出结果\机房机柜信息表.xlsx"
)
EXCEL_PATH = Path(os.environ.get("POD_CABINET_XLSX") or DEFAULT_EXCEL_PATH)

# Chinese header → schema apiName. Historical aliases are intentionally not accepted.
HEADER_MAP = {
    "PoD名称": "podName",
    "机房名称": "machineRoomName",
    "计算柜": "computeCabinets",
    "总线柜": "busbarCabinets",
    "参数面Leaf柜": "paramLeafCabinets",
    "样本面Leaf柜": "sampleLeafCabinets",
    "业务面Leaf柜": "businessLeafCabinets",
    "管理面柜": "mgmtCabinets",
}
VERSION_HEADER = "预案版本号"
REQUIRED_HEADERS = tuple(HEADER_MAP) + (VERSION_HEADER,)
CABINET_FIELDS = (
    "computeCabinets",
    "busbarCabinets",
    "paramLeafCabinets",
    "sampleLeafCabinets",
    "businessLeafCabinets",
    "mgmtCabinets",
)


def _load_object_schema() -> dict:
    schemas = json.loads(OBJECT_TYPES_PATH.read_text(encoding="utf-8"))
    schema = schemas.get(OBJECT_TYPE)
    if not isinstance(schema, dict):
        raise SystemExit(f"ObjectType {OBJECT_TYPE} not found in {OBJECT_TYPES_PATH}")
    return schema


def _schema_columns(schema: dict) -> list[str]:
    columns: list[str] = []
    for fallback, prop in (schema.get("properties") or {}).items():
        if not isinstance(prop, dict):
            continue
        column = str(prop.get("apiName") or fallback).strip()
        if column:
            validate_sql_identifier(column, kind="column")
            columns.append(column)
    return columns


def _split_tokens(value: str) -> list[str]:
    # Cabinet cells may pack multiple codes separated by an English/Chinese comma or 、.
    return [t.strip() for t in re.split(r"[,，、]", value) if t.strip()]


def _validate_headers(headers: list[str]) -> None:
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise SystemExit(
            f"Missing required headers in {EXCEL_PATH.name}: {', '.join(missing)}; got {headers}"
        )


def _is_non_draft_version(value: object) -> bool:
    version = str(value).strip() if value is not None else ""
    return bool(version) and version != "草稿"


def _read_pod_rows() -> list[dict]:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in raw_rows[0]]
    _validate_headers(headers)
    version_i = headers.index(VERSION_HEADER)

    # Aggregate by podName, preserving first-seen order.
    order: list[str] = []
    agg: dict[str, dict] = {}
    for raw in raw_rows[1:]:
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        if version_i >= len(raw) or not _is_non_draft_version(raw[version_i]):
            continue
        rec: dict = {}
        for header, value in zip(headers, raw):
            api = HEADER_MAP.get(header)
            if api and value is not None and str(value).strip() != "":
                rec[api] = str(value).strip()
        pod = rec.get("podName")
        if not pod:
            continue
        if pod not in agg:
            order.append(pod)
            agg[pod] = {"podName": pod, "machineRoomName": rec.get("machineRoomName")}
            for field in CABINET_FIELDS:
                agg[pod][field] = []  # ordered-unique token collector
        cur = agg[pod]
        if not cur.get("machineRoomName") and rec.get("machineRoomName"):
            cur["machineRoomName"] = rec["machineRoomName"]
        for field in CABINET_FIELDS:
            cell = rec.get(field)
            if not cell:
                continue
            for token in _split_tokens(cell):
                if token not in cur[field]:
                    cur[field].append(token)

    rows: list[dict] = []
    for i, pod in enumerate(order, start=1):
        cur = agg[pod]
        row = {
            "podCabinetLayoutId": f"pod_cabinet_layout-{i:04d}",
            "podName": pod,
            "status": "ACTIVE",
            "machineRoomName": cur.get("machineRoomName"),
            "sourceFile": EXCEL_PATH.name,
        }
        for field in CABINET_FIELDS:
            row[field] = ",".join(cur[field]) if cur[field] else None
        rows.append(row)
    return rows


def _sql_value(value):
    # Empty -> SQL NULL: datetime columns reject '' under Dolt strict mode; nullable optional
    # columns take NULL cleanly (required columns always carry a value upstream).
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    return value


def _fetch_column_names(conn, table_name: str) -> set[str]:
    result = conn.execute(
        _sql_text(
            "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = :t"
        ),
        {"t": table_name},
    )
    return {str(row[0]) for row in result.fetchall() if row[0] is not None}


def main() -> int:
    if not EXCEL_PATH.exists():
        raise SystemExit(f"Source workbook not found: {EXCEL_PATH}")

    schema = _load_object_schema()
    table_name = object_type_table_name(OBJECT_TYPE, schema)
    columns = _schema_columns(schema)
    rows = _read_pod_rows()
    if not rows:
        raise SystemExit(f"No rows read from {EXCEL_PATH}")

    engine = _resolve_engine(None)
    if engine is None:
        raise SystemExit(
            "Dolt is not reachable. Set DOLT_DATABASE_URL (ontology/.env) and start the dolt sql-server."
        )

    with engine.connect() as conn:
        conn.execute(_sql_text(build_create_table_sql(OBJECT_TYPE, schema)))
        existing = _fetch_column_names(conn, table_name)
        add_statements, _added = build_add_column_statements(table_name, schema, existing)
        for statement in add_statements:
            conn.execute(_sql_text(statement))
        conn.commit()

        conn.execute(_sql_text(f"DELETE FROM {quote_ident(table_name)}"))
        column_sql = ", ".join(quote_ident(c) for c in columns)
        values_sql = ", ".join(f":{c}" for c in columns)
        insert_sql = f"INSERT INTO {quote_ident(table_name)} ({column_sql}) VALUES ({values_sql})"
        for row in rows:
            conn.execute(_sql_text(insert_sql), {c: _sql_value(row.get(c)) for c in columns})
        conn.commit()

        try:
            conn.execute(
                _sql_text("CALL dolt_commit('-A', '-m', :m)"),
                {"m": f"Load {table_name} from {EXCEL_PATH.name} ({len(rows)} PoD rows)"},
            )
            conn.commit()
        except Exception as exc:  # nothing-to-commit is fine
            if "nothing to commit" not in str(exc).lower():
                raise

    print(f"Loaded {len(rows)} PoD rows into Dolt `{table_name}` from {EXCEL_PATH.name}:")
    for row in rows:
        cabinets = " ".join(
            f"{field}={row[field]}" for field in CABINET_FIELDS if row.get(field)
        ) or "(机柜列为空)"
        print(f"  {row['podCabinetLayoutId']:24} {row['podName']:18} 机房={row.get('machineRoomName','')!s:10} {cabinets}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
