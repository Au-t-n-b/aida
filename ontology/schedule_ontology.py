#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Local Ontology/OSDK-style facade for delivery schedule objects.

The facade owns object identity, object collections, link traversal, and
function entrypoints. Existing scheduling algorithms remain the source of
truth for forward/backward date calculation.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from schedule_plan_activities import (
    Activity,
    FIELD_DEPENDENCY,
    FIELD_ROW_ID,
    build_activities,
    normalize_activity_key,
    normalize_text,
    parse_duration_days,
    resolve_target_key_list,
    resolve_user_activity_key,
    split_management_units,
)

FIELD_STANDARD_DURATION = "标准工期"
FIELD_LIMIT_DURATION = "极限工期"
FIELD_RISK_BELOW_STANDARD = "低于标准工期的风险"
FIELD_EQUIPMENT_LIST = "设备型号&数量的列表"
FIELD_REMOTE_TEAM = "远程团队"
FIELD_ONSITE_TEAM = "现场团队"
FIELD_OWNER = "责任人"
FIELD_ORIGINAL_MANAGEMENT_UNIT = "原始管理单元"
ROW_KEY_SEP = "::"


def normalize_project_key(value: str) -> str:
    key = normalize_text(value).lower()
    key = re.sub(r"\s+", "_", key)
    key = re.sub(r"[^0-9a-zA-Z_\-\u4e00-\u9fff]+", "_", key)
    key = re.sub(r"_+", "_", key).strip("_")
    return key or "DeliveryProject"


def default_project_key(source_path: Path | None) -> str:
    if source_path is None:
        return "DeliveryProject"
    return normalize_project_key(source_path.stem)


def make_global_row_key(project_key: str, local_row_key: str) -> str:
    return f"{project_key}{ROW_KEY_SEP}{local_row_key}"


def normalize_pod_name(value: str) -> str:
    return normalize_activity_key(value)


def make_pod_key(project_key: str, pod_name: str) -> str:
    return f"{project_key}{ROW_KEY_SEP}{normalize_pod_name(pod_name)}"


def csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def xlsx_rows(path: Path) -> list[dict[str, str]]:
    csv_fallback = path.with_suffix(".csv")
    if csv_fallback.exists():
        return csv_rows(csv_fallback)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            f"Unable to read `{path.name}`: openpyxl is not installed and fallback CSV `{csv_fallback.name}` was not found."
        ) from exc

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = [normalize_text(str(value)) if value is not None else "" for value in rows[0]]
    output: list[dict[str, str]] = []
    for values in rows[1:]:
        row = {}
        for index, column in enumerate(header):
            if not column:
                continue
            cell_value = values[index] if index < len(values) else ""
            row[column] = "" if cell_value is None else str(cell_value)
        if any(normalize_text(str(value)) for value in row.values()):
            output.append(row)
    return output


def load_rows(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return csv_rows(path)
    if suffix == ".xlsx":
        return xlsx_rows(path)
    raise RuntimeError(f"Unsupported file type: `{path.suffix}`. Only .csv / .xlsx are supported.")


def build_successors(activities: dict[str, Activity]) -> dict[str, list[str]]:
    successors: defaultdict[str, list[str]] = defaultdict(list)
    for activity in activities.values():
        for dependency_key in activity.dependency_keys:
            successors[dependency_key].append(activity.key)
    return dict(successors)


def _coerce_duration_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return f"{value}天" if value >= 0 else ""
    text = normalize_text(str(value))
    if not text:
        return ""
    match = re.fullmatch(r"(\d+)\s*(?:天|d|D)?", text)
    if match:
        return f"{int(match.group(1))}天"
    return text


def _localize_duration_overrides(
    overrides: dict[str, object] | None,
    project_key: str,
) -> dict[str, str]:
    if not overrides:
        return {}
    prefix = f"{project_key}{ROW_KEY_SEP}"
    localized: dict[str, str] = {}
    for raw_key, raw_value in overrides.items():
        key = normalize_text(str(raw_key))
        if not key:
            continue
        local_key = key[len(prefix):] if key.startswith(prefix) else key
        duration_text = _coerce_duration_text(raw_value)
        if duration_text:
            localized[local_key] = duration_text
    return localized


def _apply_effective_sla_overrides(
    activities: dict[str, Activity],
    overrides_by_local_key: dict[str, str],
    warnings: list[str],
) -> dict[str, str]:
    applied: dict[str, str] = {}
    for local_key, duration_text in overrides_by_local_key.items():
        activity = activities.get(local_key)
        if activity is None:
            warnings.append(f"本体 effective SLA 指向的活动 `{local_key}` 未在计划中找到，已忽略。")
            continue
        duration_days = parse_duration_days(duration_text, warnings, activity.name)
        if duration_days <= 0:
            continue
        activity.sla_raw = duration_text
        activity.duration_days = duration_days
        applied[local_key] = duration_text
    return applied


@dataclass(frozen=True)
class ScheduleFunctionRun:
    results: dict[str, Any]
    warnings: list[str]
    metadata: dict[str, Any]


@dataclass(frozen=True)
class DeliveryProject:
    _ontology: LocalScheduleOntology

    @property
    def project_key(self) -> str:
        return self._ontology.project_key

    @property
    def projectKey(self) -> str:
        return self.project_key

    @property
    def project_name(self) -> str:
        return self._ontology.project_name

    @property
    def projectName(self) -> str:
        return self.project_name

    @property
    def source_file(self) -> str:
        return self._ontology.source_path.name if self._ontology.source_path else ""

    @property
    def sourceFile(self) -> str:
        return self.source_file

    @property
    def source_updated_at(self) -> str:
        return self._ontology.source_updated_at

    @property
    def sourceUpdatedAt(self) -> str:
        return self.source_updated_at

    @property
    def plan_version(self) -> str:
        return self._ontology.plan_version

    @property
    def planVersion(self) -> str:
        return self.plan_version

    @property
    def description(self) -> str:
        return self._ontology.project_description

    def plan_rows(self) -> list[DeliveryPlanRow]:
        return self._ontology.objects.DeliveryPlanRow.all()

    def pods(self) -> list[DeliveryPod]:
        return self._ontology.objects.DeliveryPod.all()

    def to_schema_object(self) -> dict[str, str]:
        return {
            "projectKey": self.project_key,
            "projectName": self.project_name,
            "sourceFile": self.source_file,
            "sourceUpdatedAt": self.source_updated_at,
            "planVersion": self.plan_version,
            "description": self.description,
        }


@dataclass(frozen=True)
class DeliveryPod:
    _ontology: LocalScheduleOntology
    _pod_name: str

    @property
    def pod_key(self) -> str:
        return make_pod_key(self.project_key, self.pod_name)

    @property
    def podKey(self) -> str:
        return self.pod_key

    @property
    def project_key(self) -> str:
        return self._ontology.project_key

    @property
    def projectKey(self) -> str:
        return self.project_key

    @property
    def pod_name(self) -> str:
        return self._pod_name

    @property
    def podName(self) -> str:
        return self.pod_name

    @property
    def management_unit(self) -> str:
        return self._pod_name

    @property
    def managementUnit(self) -> str:
        return self.management_unit

    @property
    def batch(self) -> str:
        values = [
            row.batch
            for row in self.plan_rows()
            if row.batch
        ]
        return values[0] if values else ""

    @property
    def description(self) -> str:
        return ""

    def project(self) -> DeliveryProject:
        return self._ontology.project()

    def plan_rows(self) -> list[DeliveryPlanRow]:
        return self._ontology.rows_for_pod(self.pod_name)

    def to_schema_object(self) -> dict[str, str]:
        return {
            "podKey": self.pod_key,
            "projectKey": self.project_key,
            "podName": self.pod_name,
            "managementUnit": self.management_unit,
            "batch": self.batch,
            "description": self.description,
        }


@dataclass(frozen=True)
class DeliveryPlanRow:
    _ontology: LocalScheduleOntology
    _activity: Activity

    @property
    def activity(self) -> Activity:
        return self._activity

    @property
    def row_key(self) -> str:
        return self._ontology.to_global_row_key(self.local_row_key)

    @property
    def rowKey(self) -> str:
        return self.row_key

    @property
    def local_row_key(self) -> str:
        return self._activity.row_key

    @property
    def localRowKey(self) -> str:
        return self.local_row_key

    @property
    def project_key(self) -> str:
        return self._ontology.project_key

    @property
    def projectKey(self) -> str:
        return self.project_key

    @property
    def activity_name(self) -> str:
        return self._activity.name

    @property
    def activityName(self) -> str:
        return self.activity_name

    @property
    def activity_id(self) -> str:
        return self._activity.activity_id

    @property
    def activityId(self) -> str:
        return self.activity_id

    @property
    def management_unit(self) -> str:
        return self._activity.unit

    @property
    def managementUnit(self) -> str:
        return self.management_unit

    @property
    def start_date(self) -> str:
        return self._activity.planned_start

    @property
    def startDate(self) -> str:
        return self.start_date

    @property
    def end_date(self) -> str:
        return self._activity.planned_end

    @property
    def endDate(self) -> str:
        return self.end_date

    @property
    def sla(self) -> str:
        return self._activity.sla_raw

    @property
    def duration_days(self) -> int:
        return self._activity.duration_days

    @property
    def batch(self) -> str:
        return self._activity.batch

    @property
    def dependency_activities(self) -> list[str]:
        return list(self._activity.dependency_names)

    @property
    def dependencyActivities(self) -> str:
        return normalize_text(self.raw_row.get(FIELD_DEPENDENCY, ""))

    @property
    def unresolved_dependencies(self) -> list[str]:
        return list(self._activity.unresolved_dependencies)

    @property
    def raw_row(self) -> dict[str, str]:
        return dict(self._ontology.raw_rows_by_key.get(self.local_row_key, {}))

    def depends_on(self) -> list[DeliveryPlanRow]:
        return [self._ontology.row(key) for key in self._activity.dependency_keys]

    def is_prerequisite_for(self) -> list[DeliveryPlanRow]:
        return [self._ontology.row(key) for key in self._ontology.successors.get(self.local_row_key, [])]

    def project(self) -> DeliveryProject:
        return self._ontology.project()

    def pods(self) -> list[DeliveryPod]:
        return self._ontology.pods_for_row(self)

    def to_schema_object(self) -> dict[str, str]:
        raw = self.raw_row
        standard_duration = (
            self._ontology.effective_standard_duration_by_key.get(self.local_row_key)
            or normalize_text(raw.get(FIELD_STANDARD_DURATION, ""))
        )
        limit_duration = (
            self._ontology.effective_limit_duration_by_key.get(self.local_row_key)
            or normalize_text(raw.get(FIELD_LIMIT_DURATION, ""))
        )
        return {
            "rowKey": self.row_key,
            "projectKey": self.project_key,
            "localRowKey": self.local_row_key,
            "serial": normalize_text(raw.get(FIELD_ROW_ID, "")),
            "activityId": self.activity_id,
            "managementUnit": self.management_unit,
            "activityName": self.activity_name,
            "startDate": self.start_date,
            "endDate": self.end_date,
            "sla": self.sla,
            "standardDuration": standard_duration,
            "limitDuration": limit_duration,
            "riskBelowStandard": normalize_text(raw.get(FIELD_RISK_BELOW_STANDARD, "")),
            "batch": self.batch,
            "equipmentList": normalize_text(raw.get(FIELD_EQUIPMENT_LIST, "")),
            "remoteTeam": normalize_text(raw.get(FIELD_REMOTE_TEAM, "")),
            "onsiteTeam": normalize_text(raw.get(FIELD_ONSITE_TEAM, "")),
            "owner": normalize_text(raw.get(FIELD_OWNER, "")),
            "originalManagementUnit": normalize_text(raw.get(FIELD_ORIGINAL_MANAGEMENT_UNIT, "")),
            "dependencyActivities": normalize_text(raw.get(FIELD_DEPENDENCY, "")),
        }


class DeliveryPlanRowSet:
    def __init__(self, ontology: LocalScheduleOntology) -> None:
        self._ontology = ontology

    def __iter__(self) -> Iterator[DeliveryPlanRow]:
        return iter(self.all())

    def __len__(self) -> int:
        return len(self._ontology.activities)

    def all(self) -> list[DeliveryPlanRow]:
        keys = sorted(
            self._ontology.activities,
            key=lambda key: self._ontology.activities[key].row_order,
        )
        return [self.get(key) for key in keys]

    def get(self, row_key: str) -> DeliveryPlanRow:
        return self._ontology.row(row_key)

    def get_or_none(self, row_key: str) -> DeliveryPlanRow | None:
        local_key = self._ontology.to_local_row_key(row_key)
        if local_key not in self._ontology.activities:
            return None
        return self.get(local_key)

    def resolve_key(self, user_input: str) -> str:
        return resolve_user_activity_key(user_input, self._ontology.activities)

    def resolve(self, user_input: str) -> DeliveryPlanRow:
        return self.get(self.resolve_key(user_input))

    def resolve_many(self, raw_targets: str) -> list[DeliveryPlanRow]:
        return [self.get(key) for key in resolve_target_key_list(raw_targets, self._ontology.activities)]

    def by_activity_name(self, activity_name: str) -> list[DeliveryPlanRow]:
        key = normalize_activity_key(activity_name)
        rows = [
            self.get(activity.row_key)
            for activity in self._ontology.activities.values()
            if normalize_activity_key(activity.name) == key
        ]
        return sorted(rows, key=lambda row: row.activity.row_order)


class DeliveryPodSet:
    def __init__(self, ontology: LocalScheduleOntology) -> None:
        self._ontology = ontology

    def __iter__(self) -> Iterator[DeliveryPod]:
        return iter(self.all())

    def __len__(self) -> int:
        return len(self._ontology.pod_names)

    def all(self) -> list[DeliveryPod]:
        return [self.get(pod_name) for pod_name in self._ontology.pod_names]

    def get(self, pod_key: str) -> DeliveryPod:
        pod_name = self._ontology.to_pod_name(pod_key)
        if pod_name not in self._ontology.pod_memberships:
            raise KeyError(f"Unknown DeliveryPod podKey: {pod_key}")
        return DeliveryPod(self._ontology, pod_name)

    def get_or_none(self, pod_key: str) -> DeliveryPod | None:
        try:
            return self.get(pod_key)
        except KeyError:
            return None

    def by_management_unit(self, management_unit: str) -> DeliveryPod | None:
        pod_name = normalize_pod_name(management_unit)
        return self.get_or_none(make_pod_key(self._ontology.project_key, pod_name))


class ScheduleObjects:
    def __init__(self, ontology: LocalScheduleOntology) -> None:
        self.DeliveryProject = DeliveryProjectSet(ontology)
        self.DeliveryPod = DeliveryPodSet(ontology)
        self.DeliveryPlanRow = DeliveryPlanRowSet(ontology)


class DeliveryProjectSet:
    def __init__(self, ontology: LocalScheduleOntology) -> None:
        self._ontology = ontology

    def __iter__(self) -> Iterator[DeliveryProject]:
        return iter(self.all())

    def __len__(self) -> int:
        return 1

    def all(self) -> list[DeliveryProject]:
        return [self.get(self._ontology.project_key)]

    def get(self, project_key: str) -> DeliveryProject:
        if project_key != self._ontology.project_key:
            raise KeyError(f"Unknown DeliveryProject projectKey: {project_key}")
        return self._ontology.project()

    def get_or_none(self, project_key: str) -> DeliveryProject | None:
        if project_key != self._ontology.project_key:
            return None
        return self.get(project_key)


class ScheduleLinks:
    def __init__(self, ontology: LocalScheduleOntology) -> None:
        self._ontology = ontology

    def depends_on(self, row_or_key: DeliveryPlanRow | str) -> list[DeliveryPlanRow]:
        return self._ontology.row(row_or_key).depends_on()

    def is_prerequisite_for(self, row_or_key: DeliveryPlanRow | str) -> list[DeliveryPlanRow]:
        return self._ontology.row(row_or_key).is_prerequisite_for()

    def plan_rows(self, project_or_key: DeliveryProject | str) -> list[DeliveryPlanRow]:
        project = self._ontology.project(project_or_key)
        return project.plan_rows()

    def project(self, row_or_key: DeliveryPlanRow | str) -> DeliveryProject:
        return self._ontology.row(row_or_key).project()

    def pods(self, row_or_project_or_key: DeliveryPlanRow | DeliveryProject | str) -> list[DeliveryPod]:
        if isinstance(row_or_project_or_key, DeliveryProject):
            return row_or_project_or_key.pods()
        if isinstance(row_or_project_or_key, DeliveryPlanRow):
            return row_or_project_or_key.pods()
        if row_or_project_or_key == self._ontology.project_key:
            return self._ontology.project().pods()
        return self._ontology.row(row_or_project_or_key).pods()


class ScheduleFunctions:
    def __init__(self, ontology: LocalScheduleOntology) -> None:
        self._ontology = ontology

    def forward_schedule(
        self,
        *,
        explicit_anchors: dict[str, Any],
        use_plan_roots: bool = True,
        compute: Callable[[dict[str, Activity], dict[str, Any], bool, list[str]], tuple[dict[str, Any], set[str]]] | None = None,
    ) -> ScheduleFunctionRun:
        if compute is None:
            from llm_openrouter_forwardschedule import compute_forward_schedule as compute

        warnings = self._ontology.copy_warnings()
        local_anchors = {
            self._ontology.to_local_row_key(row_key): anchor
            for row_key, anchor in explicit_anchors.items()
        }
        results, root_fallback_keys = compute(
            self._ontology.activities,
            local_anchors,
            use_plan_roots,
            warnings,
        )
        return ScheduleFunctionRun(
            results=results,
            warnings=warnings,
            metadata={
                "project_key": self._ontology.project_key,
                "root_fallback_keys": root_fallback_keys,
                "global_result_keys": {
                    local_key: self._ontology.to_global_row_key(local_key)
                    for local_key in results
                },
            },
        )

    def back_schedule(
        self,
        *,
        targets: str | Iterable[str],
        anchor_date: date,
        compute: Callable[[dict[str, Activity], list[str], date, list[str]], dict[str, Any]] | None = None,
    ) -> ScheduleFunctionRun:
        if compute is None:
            from llm_openrouter_backschedule import compute_back_schedule as compute

        if isinstance(targets, str):
            target_values = resolve_target_key_list(targets, self._ontology.activities)
        else:
            target_values = [
                self._ontology.to_local_row_key(target)
                if self._ontology.to_local_row_key(target) in self._ontology.activities
                else resolve_user_activity_key(target, self._ontology.activities)
                for target in targets
            ]

        target_keys: list[str] = []
        seen: set[str] = set()
        for target_key in target_values:
            if target_key in seen:
                continue
            target_keys.append(target_key)
            seen.add(target_key)

        warnings = self._ontology.copy_warnings()
        results = compute(self._ontology.activities, target_keys, anchor_date, warnings)
        return ScheduleFunctionRun(
            results=results,
            warnings=warnings,
            metadata={
                "project_key": self._ontology.project_key,
                "target_keys": target_keys,
                "global_target_keys": [
                    self._ontology.to_global_row_key(target_key)
                    for target_key in target_keys
                ],
                "global_result_keys": {
                    local_key: self._ontology.to_global_row_key(local_key)
                    for local_key in results
                },
            },
        )


class LocalScheduleOntology:
    def __init__(
        self,
        *,
        rows: Iterable[dict[str, str]],
        activities: dict[str, Activity],
        warnings: Iterable[str],
        source_path: Path | None = None,
        project_key: str | None = None,
        project_name: str | None = None,
        plan_version: str = "",
        project_description: str = "",
        effective_standard_duration_by_key: dict[str, str] | None = None,
        effective_limit_duration_by_key: dict[str, str] | None = None,
    ) -> None:
        self.source_path = source_path
        self.project_key = normalize_project_key(project_key or default_project_key(source_path))
        self.project_name = project_name or (source_path.stem if source_path else self.project_key)
        self.plan_version = plan_version
        self.project_description = project_description
        self.effective_standard_duration_by_key = dict(effective_standard_duration_by_key or {})
        self.effective_limit_duration_by_key = dict(effective_limit_duration_by_key or {})
        if source_path and source_path.exists():
            self.source_updated_at = datetime.fromtimestamp(source_path.stat().st_mtime).isoformat(timespec="seconds")
        else:
            self.source_updated_at = ""
        self.rows = list(rows)
        self.activities = activities
        self.warnings = list(warnings)
        self.successors = build_successors(activities)
        self.pod_memberships = self._build_pod_memberships()
        self.pod_names = sorted(self.pod_memberships)
        self.raw_rows_by_key = {
            activity.key: self.rows[activity.row_order]
            for activity in activities.values()
            if 0 <= activity.row_order < len(self.rows)
        }
        self.objects = ScheduleObjects(self)
        self.links = ScheduleLinks(self)
        self.functions = ScheduleFunctions(self)

    @classmethod
    def from_rows(
        cls,
        rows: Iterable[dict[str, str]],
        *,
        source_path: Path | None = None,
        project_key: str | None = None,
        project_name: str | None = None,
        plan_version: str = "",
        project_description: str = "",
        effective_sla_by_row_key: dict[str, object] | None = None,
        effective_limit_duration_by_row_key: dict[str, object] | None = None,
    ) -> LocalScheduleOntology:
        prepared_rows = list(rows)
        effective_project_key = normalize_project_key(project_key or default_project_key(source_path))
        activities, warnings = build_activities(prepared_rows)
        effective_standard_duration_by_key = _apply_effective_sla_overrides(
            activities,
            _localize_duration_overrides(effective_sla_by_row_key, effective_project_key),
            warnings,
        )
        effective_limit_duration_by_key = _localize_duration_overrides(
            effective_limit_duration_by_row_key,
            effective_project_key,
        )
        return cls(
            rows=prepared_rows,
            activities=activities,
            warnings=warnings,
            source_path=source_path,
            project_key=effective_project_key,
            project_name=project_name,
            plan_version=plan_version,
            project_description=project_description,
            effective_standard_duration_by_key=effective_standard_duration_by_key,
            effective_limit_duration_by_key=effective_limit_duration_by_key,
        )

    @classmethod
    def from_file(
        cls,
        path: Path,
        *,
        project_key: str | None = None,
        project_name: str | None = None,
        plan_version: str = "",
        project_description: str = "",
        effective_sla_by_row_key: dict[str, object] | None = None,
        effective_limit_duration_by_row_key: dict[str, object] | None = None,
    ) -> LocalScheduleOntology:
        source_path = path.expanduser().resolve()
        if not source_path.exists():
            raise RuntimeError(f"File not found: {source_path}")
        return cls.from_rows(
            load_rows(source_path),
            source_path=source_path,
            project_key=project_key,
            project_name=project_name,
            plan_version=plan_version,
            project_description=project_description,
            effective_sla_by_row_key=effective_sla_by_row_key,
            effective_limit_duration_by_row_key=effective_limit_duration_by_row_key,
        )

    def row(self, row_or_key: DeliveryPlanRow | str) -> DeliveryPlanRow:
        if isinstance(row_or_key, DeliveryPlanRow):
            return row_or_key
        local_key = self.to_local_row_key(row_or_key)
        try:
            activity = self.activities[local_key]
        except KeyError as exc:
            raise KeyError(f"Unknown DeliveryPlanRow rowKey: {row_or_key}") from exc
        return DeliveryPlanRow(self, activity)

    def project(self, project_or_key: DeliveryProject | str | None = None) -> DeliveryProject:
        if isinstance(project_or_key, DeliveryProject):
            return project_or_key
        if project_or_key is not None and project_or_key != self.project_key:
            raise KeyError(f"Unknown DeliveryProject projectKey: {project_or_key}")
        return DeliveryProject(self)

    def _build_pod_memberships(self) -> dict[str, list[str]]:
        memberships: defaultdict[str, list[str]] = defaultdict(list)
        for activity in self.activities.values():
            for unit in split_management_units(activity.unit):
                pod_name = normalize_pod_name(unit)
                if activity.key not in memberships[pod_name]:
                    memberships[pod_name].append(activity.key)
        return {
            pod_name: sorted(keys, key=lambda key: self.activities[key].row_order)
            for pod_name, keys in memberships.items()
        }

    def rows_for_pod(self, pod_or_key: DeliveryPod | str) -> list[DeliveryPlanRow]:
        if isinstance(pod_or_key, DeliveryPod):
            pod_name = pod_or_key.pod_name
        else:
            pod_name = self.to_pod_name(pod_or_key)
        return [
            self.row(local_key)
            for local_key in self.pod_memberships.get(pod_name, [])
        ]

    def pods_for_row(self, row_or_key: DeliveryPlanRow | str) -> list[DeliveryPod]:
        row = self.row(row_or_key)
        return [
            DeliveryPod(self, normalize_pod_name(unit))
            for unit in split_management_units(row.management_unit)
            if normalize_pod_name(unit) in self.pod_memberships
        ]

    def to_pod_name(self, pod_key: str) -> str:
        prefix = f"{self.project_key}{ROW_KEY_SEP}"
        if pod_key.startswith(prefix):
            return pod_key[len(prefix):]
        return normalize_pod_name(pod_key)

    def to_global_row_key(self, local_row_key: str) -> str:
        local_key = self.to_local_row_key(local_row_key)
        return make_global_row_key(self.project_key, local_key)

    def to_local_row_key(self, row_key: str) -> str:
        prefix = f"{self.project_key}{ROW_KEY_SEP}"
        if row_key.startswith(prefix):
            return row_key[len(prefix):]
        return row_key

    def copy_warnings(self) -> list[str]:
        return list(self.warnings)
