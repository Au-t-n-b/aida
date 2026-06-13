"""Load 维保SLA.xlsx into Dolt as SlaRequirement.

Current source of truth:
  temp/京东三期/早期介入/交付预案/输出结果/维保SLA.xlsx, sheet 维保SLA

Only the new header dialect is accepted:
  问题等级 / 服务覆盖时间 / 响应时间 / 回复时间 / 解决时间 / 硬件支持 / 服务类型 / 预案版本号

`预案版本号` is used only as a load filter: blank and `草稿` rows are skipped, and the
version value is not stored in the ontology row. Historical columns such as `定义` and
`恢复时间` are not kept for compatibility; the loader drops their Dolt columns when present.

Run (Dolt sql-server must be up; DOLT_DATABASE_URL from ontology/.env):
    cd ontology && python load_sla_requirement_to_dolt.py
    # override the source file:
    SLA_REQUIREMENT_XLSX="D:\\path\\to\\维保SLA.xlsx" python load_sla_requirement_to_dolt.py
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

ONTOLOGY_DIR = Path(__file__).resolve().parent
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))

try:  # load DOLT_DATABASE_URL et al. from ontology/.env when run standalone
    from dotenv import load_dotenv

    load_dotenv(ONTOLOGY_DIR / ".env")
except Exception:  # pragma: no cover - dotenv optional
    pass

import openpyxl  # noqa: E402

from dolt_schema_sync import (  # noqa: E402
    _resolve_engine,
    _sql_text,
    build_add_column_statements,
    build_create_table_sql,
    object_type_table_name,
    quote_ident,
    validate_sql_identifier,
)


OBJECT_TYPE = "SlaRequirement"
DEFAULT_EXCEL_PATH = Path(
    r"D:\Code\aida\temp\京东三期\早期介入\交付预案\输出结果\维保SLA.xlsx"
)
EXCEL_PATH = Path(os.environ.get("SLA_REQUIREMENT_XLSX") or DEFAULT_EXCEL_PATH)
SHEET = "维保SLA"
OBJECT_TYPES_PATH = ONTOLOGY_DIR / "schema" / "object-types.json"

HEADER_MAP = {
    "问题等级": "issueLevel",
    "服务覆盖时间": "coverageWindow",
    "响应时间": "responseTime",
    "回复时间": "replyTime",
    "解决时间": "resolutionTime",
    "硬件支持": "hardwareSupport",
    "服务类型": "serviceType",
}
VERSION_HEADER = "预案版本号"
REQUIRED_HEADERS = tuple(HEADER_MAP) + (VERSION_HEADER,)
STALE_COLUMNS = ("definition", "recoveryTime")


def _load_object_schema() -> dict:
    schemas = json.loads(OBJECT_TYPES_PATH.read_text(encoding="utf-8"))
    schema = schemas.get(OBJECT_TYPE)
    if not isinstance(schema, dict):
        raise SystemExit(f"ObjectType {OBJECT_TYPE} not found in {OBJECT_TYPES_PATH}")
    return schema


def _schema_columns(schema: dict) -> list[str]:
    columns: list[str] = []
    for fallback, prop in (schema.get("properties") or {}).items():
        if not isinstance(prop, dict):
            continue
        column = str(prop.get("apiName") or fallback).strip()
        if column:
            validate_sql_identifier(column, kind="column")
            columns.append(column)
    return columns


def _validate_headers(headers: list[str]) -> None:
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise SystemExit(
            f"Missing required headers in {EXCEL_PATH.name}: {', '.join(missing)}; got {headers}"
        )


def _is_non_draft_version(value: object) -> bool:
    version = str(value).strip() if value is not None else ""
    return bool(version) and version != "草稿"


def _read_rows() -> list[dict]:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in raw_rows[0]]
    _validate_headers(headers)
    idx = {header: i for i, header in enumerate(headers)}
    version_i = idx[VERSION_HEADER]

    rows: list[dict] = []
    for raw in raw_rows[1:]:
        if all(value is None or str(value).strip() == "" for value in raw):
            continue
        if version_i >= len(raw) or not _is_non_draft_version(raw[version_i]):
            continue

        row_number = len(rows) + 1
        record: dict = {
            "slaRequirementId": f"sla_requirement-{row_number:04d}",
            "status": "ACTIVE",
            "sourceFile": EXCEL_PATH.name,
        }
        for header, api_name in HEADER_MAP.items():
            col_i = idx[header]
            value = raw[col_i] if col_i < len(raw) else None
            if value is not None and str(value).strip() != "":
                record[api_name] = str(value).strip()
        issue_level = str(record.get("issueLevel") or "NA").strip()
        service_type = str(record.get("serviceType") or "SLA").strip()
        record["slaName"] = f"{service_type}-{issue_level}-{row_number:02d}"
        rows.append(record)
    return rows


def _sql_value(value):
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    return value


def _fetch_column_names(conn, table_name: str) -> set[str]:
    result = conn.execute(
        _sql_text(
            "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = :t"
        ),
        {"t": table_name},
    )
    return {str(row[0]) for row in result.fetchall() if row[0] is not None}


def _drop_stale_columns(conn, table_name: str, existing_columns: set[str]) -> list[str]:
    dropped: list[str] = []
    for column in STALE_COLUMNS:
        if column not in existing_columns:
            continue
        conn.execute(_sql_text(f"ALTER TABLE {quote_ident(table_name)} DROP COLUMN {quote_ident(column)}"))
        dropped.append(column)
    return dropped


def main() -> int:
    if not EXCEL_PATH.exists():
        raise SystemExit(f"Source workbook not found: {EXCEL_PATH}")

    schema = _load_object_schema()
    table_name = object_type_table_name(OBJECT_TYPE, schema)
    columns = _schema_columns(schema)
    rows = _read_rows()
    if not rows:
        raise SystemExit(f"No rows read from {EXCEL_PATH}")

    engine = _resolve_engine(None)
    if engine is None:
        raise SystemExit(
            "Dolt is not reachable. Set DOLT_DATABASE_URL (ontology/.env) and start the dolt sql-server."
        )

    with engine.connect() as conn:
        conn.execute(_sql_text(build_create_table_sql(OBJECT_TYPE, schema)))
        existing = _fetch_column_names(conn, table_name)
        add_statements, _added = build_add_column_statements(table_name, schema, existing)
        for statement in add_statements:
            conn.execute(_sql_text(statement))
        if add_statements:
            existing = _fetch_column_names(conn, table_name)
        dropped = _drop_stale_columns(conn, table_name, existing)
        conn.commit()

        conn.execute(_sql_text(f"DELETE FROM {quote_ident(table_name)}"))
        column_sql = ", ".join(quote_ident(c) for c in columns)
        values_sql = ", ".join(f":{c}" for c in columns)
        insert_sql = f"INSERT INTO {quote_ident(table_name)} ({column_sql}) VALUES ({values_sql})"
        for row in rows:
            conn.execute(_sql_text(insert_sql), {c: _sql_value(row.get(c)) for c in columns})
        conn.commit()

        try:
            conn.execute(
                _sql_text("CALL dolt_commit('-A', '-m', :m)"),
                {"m": f"Load {table_name} from {EXCEL_PATH.name} ({len(rows)} rows)"},
            )
            conn.commit()
        except Exception as exc:  # nothing-to-commit is fine
            if "nothing to commit" not in str(exc).lower():
                raise

    print(f"Loaded {len(rows)} rows into Dolt `{table_name}` from {EXCEL_PATH.name}:")
    if dropped:
        print(f"  dropped stale columns: {', '.join(dropped)}")
    for row in rows:
        print(
            f"  {row['slaRequirementId']:24} {row.get('serviceType',''):8} "
            f"{row.get('issueLevel',''):8} coverage={row.get('coverageWindow','')} "
            f"response={row.get('responseTime','')} reply={row.get('replyTime','')} "
            f"resolution={row.get('resolutionTime','')}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
