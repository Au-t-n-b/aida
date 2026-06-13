"""Tests for publishing derived contingency findings into project-management Excel files.

Run (PowerShell, repo root):
    ontology\\.venv\\Scripts\\python.exe ontology\\tests\\test_project_management_excel_publish.py -v
"""

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


RISK_HEADERS = [
    "id",
    "modified",
    "project_id",
    "issue_id",
    "name",
    "rule_id",
    "description",
    "level",
    "owner",
    "effect",
    "progress",
    "status",
    "start_time",
    "end_time",
    "planned_complete_time",
    "agent_name",
    "risk_type",
    "closed_notes",
    "SOURCE",
    "risk_code",
    "source_system_url",
    "CREATED_AT",
    "UPDATED_AT",
]

RISK_MEASURE_HEADERS = [
    "MEASURE_ID",
    "FOREIGN_ID",
    "task_id",
    "PROJECT_ID",
    "ISSUE_ID",
    "MEASURE_NAME",
    "RESPONSIBLE_PERSON",
    "PLANNED_COMPLETE_TIME",
    "REAL_COMPLETE_TIME",
    "PROGRESS",
    "STATUS",
]

TASK_HEADERS = [
    "任务ID",
    "编号",
    "*任务名称",
    "任务描述",
    "*责任人",
    "标签",
    "优先级",
    "状态",
    "下发时间",
    "*计划完成时间",
    "实际完成时间",
    "进展",
    "操作人员",
    "关联的风险",
    "备注",
]

TASK_PROGRESS_HEADERS = ["id", "task_id", "project_id", "dispatch_time", "progress", "operator"]

ISSUE_HEADERS = [
    "ID",
    "CUSTOM_CODE",
    "PROJECT_ID",
    "DESCRIPTION",
    "PRIMARY_ISSUE_TYPE",
    "SECONDARY_ISSUE_TYPE",
    "SEVERITY",
    "SOURCE",
    "STATUS",
    "ASSIGNEE_ID",
    "PLANNED_COMPLETE_AT",
    "ACTUAL_CLOSED_AT",
    "CREATED_BY",
    "CREATED_AT",
    "UPDATED_AT",
    "AGENT_NAME",
]

ISSUE_PROGRESS_HEADERS = ["ID", "PROBLEM_ID", "CONTENT", "LOG_TYPE", "CREATED_BY", "CREATED_AT"]

PROJECT_ID = "6336ff05cffc4144aa6d96a87e66a36e"
NOW = datetime(2026, 6, 12, 9, 30, 0)


def _write_book(path: Path, headers: list[str], rows: list[list[object]] | None = None, *, sheet: str = "Sheet1", blanks: int = 0) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows or []:
        ws.append(row)
    for _ in range(blanks):
        ws.append([None] * len(headers))
    wb.save(path)


def _make_target_tree(root: Path) -> Path:
    base = root / "项目管理"
    _write_book(
        base / "风险" / "输出结果" / "风险信息.xlsx",
        RISK_HEADERS,
        [[100, None, PROJECT_ID, "seed-issue", "seed", None, None, "低", None, None, None, None, None, None, None, None, None, None, None, "R-SEED", None, None, None]],
    )
    _write_book(base / "风险" / "输出结果" / "风险信息_应对措施.xlsx", RISK_MEASURE_HEADERS, [[200, 100, None, PROJECT_ID, "seed-issue", "seed measure", None, None, None, None, None]])
    _write_book(
        base / "任务" / "输出结果" / "任务信息.xlsx",
        TASK_HEADERS,
        [[300, "TASK-SEED", "seed task", None, "PM", None, "低", "处理中", None, None, None, None, None, "R-SEED", None]],
        sheet="任务详情",
        blanks=5,
    )
    _write_book(base / "任务" / "输出结果" / "任务信息_历史进展.xlsx", TASK_PROGRESS_HEADERS, [[400, 300, PROJECT_ID, None, "seed progress", "tester"]])
    _write_book(base / "问题" / "输出结果" / "问题信息.xlsx", ISSUE_HEADERS, [[500, "ISSUE-SEED", PROJECT_ID, "seed issue", None, None, "low", None, "processing", None, None, None, None, None, None, None]])
    _write_book(base / "问题" / "输出结果" / "问题信息_历史进展.xlsx", ISSUE_PROGRESS_HEADERS, [[600, 500, "seed issue progress", "feedback", "tester", None]])
    return base


def _table_rows(path: Path, *, sheet: str | None = None) -> list[dict[str, object]]:
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    headers = [str(cell.value) for cell in ws[1]]
    rows: list[dict[str, object]] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and str(value) != "" for value in raw):
            continue
        rows.append(dict(zip(headers, raw)))
    return rows


def _sample_generation() -> dict[str, object]:
    return {
        "planId": "JD_A3_CONTINGENCY",
        "projectKey": "京东",
        "assessmentId": "JD_A3_DELIVERABILITY",
        "riskCount": 2,
        "risks": [
            {
                "riskId": "R-HARD",
                "riskName": "EOM生命周期硬阻塞",
                "riskPoint": "EOM风险",
                "riskType": "生命周期风险",
                "severity": "重大",
                "owner": "架构师",
                "mitigationPlan": "确认替代版本与升级窗口。",
                "impact": "版本不可交付。",
                "description": "产品生命周期命中 EOM。",
                "provenance": {"rule": {"ruleId": "RULE-LC", "riskSubCategory": "DEV-LC-01"}},
            },
            {
                "riskId": "R-TRACK",
                "riskName": "光模块补齐跟踪",
                "riskPoint": "光模块丢失风险",
                "riskType": "任务跟踪风险",
                "severity": "中",
                "owner": "交付PM",
                "mitigationPlan": "补齐光模块并跟踪采购到货。",
                "impact": "影响联调进度。",
                "description": "光模块配置缺失。",
                "provenance": {"rule": {"ruleId": "RULE-TRK", "riskSubCategory": "TRK-04"}},
            },
        ],
    }


class ProjectManagementExcelPublishTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.base = _make_target_tree(Path(tmp.name))

    def _target(self):
        import project_management_excel

        return project_management_excel.ProjectManagementExcelTarget(
            project_key="京东",
            base_dir=self.base,
            file_center_project_id=PROJECT_ID,
        )

    def test_publish_writes_risk_task_issue_rows_and_is_idempotent(self):
        import project_management_excel

        first = project_management_excel.publish_findings_to_project_management_excel(
            _sample_generation(),
            target=self._target(),
            write=True,
            now=NOW,
        )
        self.assertEqual(first["riskCount"], 2)
        self.assertEqual(first["riskRowsWritten"], 2)
        self.assertEqual(first["taskRowsWritten"], 1)
        self.assertEqual(first["issueRowsWritten"], 1)
        self.assertEqual(len(first["updatedFiles"]), 6)

        risk_rows = _table_rows(self.base / "风险" / "输出结果" / "风险信息.xlsx")
        self.assertEqual([row["risk_code"] for row in risk_rows].count("R-HARD"), 1)
        self.assertEqual([row["risk_code"] for row in risk_rows].count("R-TRACK"), 1)
        self.assertEqual(next(row for row in risk_rows if row["risk_code"] == "R-HARD")["level"], "高")

        task_rows = _table_rows(self.base / "任务" / "输出结果" / "任务信息.xlsx", sheet="任务详情")
        track_task = next(row for row in task_rows if row["关联的风险"] == "R-TRACK")
        self.assertEqual(track_task["任务ID"], 301)
        self.assertEqual(track_task["*任务名称"], "处置：光模块补齐跟踪")
        self.assertEqual(track_task["标签"], "预案决策")

        issue_rows = _table_rows(self.base / "问题" / "输出结果" / "问题信息.xlsx")
        hard_issue = next(row for row in issue_rows if row["CUSTOM_CODE"] != "ISSUE-SEED")
        self.assertEqual(hard_issue["SEVERITY"], "high")
        self.assertEqual(hard_issue["SECONDARY_ISSUE_TYPE"], "EOM风险")

        project_management_excel.publish_findings_to_project_management_excel(
            _sample_generation(),
            target=self._target(),
            write=True,
            now=NOW,
        )
        self.assertEqual(
            [row["risk_code"] for row in _table_rows(self.base / "风险" / "输出结果" / "风险信息.xlsx")].count("R-HARD"),
            1,
        )
        self.assertEqual(
            [row["关联的风险"] for row in _table_rows(self.base / "任务" / "输出结果" / "任务信息.xlsx", sheet="任务详情")].count("R-TRACK"),
            1,
        )
        self.assertEqual(
            sum(1 for row in _table_rows(self.base / "问题" / "输出结果" / "问题信息.xlsx") if row["CUSTOM_CODE"] != "ISSUE-SEED"),
            1,
        )

    def test_action_validate_dry_run_and_apply_route_to_excel_publisher(self):
        import backend_app
        import data_connector
        import project_management_excel

        payload = {
            "planId": "JD_A3_CONTINGENCY",
            "projectKey": "京东",
            "assessmentId": "JD_A3_DELIVERABILITY",
            "referenceDate": "2026-06-12",
        }
        with (
            mock.patch.object(data_connector, "derive_contingency_risks_for_plan", return_value=_sample_generation()),
            mock.patch.object(project_management_excel, "resolve_project_management_excel_target", return_value=self._target()),
        ):
            dry_run = backend_app.validate_action_type("PublishContingencyFindings", payload)
            self.assertEqual(dry_run["riskRowsWritten"], 2)
            self.assertEqual(len(_table_rows(self.base / "风险" / "输出结果" / "风险信息.xlsx")), 1)

            applied = backend_app.apply_action_type("PublishContingencyFindings", payload)
            self.assertEqual(applied["riskRowsWritten"], 2)
            self.assertEqual(len(_table_rows(self.base / "风险" / "输出结果" / "风险信息.xlsx")), 3)


if __name__ == "__main__":
    unittest.main()
