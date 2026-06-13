import importlib
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


ONTOLOGY_DIR = Path(__file__).resolve().parents[1]
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))


NEW_HEADERS = [
    "问题等级",
    "服务覆盖时间",
    "响应时间",
    "回复时间",
    "解决时间",
    "硬件支持",
    "服务类型",
    "预案版本号",
]


def _build_book(rows: list[list[object]]) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "维保SLA"
    for row in rows:
        worksheet.append(row)
    return workbook


class SlaRequirementLoaderTests(unittest.TestCase):
    def test_loader_uses_new_headers_and_skips_drafts(self):
        loader = importlib.import_module("load_sla_requirement_to_dolt")
        workbook = _build_book(
            [
                NEW_HEADERS,
                ["紧急问题", "7×24", "1 小时", "30 分钟", "2 小时", "现场+远程", "标准+", "草稿"],
                ["重要问题", "5×8", "2 小时", "1 小时", "4 小时", "", "标准+", ""],
                ["一般问题", "5×8", "4 小时", "2 小时", "8 小时", "远程", "标准+", "V1.0"],
            ]
        )

        old_path = loader.EXCEL_PATH
        loader.EXCEL_PATH = Path("维保SLA.xlsx")
        try:
            with patch.object(loader.openpyxl, "load_workbook", return_value=workbook):
                rows = loader._read_rows()
        finally:
            loader.EXCEL_PATH = old_path

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["issueLevel"], "一般问题")
        self.assertEqual(rows[0]["coverageWindow"], "5×8")
        self.assertEqual(rows[0]["responseTime"], "4 小时")
        self.assertEqual(rows[0]["replyTime"], "2 小时")
        self.assertEqual(rows[0]["resolutionTime"], "8 小时")
        self.assertEqual(rows[0]["hardwareSupport"], "远程")
        self.assertEqual(rows[0]["serviceType"], "标准+")
        self.assertNotIn("proposalVersion", rows[0])
        self.assertNotIn("planVersion", rows[0])
        self.assertNotIn("recoveryTime", rows[0])
        self.assertNotIn("definition", rows[0])

    def test_loader_rejects_historical_headers(self):
        loader = importlib.import_module("load_sla_requirement_to_dolt")
        workbook = _build_book(
            [
                ["问题级别", "定义", "覆盖时段", "响应时间", "恢复时间", "解决时间", "硬件支持"],
                ["紧急问题", "旧定义", "7×24", "1 小时", "2 小时", "4 小时", "远程"],
            ]
        )

        old_path = loader.EXCEL_PATH
        loader.EXCEL_PATH = Path("old.xlsx")
        try:
            with patch.object(loader.openpyxl, "load_workbook", return_value=workbook):
                with self.assertRaisesRegex(SystemExit, "Missing required headers"):
                    loader._read_rows()
        finally:
            loader.EXCEL_PATH = old_path

    def test_schema_uses_new_sla_columns_only(self):
        object_types = json.loads((ONTOLOGY_DIR / "schema" / "object-types.json").read_text(encoding="utf-8"))
        properties = object_types["SlaRequirement"]["properties"]

        expected_source_columns = {
            "issueLevel": "问题等级",
            "coverageWindow": "服务覆盖时间",
            "responseTime": "响应时间",
            "replyTime": "回复时间",
            "resolutionTime": "解决时间",
            "hardwareSupport": "硬件支持",
            "serviceType": "服务类型",
        }
        for field, source_column in expected_source_columns.items():
            self.assertEqual(properties[field].get("sourceColumnName"), source_column)

        self.assertNotIn("definition", properties)
        self.assertNotIn("recoveryTime", properties)
        self.assertNotIn("proposalVersion", properties)
        self.assertNotIn("planVersion", properties)
