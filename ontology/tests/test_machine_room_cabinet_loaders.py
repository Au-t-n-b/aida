import importlib
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


ONTOLOGY_DIR = Path(__file__).resolve().parents[1]
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))


NEW_HEADERS = [
    "PoD名称",
    "机房名称",
    "计算柜",
    "总线柜",
    "参数面Leaf柜",
    "业务面Leaf柜",
    "管理面柜",
    "样本面Leaf柜",
    "预案版本号",
]


def _build_book(rows: list[list[object]]) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    for row in rows:
        worksheet.append(row)
    return workbook


class MachineRoomCabinetLoaderTests(unittest.TestCase):
    def test_pod_cabinet_layout_loader_uses_new_headers_and_skips_drafts(self):
        loader = importlib.import_module("load_pod_cabinet_layout_to_dolt")
        workbook = _build_book(
            [
                NEW_HEADERS,
                ["POD1", "401", "A01,A02", "A03", "A04", "A05", "A06", "A07", "草稿"],
                ["POD1", "401", "A01,A02", "A03", "A04", "A05", "A06", "A07", "V1.0"],
            ]
        )

        old_path = loader.EXCEL_PATH
        loader.EXCEL_PATH = Path("机房机柜信息表.xlsx")
        try:
            with patch.object(loader.openpyxl, "load_workbook", return_value=workbook):
                rows = loader._read_pod_rows()
        finally:
            loader.EXCEL_PATH = old_path

        self.assertEqual(len(rows), 1)
        self.assertNotIn("proposalVersion", rows[0])
        self.assertNotIn("planVersion", rows[0])
        self.assertEqual(rows[0]["podName"], "POD1")
        self.assertEqual(rows[0]["machineRoomName"], "401")
        self.assertEqual(rows[0]["busbarCabinets"], "A03")
        self.assertEqual(rows[0]["paramLeafCabinets"], "A04")
        self.assertEqual(rows[0]["businessLeafCabinets"], "A05")
        self.assertEqual(rows[0]["mgmtCabinets"], "A06")
        self.assertEqual(rows[0]["sampleLeafCabinets"], "A07")

    def test_pod_cabinet_layout_loader_rejects_historical_headers(self):
        loader = importlib.import_module("load_pod_cabinet_layout_to_dolt")
        workbook = _build_book(
            [
                ["PoD", "机房", "计算柜", "母线柜", "参数Leaf柜", "业务Leaf柜", "网管柜", "样本Leaf柜"],
                ["POD1", "401", "A01", "A03", "A04", "A05", "A06", "A07"],
            ]
        )

        old_path = loader.EXCEL_PATH
        loader.EXCEL_PATH = Path("old.xlsx")
        try:
            with patch.object(loader.openpyxl, "load_workbook", return_value=workbook):
                with self.assertRaisesRegex(SystemExit, "Missing required headers"):
                    loader._read_pod_rows()
        finally:
            loader.EXCEL_PATH = old_path

    def test_machine_room_config_loader_filters_non_draft_rows_without_version_fields(self):
        loader = importlib.import_module("load_machine_room_config_to_dolt")
        workbook = _build_book(
            [
                NEW_HEADERS,
                ["POD1", "401", "A01,A02", "A03", "A04", "A05", "A06", "", "草稿"],
                ["POD1", "401", "A01,A02", "A03", "A04", "A05", "A06", "", "V1.0"],
            ]
        )

        old_path = loader.EXCEL_PATH
        loader.EXCEL_PATH = Path("机房机柜信息表.xlsx")
        try:
            with patch.object(loader.openpyxl, "load_workbook", return_value=workbook):
                rows = loader._read_rows()
        finally:
            loader.EXCEL_PATH = old_path

        self.assertEqual([row["cabinetType"] for row in rows], ["计算柜", "总线柜", "参数面Leaf柜", "业务面Leaf柜", "管理面柜"])
        self.assertTrue(all("proposalVersion" not in row and "planVersion" not in row for row in rows))
        self.assertEqual(rows[0]["machineRoomId"], "MRC-POD1-COMP")
        self.assertEqual(rows[0]["quantity"], 2)
        self.assertEqual(rows[0]["remark"], "A01,A02")

    def test_pod_cabinet_layout_schema_uses_new_source_column_names(self):
        object_types = json.loads((ONTOLOGY_DIR / "schema" / "object-types.json").read_text(encoding="utf-8"))
        properties = object_types["PodCabinetLayout"]["properties"]

        expected_source_columns = {
            "podName": "PoD名称",
            "machineRoomName": "机房名称",
            "computeCabinets": "计算柜",
            "busbarCabinets": "总线柜",
            "paramLeafCabinets": "参数面Leaf柜",
            "businessLeafCabinets": "业务面Leaf柜",
            "mgmtCabinets": "管理面柜",
            "sampleLeafCabinets": "样本面Leaf柜",
        }
        for field, source_column in expected_source_columns.items():
            self.assertEqual(properties[field].get("sourceColumnName"), source_column)

        self.assertEqual(properties["busbarCabinets"]["displayMetadata"]["displayName"], "总线柜")
        self.assertEqual(properties["mgmtCabinets"]["displayMetadata"]["displayName"], "管理面柜")
