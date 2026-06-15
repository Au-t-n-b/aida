from __future__ import annotations

import hashlib
import os
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


BASE_DIR = Path(__file__).resolve().parent
PROJECT_KEY_JD = "京东"
JD_PROJECT_ID = "6336ff05cffc4144aa6d96a87e66a36e"
AGENT_NAME = "ontology_decision_agent"
AUTO_SOURCE = "本体决策自动识别"
AUTO_PROGRESS_PREFIX = "由数字孪生预案决策自动生成"

RISK_FILE = Path("风险") / "输出结果" / "风险信息.xlsx"
RISK_MEASURE_FILE = Path("风险") / "输出结果" / "风险信息_应对措施.xlsx"
TASK_FILE = Path("任务") / "输出结果" / "任务信息.xlsx"
TASK_PROGRESS_FILE = Path("任务") / "输出结果" / "任务信息_历史进展.xlsx"
ISSUE_FILE = Path("问题") / "输出结果" / "问题信息.xlsx"
ISSUE_PROGRESS_FILE = Path("问题") / "输出结果" / "问题信息_历史进展.xlsx"

EXPECTED_FILES = (
    RISK_FILE,
    RISK_MEASURE_FILE,
    TASK_FILE,
    TASK_PROGRESS_FILE,
    ISSUE_FILE,
    ISSUE_PROGRESS_FILE,
)

LIFECYCLE_RISK_POINTS = {"EOM风险", "TR5风险", "EOS风险", "ESS风险"}
TRACKING_RISK_POINTS = {
    "海外伙伴维护技能不足风险",
    "首个智算验收标准未确认风险",
    "海外伙伴交付能力不足风险",
    "光模块丢失风险",
    "分包商能力不足风险",
}


@dataclass(frozen=True)
class ProjectManagementExcelTarget:
    project_key: str
    base_dir: Path
    file_center_project_id: str

    def path(self, relative_path: Path) -> Path:
        return self.base_dir / relative_path


def resolve_project_management_excel_target(project_key: str) -> ProjectManagementExcelTarget:
    """Resolve a project key to the current project-management Excel target.

    Kept intentionally thin: later file-center migration can swap this resolver without touching
    risk/task/issue mapping logic.
    """
    key = str(project_key or "").strip()
    if key != PROJECT_KEY_JD:
        raise ValueError(f"Unsupported projectKey `{project_key}` for project-management Excel publish.")
    target = ProjectManagementExcelTarget(
        project_key=PROJECT_KEY_JD,
        base_dir=BASE_DIR / "contingency-data" / "京东三期" / "项目管理",
        file_center_project_id=JD_PROJECT_ID,
    )
    missing = [str(target.path(relative)) for relative in EXPECTED_FILES if not target.path(relative).exists()]
    if missing:
        raise FileNotFoundError("Project-management Excel template(s) not found: " + ", ".join(missing))
    return target


def publish_findings_to_project_management_excel(
    generation: dict[str, Any],
    *,
    target: ProjectManagementExcelTarget | None = None,
    write: bool = True,
    now: datetime | None = None,
) -> dict[str, Any]:
    project_key = _text(generation.get("projectKey") or generation.get("project_key"))
    target = target or resolve_project_management_excel_target(project_key)
    risks = [risk for risk in (generation.get("risks") or []) if isinstance(risk, dict)]
    timestamp = (now or datetime.now()).replace(microsecond=0)
    date_text = timestamp.date().isoformat()
    timestamp_text = timestamp.isoformat(sep=" ")

    final_paths = [target.path(relative) for relative in EXPECTED_FILES]
    if write:
        _assert_files_replaceable(final_paths)
    books = _load_workbooks(target)
    risk_ws = books[RISK_FILE].active
    risk_measure_ws = books[RISK_MEASURE_FILE].active
    task_ws = books[TASK_FILE]["任务详情"]
    task_progress_ws = books[TASK_PROGRESS_FILE].active
    issue_ws = books[ISSUE_FILE].active
    issue_progress_ws = books[ISSUE_PROGRESS_FILE].active

    _require_headers(risk_ws, RISK_HEADERS)
    _require_headers(risk_measure_ws, RISK_MEASURE_HEADERS)
    _require_headers(task_ws, TASK_HEADERS)
    _require_headers(task_progress_ws, TASK_PROGRESS_HEADERS)
    _require_headers(issue_ws, ISSUE_HEADERS)
    _require_headers(issue_progress_ws, ISSUE_PROGRESS_HEADERS)

    task_rows_written = 0
    issue_rows_written = 0
    risk_rows_written = 0

    for risk in risks:
        risk_id = _risk_id(risk)
        if not risk_id:
            continue
        risk_row, risk_numeric_id, is_new_risk = _upsert_risk_row(
            risk_ws,
            risk,
            target,
            timestamp_text=timestamp_text,
            date_text=date_text,
        )
        if risk_row and risk_numeric_id is not None:
            risk_rows_written += 1

        is_hard = is_hard_blocker_risk(risk)
        task_numeric_id: int | None = None
        if is_hard:
            issue_numeric_id, _ = _upsert_issue_row(
                issue_ws,
                risk,
                target,
                timestamp_text=timestamp_text,
            )
            _upsert_issue_progress_row(
                issue_progress_ws,
                issue_numeric_id,
                risk,
                timestamp_text=timestamp_text,
            )
            issue_rows_written += 1
        else:
            task_numeric_id, _ = _upsert_task_row(
                task_ws,
                risk,
                target,
                timestamp_text=timestamp_text,
                date_text=date_text,
            )
            _upsert_task_progress_row(
                task_progress_ws,
                task_numeric_id,
                risk,
                target,
                timestamp_text=timestamp_text,
            )
            task_rows_written += 1

        if _text(risk.get("mitigationPlan") or risk.get("mitigation_plan")):
            _upsert_risk_measure_row(
                risk_measure_ws,
                risk,
                target,
                risk_numeric_id,
                task_numeric_id,
            )

    updated_files = [str(target.path(relative)) for relative in EXPECTED_FILES]
    if write:
        _save_workbooks_atomically(books, target)

    return {
        "riskCount": len(risks),
        "riskRowsWritten": risk_rows_written,
        "taskRowsWritten": task_rows_written,
        "issueRowsWritten": issue_rows_written,
        "updatedFiles": updated_files,
        "dryRun": not write,
    }


RISK_HEADERS = [
    "id",
    "modified",
    "project_id",
    "issue_id",
    "name",
    "rule_id",
    "description",
    "level",
    "owner",
    "effect",
    "progress",
    "status",
    "start_time",
    "end_time",
    "planned_complete_time",
    "agent_name",
    "risk_type",
    "closed_notes",
    "SOURCE",
    "risk_code",
    "source_system_url",
    "CREATED_AT",
    "UPDATED_AT",
]

RISK_MEASURE_HEADERS = [
    "MEASURE_ID",
    "FOREIGN_ID",
    "task_id",
    "PROJECT_ID",
    "ISSUE_ID",
    "MEASURE_NAME",
    "RESPONSIBLE_PERSON",
    "PLANNED_COMPLETE_TIME",
    "REAL_COMPLETE_TIME",
    "PROGRESS",
    "STATUS",
]

TASK_HEADERS = [
    "任务ID",
    "编号",
    "*任务名称",
    "任务描述",
    "*责任人",
    "标签",
    "优先级",
    "状态",
    "下发时间",
    "*计划完成时间",
    "实际完成时间",
    "进展",
    "操作人员",
    "关联的风险",
    "备注",
]

TASK_PROGRESS_HEADERS = ["id", "task_id", "project_id", "dispatch_time", "progress", "operator"]

ISSUE_HEADERS = [
    "ID",
    "CUSTOM_CODE",
    "PROJECT_ID",
    "DESCRIPTION",
    "PRIMARY_ISSUE_TYPE",
    "SECONDARY_ISSUE_TYPE",
    "SEVERITY",
    "SOURCE",
    "STATUS",
    "ASSIGNEE_ID",
    "PLANNED_COMPLETE_AT",
    "ACTUAL_CLOSED_AT",
    "CREATED_BY",
    "CREATED_AT",
    "UPDATED_AT",
    "AGENT_NAME",
]

ISSUE_PROGRESS_HEADERS = ["ID", "PROBLEM_ID", "CONTENT", "LOG_TYPE", "CREATED_BY", "CREATED_AT"]


def _load_workbooks(target: ProjectManagementExcelTarget):
    return {relative: load_workbook(target.path(relative)) for relative in EXPECTED_FILES}


def _save_workbooks_atomically(books: dict[Path, Any], target: ProjectManagementExcelTarget) -> None:
    staged: list[tuple[Path, Path]] = []
    final_paths = [target.path(relative) for relative in books]
    try:
        for relative, workbook in books.items():
            final_path = target.path(relative)
            tmp_path = final_path.with_name(f".{final_path.stem}.{uuid.uuid4().hex}.tmp{final_path.suffix}")
            workbook.save(tmp_path)
            staged.append((tmp_path, final_path))
        for workbook in books.values():
            close = getattr(workbook, "close", None)
            if callable(close):
                close()
        _assert_files_replaceable(final_paths)
        for tmp_path, final_path in staged:
            try:
                os.replace(tmp_path, final_path)
            except PermissionError as exc:
                raise PermissionError(
                    f"Cannot update `{final_path}` because the workbook is locked. "
                    "Close the project-management Excel file and retry."
                ) from exc
    finally:
        for tmp_path, _ in staged:
            if tmp_path.exists():
                tmp_path.unlink(missing_ok=True)


def _assert_files_replaceable(paths: list[Path]) -> None:
    for path in paths:
        try:
            with path.open("r+b"):
                pass
        except PermissionError as exc:
            raise PermissionError(
                f"Cannot update `{path}` because the workbook is locked. "
                "Close the project-management Excel file and retry."
            ) from exc


def _header_map(ws: Worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): index
        for index, cell in enumerate(ws[1], start=1)
        if cell.value is not None and str(cell.value).strip()
    }


def _require_headers(ws: Worksheet, headers: list[str]) -> None:
    present = _header_map(ws)
    missing = [header for header in headers if header not in present]
    if missing:
        raise ValueError(f"{ws.title} is missing required header(s): {', '.join(missing)}")


def _is_nonempty(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _last_nonempty_row(ws: Worksheet) -> int:
    last = 1
    for row_idx in range(2, ws.max_row + 1):
        if any(_is_nonempty(ws.cell(row_idx, col_idx).value) for col_idx in range(1, ws.max_column + 1)):
            last = row_idx
    return last


def _append_or_reuse_blank_row(ws: Worksheet) -> int:
    row_idx = _last_nonempty_row(ws) + 1
    if row_idx > ws.max_row:
        _copy_row_style(ws, max(2, row_idx - 1), row_idx)
    return row_idx


def _copy_row_style(ws: Worksheet, source_row: int, target_row: int) -> None:
    if source_row < 1 or source_row == target_row:
        return
    for col_idx in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col_idx)
        dst = ws.cell(target_row, col_idx)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)


def _next_numeric_id(ws: Worksheet, header: str) -> int:
    headers = _header_map(ws)
    col_idx = headers[header]
    current_max = 0
    for row_idx in range(2, ws.max_row + 1):
        value = ws.cell(row_idx, col_idx).value
        if value is None or str(value).strip() == "":
            continue
        try:
            current_max = max(current_max, int(float(str(value))))
        except ValueError:
            continue
    return current_max + 1


def _find_row(ws: Worksheet, header: str, value: Any) -> int | None:
    headers = _header_map(ws)
    col_idx = headers.get(header)
    needle = _text(value)
    if not col_idx or not needle:
        return None
    for row_idx in range(2, ws.max_row + 1):
        if _text(ws.cell(row_idx, col_idx).value) == needle:
            return row_idx
    return None


def _cell(ws: Worksheet, row_idx: int, header: str):
    return ws.cell(row_idx, _header_map(ws)[header])


def _set(ws: Worksheet, row_idx: int, header: str, value: Any) -> None:
    _cell(ws, row_idx, header).value = value


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _risk_id(risk: dict[str, Any]) -> str:
    return _text(risk.get("riskId") or risk.get("risk_id") or risk.get("id"))


def _risk_name(risk: dict[str, Any]) -> str:
    return _text(risk.get("riskName") or risk.get("risk_name") or risk.get("name"))


def _risk_point(risk: dict[str, Any]) -> str:
    return _text(risk.get("riskPoint") or risk.get("risk_point"))


def _risk_type(risk: dict[str, Any]) -> str:
    return _text(risk.get("riskType") or risk.get("risk_type"))


def _risk_description(risk: dict[str, Any]) -> str:
    return _text(risk.get("description") or risk.get("sourceSummary") or risk.get("source_summary"))


def _mitigation(risk: dict[str, Any]) -> str:
    return _text(risk.get("mitigationPlan") or risk.get("mitigation_plan"))


def _owner(risk: dict[str, Any]) -> str:
    return _text(risk.get("owner")) or "交付PM"


def _rule_info(risk: dict[str, Any]) -> dict[str, Any]:
    provenance = risk.get("provenance")
    if not isinstance(provenance, dict):
        return {}
    rule = provenance.get("rule")
    return rule if isinstance(rule, dict) else {}


def _risk_sub_category(risk: dict[str, Any]) -> str:
    return _text(_rule_info(risk).get("riskSubCategory") or _rule_info(risk).get("risk_sub_category"))


def _rule_id(risk: dict[str, Any]) -> str:
    return _text(_rule_info(risk).get("ruleId") or _rule_info(risk).get("rule_id"))


def _stable_digest(risk_id: str) -> str:
    return hashlib.sha1(risk_id.encode("utf-8")).hexdigest().upper()


def _risk_issue_id(risk_id: str) -> str:
    return f"RISK-{_stable_digest(risk_id)[:16]}"


def _task_code(risk_id: str) -> str:
    return f"TASK-{_stable_digest(risk_id)[:12]}"


def _issue_code(risk_id: str) -> str:
    return f"ISSUE-{_stable_digest(risk_id)[:12]}"


def _normalize_risk_level(value: Any) -> str:
    text = _text(value)
    if text in {"重大", "严重", "高", "HIGH", "high"}:
        return "高"
    if text in {"中", "MEDIUM", "medium"}:
        return "中"
    if text in {"低", "LOW", "low"}:
        return "低"
    return text or "中"


def _priority(value: Any) -> str:
    return _normalize_risk_level(value)


def _issue_severity(value: Any) -> str:
    level = _normalize_risk_level(value)
    if level == "高":
        return "high"
    if level == "中":
        return "medium"
    return "low"


def is_hard_blocker_risk(risk: dict[str, Any]) -> bool:
    risk_point = _risk_point(risk)
    if risk_point in TRACKING_RISK_POINTS:
        return False
    category = _risk_sub_category(risk)
    if category.startswith("SVC-MIS-"):
        return True
    if category == "DEV-CMP-01":
        return True
    if category in {f"DEV-LC-0{idx}" for idx in range(1, 6)}:
        return True
    if risk_point in LIFECYCLE_RISK_POINTS:
        return True
    text = " ".join([category, risk_point, _risk_name(risk), _risk_description(risk)])
    return any(keyword in text for keyword in ("服务漏配", "部件不兼容", "端口不一致"))


def _upsert_risk_row(
    ws: Worksheet,
    risk: dict[str, Any],
    target: ProjectManagementExcelTarget,
    *,
    timestamp_text: str,
    date_text: str,
) -> tuple[int, int | None, bool]:
    risk_id = _risk_id(risk)
    row_idx = _find_row(ws, "risk_code", risk_id)
    is_new = row_idx is None
    if row_idx is None:
        row_idx = _append_or_reuse_blank_row(ws)
        risk_numeric_id = _next_numeric_id(ws, "id")
        _set(ws, row_idx, "id", risk_numeric_id)
        _set(ws, row_idx, "CREATED_AT", timestamp_text)
    else:
        risk_numeric_id = _cell(ws, row_idx, "id").value
        if not _text(_cell(ws, row_idx, "CREATED_AT").value):
            _set(ws, row_idx, "CREATED_AT", timestamp_text)

    _set(ws, row_idx, "modified", 1)
    _set(ws, row_idx, "project_id", target.file_center_project_id)
    _set(ws, row_idx, "issue_id", _risk_issue_id(risk_id))
    _set(ws, row_idx, "name", _risk_name(risk))
    _set(ws, row_idx, "rule_id", _rule_id(risk))
    _set(ws, row_idx, "description", _risk_description(risk))
    _set(ws, row_idx, "level", _normalize_risk_level(risk.get("severity")))
    _set(ws, row_idx, "owner", _owner(risk))
    _set(ws, row_idx, "effect", _text(risk.get("impact")))
    _set(ws, row_idx, "progress", "新增")
    _set(ws, row_idx, "status", "处理中")
    _set(ws, row_idx, "start_time", date_text)
    _set(ws, row_idx, "end_time", None)
    _set(ws, row_idx, "planned_complete_time", None)
    _set(ws, row_idx, "agent_name", AGENT_NAME)
    _set(ws, row_idx, "risk_type", _risk_type(risk))
    _set(ws, row_idx, "closed_notes", _mitigation(risk))
    _set(ws, row_idx, "SOURCE", AUTO_SOURCE)
    _set(ws, row_idx, "risk_code", risk_id)
    _set(ws, row_idx, "source_system_url", None)
    _set(ws, row_idx, "UPDATED_AT", timestamp_text)
    return row_idx, int(risk_numeric_id) if risk_numeric_id is not None and str(risk_numeric_id).strip().isdigit() else risk_numeric_id, is_new


def _upsert_risk_measure_row(
    ws: Worksheet,
    risk: dict[str, Any],
    target: ProjectManagementExcelTarget,
    risk_numeric_id: int | None,
    task_numeric_id: int | None,
) -> tuple[int, int, bool]:
    issue_id = _risk_issue_id(_risk_id(risk))
    row_idx = _find_row(ws, "ISSUE_ID", issue_id)
    is_new = row_idx is None
    if row_idx is None:
        row_idx = _append_or_reuse_blank_row(ws)
        measure_id = _next_numeric_id(ws, "MEASURE_ID")
        _set(ws, row_idx, "MEASURE_ID", measure_id)
    else:
        measure_id = _cell(ws, row_idx, "MEASURE_ID").value or _next_numeric_id(ws, "MEASURE_ID")
    _set(ws, row_idx, "FOREIGN_ID", risk_numeric_id)
    _set(ws, row_idx, "task_id", task_numeric_id)
    _set(ws, row_idx, "PROJECT_ID", target.file_center_project_id)
    _set(ws, row_idx, "ISSUE_ID", issue_id)
    _set(ws, row_idx, "MEASURE_NAME", _mitigation(risk))
    _set(ws, row_idx, "RESPONSIBLE_PERSON", _owner(risk))
    _set(ws, row_idx, "PLANNED_COMPLETE_TIME", None)
    _set(ws, row_idx, "REAL_COMPLETE_TIME", None)
    _set(ws, row_idx, "PROGRESS", None)
    _set(ws, row_idx, "STATUS", "待下发")
    return row_idx, int(measure_id), is_new


def _upsert_task_row(
    ws: Worksheet,
    risk: dict[str, Any],
    target: ProjectManagementExcelTarget,
    *,
    timestamp_text: str,
    date_text: str,
) -> tuple[int, bool]:
    risk_id = _risk_id(risk)
    row_idx = _find_row(ws, "关联的风险", risk_id) or _find_row(ws, "编号", _task_code(risk_id))
    is_new = row_idx is None
    if row_idx is None:
        row_idx = _append_or_reuse_blank_row(ws)
        task_id = _next_numeric_id(ws, "任务ID")
        _set(ws, row_idx, "任务ID", task_id)
    else:
        task_id = _cell(ws, row_idx, "任务ID").value or _next_numeric_id(ws, "任务ID")
    _set(ws, row_idx, "编号", _task_code(risk_id))
    _set(ws, row_idx, "*任务名称", f"处置：{_risk_name(risk)}")
    _set(ws, row_idx, "任务描述", _mitigation(risk) or _risk_description(risk))
    _set(ws, row_idx, "*责任人", _owner(risk))
    _set(ws, row_idx, "标签", "预案决策")
    _set(ws, row_idx, "优先级", _priority(risk.get("severity")))
    _set(ws, row_idx, "状态", "处理中")
    _set(ws, row_idx, "下发时间", date_text)
    _set(ws, row_idx, "*计划完成时间", None)
    _set(ws, row_idx, "实际完成时间", None)
    _set(ws, row_idx, "进展", AUTO_PROGRESS_PREFIX)
    _set(ws, row_idx, "操作人员", AGENT_NAME)
    _set(ws, row_idx, "关联的风险", risk_id)
    _set(ws, row_idx, "备注", _text(risk.get("impact")))
    return int(task_id), is_new


def _upsert_task_progress_row(
    ws: Worksheet,
    task_numeric_id: int,
    risk: dict[str, Any],
    target: ProjectManagementExcelTarget,
    *,
    timestamp_text: str,
) -> tuple[int | None, bool]:
    content = f"{AUTO_PROGRESS_PREFIX}：{_mitigation(risk) or _risk_description(risk)}"
    for row_idx in range(2, ws.max_row + 1):
        if _text(_cell(ws, row_idx, "task_id").value) == str(task_numeric_id) and _text(_cell(ws, row_idx, "progress").value) == content:
            return None, False
    row_idx = _append_or_reuse_blank_row(ws)
    progress_id = _next_numeric_id(ws, "id")
    _set(ws, row_idx, "id", progress_id)
    _set(ws, row_idx, "task_id", task_numeric_id)
    _set(ws, row_idx, "project_id", target.file_center_project_id)
    _set(ws, row_idx, "dispatch_time", timestamp_text)
    _set(ws, row_idx, "progress", content)
    _set(ws, row_idx, "operator", AGENT_NAME)
    return progress_id, True


def _upsert_issue_row(
    ws: Worksheet,
    risk: dict[str, Any],
    target: ProjectManagementExcelTarget,
    *,
    timestamp_text: str,
) -> tuple[int, bool]:
    code = _issue_code(_risk_id(risk))
    row_idx = _find_row(ws, "CUSTOM_CODE", code)
    is_new = row_idx is None
    if row_idx is None:
        row_idx = _append_or_reuse_blank_row(ws)
        issue_id = _next_numeric_id(ws, "ID")
        _set(ws, row_idx, "ID", issue_id)
        _set(ws, row_idx, "CREATED_AT", timestamp_text)
    else:
        issue_id = _cell(ws, row_idx, "ID").value or _next_numeric_id(ws, "ID")
        if not _text(_cell(ws, row_idx, "CREATED_AT").value):
            _set(ws, row_idx, "CREATED_AT", timestamp_text)
    _set(ws, row_idx, "CUSTOM_CODE", code)
    _set(ws, row_idx, "PROJECT_ID", target.file_center_project_id)
    _set(ws, row_idx, "DESCRIPTION", _risk_description(risk))
    _set(ws, row_idx, "PRIMARY_ISSUE_TYPE", _risk_type(risk) or "预案决策")
    _set(ws, row_idx, "SECONDARY_ISSUE_TYPE", _risk_point(risk))
    _set(ws, row_idx, "SEVERITY", _issue_severity(risk.get("severity")))
    _set(ws, row_idx, "SOURCE", AGENT_NAME)
    _set(ws, row_idx, "STATUS", "processing")
    _set(ws, row_idx, "ASSIGNEE_ID", _owner(risk))
    _set(ws, row_idx, "PLANNED_COMPLETE_AT", None)
    _set(ws, row_idx, "ACTUAL_CLOSED_AT", None)
    _set(ws, row_idx, "CREATED_BY", AGENT_NAME)
    _set(ws, row_idx, "UPDATED_AT", timestamp_text)
    _set(ws, row_idx, "AGENT_NAME", AGENT_NAME)
    return int(issue_id), is_new


def _upsert_issue_progress_row(
    ws: Worksheet,
    issue_numeric_id: int,
    risk: dict[str, Any],
    *,
    timestamp_text: str,
) -> tuple[int | None, bool]:
    content = f"{AUTO_PROGRESS_PREFIX}：{_mitigation(risk) or _risk_description(risk)}"
    for row_idx in range(2, ws.max_row + 1):
        if _text(_cell(ws, row_idx, "PROBLEM_ID").value) == str(issue_numeric_id) and _text(_cell(ws, row_idx, "CONTENT").value) == content:
            return None, False
    row_idx = _append_or_reuse_blank_row(ws)
    progress_id = _next_numeric_id(ws, "ID")
    _set(ws, row_idx, "ID", progress_id)
    _set(ws, row_idx, "PROBLEM_ID", issue_numeric_id)
    _set(ws, row_idx, "CONTENT", content)
    _set(ws, row_idx, "LOG_TYPE", "feedback")
    _set(ws, row_idx, "CREATED_BY", AGENT_NAME)
    _set(ws, row_idx, "CREATED_AT", timestamp_text)
    return progress_id, True
