from __future__ import annotations

import argparse
import csv
import json
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from schedule_ontology import (
    FIELD_EQUIPMENT_LIST,
    FIELD_LIMIT_DURATION,
    FIELD_ONSITE_TEAM,
    FIELD_ORIGINAL_MANAGEMENT_UNIT,
    FIELD_OWNER,
    FIELD_REMOTE_TEAM,
    FIELD_RISK_BELOW_STANDARD,
    FIELD_STANDARD_DURATION,
    LocalScheduleOntology,
)
from schedule_plan_activities import (
    FIELD_ACTIVITY_ID,
    FIELD_BATCH,
    FIELD_DEPENDENCY,
    FIELD_END,
    FIELD_NAME,
    FIELD_ROW_ID,
    FIELD_SLA,
    FIELD_START,
    FIELD_UNIT,
)


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_SOURCE_ROOT = Path("D:/Code/aida-delivery")
DEFAULT_OUTPUT_CSV = BASE_DIR / "JD_A3_delivery_plan_20260610.csv"
DEFAULT_REPORT = BASE_DIR / "data" / "JD_A3_plan_refresh_report_20260610.json"
DEFAULT_PLAN_DATA_DIR = BASE_DIR / "data" / "plan"
DEFAULT_PROJECT_KEY = "\u4eac\u4e1c"

ROOM_IMPLEMENTATION = "\u673a\u623f\u6539\u9020\u5b9e\u65bd"
ARRIVAL = "\u5230\u8d27"
POWER_ON = "\u4e0a\u7535"
CLUSTER_DEBUG = "\u96c6\u7fa4\u6027\u80fd\u8c03\u4f18"

PLAN_CSV_FIELDS = [
    FIELD_ROW_ID,
    FIELD_ACTIVITY_ID,
    FIELD_UNIT,
    FIELD_NAME,
    FIELD_START,
    FIELD_END,
    FIELD_SLA,
    FIELD_DEPENDENCY,
    FIELD_STANDARD_DURATION,
    FIELD_LIMIT_DURATION,
    FIELD_RISK_BELOW_STANDARD,
    FIELD_BATCH,
    FIELD_EQUIPMENT_LIST,
    FIELD_REMOTE_TEAM,
    FIELD_ONSITE_TEAM,
    FIELD_OWNER,
    FIELD_ORIGINAL_MANAGEMENT_UNIT,
]

AUXILIARY_PLAN_CSV_FILES = {
    "EquipmentRoom": "equipment_rooms_jd.csv",
    "EquipmentArrivalItem": "equipment_arrival_items_jd.csv",
    "DeliveryBatch": "delivery_batches_jd.csv",
    "ProjectParticipant": "project_participants_jd.csv",
    "ResourceTeam": "resource_teams_jd.csv",
    "ActivityTemplate": "activity_templates_jd.csv",
    "ScheduleRisk": "schedule_risks_jd.csv",
}

SYNC_OBJECT_TYPES = [
    "DeliveryProject",
    "DeliveryPod",
    "DeliveryPlanRow",
    "Milestone",
    *AUXILIARY_PLAN_CSV_FILES.keys(),
]


def main() -> None:
    parser = argparse.ArgumentParser(description="Refresh JD A3 plan ontology data from project Excel sources.")
    parser.add_argument("--source-root", type=Path, default=DEFAULT_SOURCE_ROOT)
    parser.add_argument("--output-csv", type=Path, default=DEFAULT_OUTPUT_CSV)
    parser.add_argument("--milestone-csv", type=Path, default=BASE_DIR / "Milestone.csv")
    parser.add_argument("--plan-data-dir", type=Path, default=DEFAULT_PLAN_DATA_DIR)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--project-key", default=DEFAULT_PROJECT_KEY)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--skip-dolt", action="store_true")
    args = parser.parse_args()

    source_bundle = load_source_bundle(args.source_root)
    converted_rows = convert_issued_rows(source_bundle)
    auxiliary_rows = build_auxiliary_plan_rows(source_bundle, args.project_key)
    schema_rows, warnings = build_schema_rows(converted_rows, args.project_key, args.output_csv.name)
    milestones = derive_milestones(schema_rows)
    report = build_report(source_bundle, converted_rows, schema_rows, warnings, milestones, auxiliary_rows)

    if not args.dry_run:
        write_plan_csv(args.output_csv, converted_rows)
        write_auxiliary_plan_csvs(args.plan_data_dir, auxiliary_rows)
        update_milestone_csv(args.milestone_csv, args.project_key, milestones)
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    if not args.dry_run and not args.skip_dolt:
        load_env(BASE_DIR / ".env")
        report["dolt"] = sync_and_validate_dolt()
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=2))


def load_source_bundle(source_root: Path) -> dict[str, Any]:
    paths = discover_source_paths(source_root)
    issued_rows = load_sheet_rows(*paths["issued_full"])
    activity_def_rows = load_sheet_rows(*paths["activity_def"])
    dependency_rows = load_sheet_rows(*paths["activity_dep"])
    baseline_rows = load_baseline_rows(*paths["baseline"])
    bundle = {
        "paths": {key: f"{path}::{sheet}" for key, (path, sheet) in paths.items()},
        "issued": issued_rows,
        "activity_defs": activity_def_rows,
        "dependencies": dependency_rows,
        "baseline": baseline_rows,
    }
    optional_roles = {
        "equipment_rooms": "equipment_rooms",
        "arrival_items": "arrival_items",
        "participants": "participants",
        "resource_teams": "resource_teams",
        "batches": "batches",
        "schedule_risks": "schedule_risks",
    }
    for path_key, bundle_key in optional_roles.items():
        bundle[bundle_key] = load_sheet_rows(*paths[path_key]) if path_key in paths else []
    return bundle


def discover_source_paths(source_root: Path) -> dict[str, tuple[Path, str]]:
    structured_paths = discover_structured_project_data_paths(source_root)
    if all(role in structured_paths for role in ("issued_full", "activity_def", "activity_dep", "baseline")):
        return structured_paths

    paths: dict[str, tuple[Path, str]] = {}
    workbooks = sorted(source_root.rglob("*.xlsx"))
    for path in workbooks:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                header = first_header_row(worksheet)
                header_set = set(header)
                if {"ID", "TODO_ID", "SERIAL_NUMBER", "OWNER"}.issubset(header_set):
                    paths["issued_full"] = (path, worksheet.title)
                if {"ACTIVITY_ID", "ACTIVITY_NAME", "SLA"}.issubset(header_set):
                    paths["activity_def"] = (path, worksheet.title)
                if {"REL_ID", "FROM_ACTIVITY_ID", "DEST_ACTIVITY_ID", "FROM_ACTIVITY_NAME", "DEST_ACTIVITY_NAME"}.issubset(header_set):
                    paths["activity_dep"] = (path, worksheet.title)
                if "SS" in header_set and "FS" in header_set and len(header) >= 30:
                    paths["baseline"] = (path, worksheet.title)
        finally:
            workbook.close()

    missing = [role for role in ("issued_full", "activity_def", "activity_dep", "baseline") if role not in paths]
    if missing:
        raise RuntimeError(f"Missing source workbook role(s): {', '.join(missing)}")
    return paths


def discover_structured_project_data_paths(source_root: Path) -> dict[str, tuple[Path, str]]:
    specs = {
        "activity_def": (
            source_root / "01_活动定义" / "A3-液冷-活动定义.xlsx",
            {"ACTIVITY_ID", "ACTIVITY_NAME", "SLA"},
        ),
        "activity_dep": (
            source_root / "02_活动依赖" / "A3-液冷-活动依赖.xlsx",
            {"REL_ID", "FROM_ACTIVITY_ID", "DEST_ACTIVITY_ID", "FROM_ACTIVITY_NAME", "DEST_ACTIVITY_NAME"},
        ),
        "baseline": (
            source_root / "04_计划基线" / "计划基线0609.xlsx",
            {"序号", "一级活动", "二级活动", "SS", "FS"},
        ),
        "equipment_rooms": (
            source_root / "05_机房机柜信息" / "机房机柜信息表.xlsx",
            {"PoD名称", "机房名称"},
        ),
        "arrival_items": (
            source_root / "06_到货表" / "04 JD三期_A3液冷到货表_260309.xlsx",
            {"ID", "管理单元", "设备类型", "型号", "到货日期"},
        ),
        "participants": (
            source_root / "07_项目人员信息" / "03 项目人员信息表_JD_0114.xlsx",
            {"姓名", "工号", "角色"},
        ),
        "resource_teams": (
            source_root / "08_施工队伍信息" / "施工队伍信息.xlsx",
            {"队伍编号", "人数", "经验等级", "在场状态"},
        ),
        "batches": (
            source_root / "09_批次信息" / "批次信息.xlsx",
            {"批次名", "上电目标日期", "上线目标日期", "该批包含的PoD"},
        ),
        "issued_full": (
            source_root / "10_输出文件" / "交付计划表.xlsx",
            {"ID", "SERIAL_NUMBER", "ACTIVITY_ID", "ACTIVITY_NAME"},
        ),
        "schedule_risks": (
            source_root / "11_风险信息" / "风险信息.xlsx",
            {"id", "project_id", "description", "level"},
        ),
    }
    paths: dict[str, tuple[Path, str]] = {}
    for role, (path, required_headers) in specs.items():
        if not path.exists():
            continue
        paths[role] = (path, find_sheet_with_headers(path, required_headers))
    return paths


def find_sheet_with_headers(path: Path, required_headers: set[str]) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            if required_headers.issubset(set(first_header_row(worksheet))):
                return worksheet.title
        return workbook.worksheets[0].title
    finally:
        workbook.close()


def first_header_row(worksheet: Any) -> list[str]:
    for values in worksheet.iter_rows(min_row=1, max_row=20, values_only=True):
        header = [string_value(value) for value in values]
        if any(header):
            return header
    return []


def load_sheet_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        iterator = iter(worksheet.iter_rows(values_only=True))
        header: list[str] = []
        for values in iterator:
            header = [string_value(value) for value in values]
            if any(header):
                break
        rows: list[dict[str, Any]] = []
        for values in iterator:
            row = {header[index]: values[index] if index < len(values) else None for index in range(len(header)) if header[index]}
            if any(string_value(value) for value in row.values()):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def build_auxiliary_plan_rows(source_bundle: dict[str, Any], project_key: str) -> dict[str, list[dict[str, str]]]:
    baseline = source_bundle.get("baseline") or {}
    return {
        "EquipmentRoom": build_equipment_room_rows(source_bundle.get("equipment_rooms") or [], project_key),
        "EquipmentArrivalItem": build_equipment_arrival_item_rows(source_bundle.get("arrival_items") or [], project_key),
        "DeliveryBatch": build_delivery_batch_rows(source_bundle.get("batches") or [], project_key),
        "ProjectParticipant": build_project_participant_rows(
            source_bundle.get("participants") or [],
            project_key,
            include_construction=bool(source_bundle.get("resource_teams")),
        ),
        "ResourceTeam": build_resource_team_rows(source_bundle.get("resource_teams") or [], project_key),
        "ActivityTemplate": build_activity_template_rows(source_bundle.get("activity_defs") or [], baseline, project_key),
        "ScheduleRisk": build_schedule_risk_rows(source_bundle.get("schedule_risks") or [], project_key),
    }


def build_equipment_room_rows(rows: list[dict[str, Any]], project_key: str) -> list[dict[str, str]]:
    grouped: dict[str, dict[str, Any]] = {}
    cabinet_columns = ["计算柜", "总线柜", "参数面Leaf柜", "样本面Leaf柜", "业务面Leaf柜", "管理面柜"]
    for row in rows:
        room_name = string_value(row.get("机房名称"))
        if not room_name:
            continue
        pod_name = string_value(row.get("PoD名称"))
        entry = grouped.setdefault(
            room_name,
            {
                "roomKey": f"{project_key}::room::{room_name}",
                "projectKey": project_key,
                "roomName": room_name,
                "podKeys": set(),
                "conditions": [],
            },
        )
        if pod_name:
            entry["podKeys"].add(pod_name)
        cabinets = {
            column: string_value(row.get(column))
            for column in cabinet_columns
            if string_value(row.get(column))
        }
        if pod_name or cabinets:
            entry["conditions"].append({"podName": pod_name, "cabinets": cabinets})
    output = []
    for entry in grouped.values():
        output.append(
            {
                "roomKey": entry["roomKey"],
                "projectKey": entry["projectKey"],
                "roomName": entry["roomName"],
                "podKeys": ",".join(sorted(entry["podKeys"])),
                "readinessStatus": "UNKNOWN",
                "roomConditionJson": json.dumps(entry["conditions"], ensure_ascii=False, sort_keys=True),
            }
        )
    return output


def build_equipment_arrival_item_rows(rows: list[dict[str, Any]], project_key: str) -> list[dict[str, str]]:
    output = []
    for index, row in enumerate(rows, start=1):
        item_id = string_value(row.get("ID")) or str(index)
        output.append(
            {
                "arrivalItemKey": f"{project_key}::arrival::{item_id}",
                "projectKey": project_key,
                "itemId": item_id,
                "managementUnit": string_value(row.get("管理单元")),
                "room": string_value(row.get("机房")),
                "podName": string_value(row.get("PoD") or row.get("Pod名")),
                "equipmentType": string_value(row.get("设备类型")),
                "model": string_value(row.get("型号")),
                "detailedConfig": string_value(row.get("总配置") or row.get("详细配置")),
                "unit": string_value(row.get("单位")),
                "quantity": string_value(row.get("数量")),
                "arrivalDate": date_value(row.get("到货日期")),
                "remark": string_value(row.get("备注")),
            }
        )
    return output


def build_delivery_batch_rows(rows: list[dict[str, Any]], project_key: str) -> list[dict[str, str]]:
    output = []
    for index, row in enumerate(rows, start=1):
        batch_name = string_value(row.get("批次名")) or f"批次{index}"
        output.append(
            {
                "batchKey": f"{project_key}::batch::{batch_name}",
                "projectKey": project_key,
                "batchName": batch_name,
                "sequence": str(index),
                "podKeys": string_value(row.get("该批包含的PoD")),
                "targetPowerOnDate": date_value(row.get("上电目标日期")),
                "targetGoLiveDate": date_value(row.get("上线目标日期")),
                "status": "ACTIVE",
            }
        )
    return output


def build_project_participant_rows(
    rows: list[dict[str, Any]],
    project_key: str,
    *,
    include_construction: bool,
) -> list[dict[str, str]]:
    output = []
    seen: set[str] = set()
    for index, row in enumerate(rows, start=1):
        name = string_value(row.get("姓名"))
        if not name:
            continue
        employee_id = string_value(row.get("工号")) or str(index)
        participant_key = f"{project_key}::participant::{employee_id}"
        if participant_key in seen:
            continue
        seen.add(participant_key)
        output.append(
            {
                "participantKey": participant_key,
                "projectKey": project_key,
                "partyType": "PERSON",
                "partyName": name,
                "responsibilityScope": string_value(row.get("角色")),
                "description": string_value(row.get("备注")),
            }
        )
    if include_construction and f"{project_key}::participant::construction" not in seen:
        output.append(
            {
                "participantKey": f"{project_key}::participant::construction",
                "projectKey": project_key,
                "partyType": "PARTNER",
                "partyName": "施工队伍",
                "responsibilityScope": "现场施工",
                "description": "由 08_施工队伍信息 自动生成的默认参与方。",
            }
        )
    return output


def build_resource_team_rows(rows: list[dict[str, Any]], project_key: str) -> list[dict[str, str]]:
    output = []
    participant_key = f"{project_key}::participant::construction"
    for index, row in enumerate(rows, start=1):
        team_id = string_value(row.get("队伍编号")) or str(index)
        output.append(
            {
                "teamKey": f"{project_key}::team::{team_id}",
                "participantKey": participant_key,
                "projectKey": project_key,
                "teamName": f"施工队伍{team_id}",
                "teamType": "CONSTRUCTION",
                "headcount": string_value(row.get("人数")),
                "experienceLevel": string_value(row.get("经验等级")),
                "roleMix": string_value(row.get("在场状态")),
            }
        )
    return output


def build_activity_template_rows(
    activity_defs: list[dict[str, Any]],
    baseline: dict[tuple[str, str], dict[str, str]],
    project_key: str,
) -> list[dict[str, str]]:
    baseline_by_id = {activity_id: value for (activity_id, _name), value in baseline.items()}
    baseline_by_name = {name: value for (_activity_id, name), value in baseline.items()}
    output = []
    seen: set[str] = set()
    for row in activity_defs:
        activity_id = normalize_activity_id(row.get("ACTIVITY_ID"))
        activity_name = string_value(row.get("ACTIVITY_NAME"))
        if not activity_id and not activity_name:
            continue
        key_part = activity_id or normalize_key(activity_name)
        template_key = f"{project_key}::activity_template::{key_part}"
        if template_key in seen:
            continue
        seen.add(template_key)
        baseline_item = baseline_by_id.get(activity_id) or baseline_by_name.get(activity_name) or {}
        output.append(
            {
                "templateKey": template_key,
                "activityType": activity_name or activity_id,
                "baseStandardSlaDays": baseline_item.get("standard", "") or string_value(row.get("SLA")),
                "baseLimitSlaDays": baseline_item.get("limit", ""),
                "resourceSensitive": "false",
                "scope": "PROJECT",
            }
        )
    return output


def build_schedule_risk_rows(rows: list[dict[str, Any]], project_key: str) -> list[dict[str, str]]:
    output = []
    for index, row in enumerate(rows, start=1):
        risk_id = string_value(row.get("id") or row.get("风险ID")) or str(index)
        category = string_value(row.get("name") or row.get("风险分类") or row.get("风险类型")) or "项目风险"
        description = string_value(row.get("description") or row.get("风险描述"))
        issue_id = string_value(row.get("issue_id") or row.get("问题ID"))
        output.append(
            {
                "riskKey": f"{project_key}::schedule_risk::{risk_id}",
                "projectKey": project_key,
                "riskType": category,
                "category": category,
                "title": description or category,
                "rootCause": description,
                "impactScopeJson": json.dumps(
                    {
                        "issueId": issue_id,
                        "ruleId": string_value(row.get("rule_id") or row.get("规则ID")),
                        "effect": string_value(row.get("effect") or row.get("风险影响")),
                        "progress": string_value(row.get("progress") or row.get("进展")),
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                ),
                "suggestedAction": string_value(row.get("关闭说明")),
                "coordinationRole": string_value(row.get("owner") or row.get("责任人")),
                "severity": string_value(row.get("level") or row.get("风险等级")),
                "status": string_value(row.get("status") or row.get("状态")),
            }
        )
    return output


def load_baseline_rows(path: Path, sheet_name: str) -> dict[tuple[str, str], dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        rows: dict[tuple[str, str], dict[str, str]] = {}
        for values in worksheet.iter_rows(min_row=3, values_only=True):
            activity_id = normalize_activity_id(values[0] if len(values) > 0 else "")
            activity_name = string_value(values[2] if len(values) > 2 else "")
            if not activity_id or not activity_name:
                continue
            rows[(activity_id, activity_name)] = {
                "standard": string_value(values[9] if len(values) > 9 else ""),
                "limit": string_value(values[10] if len(values) > 10 else ""),
                "riskName": string_value(values[14] if len(values) > 14 else ""),
                "riskDesc": string_value(values[15] if len(values) > 15 else ""),
                "riskImpact": string_value(values[16] if len(values) > 16 else ""),
                "riskPlan": string_value(values[17] if len(values) > 17 else ""),
            }
        return rows
    finally:
        workbook.close()


def convert_issued_rows(source_bundle: dict[str, Any]) -> list[dict[str, str]]:
    activity_defs = source_bundle["activity_defs"]
    dependencies = source_bundle["dependencies"]
    baseline = source_bundle["baseline"]
    sla_by_id = {normalize_activity_id(row.get("ACTIVITY_ID")): string_value(row.get("SLA")) for row in activity_defs}
    baseline_by_id = {activity_id: value for (activity_id, _name), value in baseline.items()}
    baseline_by_name = {name: value for (_activity_id, name), value in baseline.items()}
    deps_by_dest: dict[str, list[str]] = defaultdict(list)
    for row in dependencies:
        dest = normalize_activity_id(row.get("DEST_ACTIVITY_ID"))
        source_name = string_value(row.get("FROM_ACTIVITY_NAME"))
        if dest and source_name and source_name not in deps_by_dest[dest]:
            deps_by_dest[dest].append(source_name)

    converted_rows: list[dict[str, str]] = []
    for source_row in source_bundle["issued"]:
        activity_id = normalize_activity_id(source_row.get("ACTIVITY_ID"))
        activity_name = string_value(source_row.get("ACTIVITY_NAME"))
        baseline_item = (
            baseline.get((activity_id, activity_name))
            or baseline_by_id.get(activity_id)
            or baseline_by_name.get(activity_name)
            or {}
        )
        risk_bits = [
            baseline_item.get("riskName", ""),
            baseline_item.get("riskDesc", ""),
            baseline_item.get("riskImpact", ""),
            baseline_item.get("riskPlan", ""),
        ]
        converted_rows.append(
            {
                FIELD_ROW_ID: string_value(source_row.get("SERIAL_NUMBER") or source_row.get("ID")),
                FIELD_ACTIVITY_ID: activity_id,
                FIELD_UNIT: string_value(source_row.get("MANAGEMENT_UNIT")),
                FIELD_NAME: activity_name,
                FIELD_START: date_value(source_row.get("START_DATE")),
                FIELD_END: date_value(source_row.get("END_DATE")),
                FIELD_SLA: sla_by_id.get(activity_id, ""),
                FIELD_DEPENDENCY: ",".join(deps_by_dest.get(activity_id, [])),
                FIELD_STANDARD_DURATION: baseline_item.get("standard", ""),
                FIELD_LIMIT_DURATION: baseline_item.get("limit", ""),
                FIELD_RISK_BELOW_STANDARD: ";".join(value for value in risk_bits if value),
                FIELD_BATCH: "",
                FIELD_EQUIPMENT_LIST: string_value(source_row.get("FORMAT_INSTRUCTION") or source_row.get("RAW_EQUIPMENT_LIST")),
                FIELD_REMOTE_TEAM: string_value(source_row.get("SRC_AGENT")),
                FIELD_ONSITE_TEAM: string_value(source_row.get("TARGET_AGENT")),
                FIELD_OWNER: string_value(source_row.get("PRINCIPAL") or source_row.get("OWNER")),
                FIELD_ORIGINAL_MANAGEMENT_UNIT: string_value(source_row.get("REAL_MANAGEMENT_UNIT")),
            }
        )
    return converted_rows


def build_schema_rows(
    converted_rows: list[dict[str, str]],
    project_key: str,
    source_file: str,
) -> tuple[list[dict[str, str]], list[str]]:
    ontology = LocalScheduleOntology.from_rows(
        converted_rows,
        project_key=project_key,
        project_name=project_key,
        source_path=Path(source_file),
    )
    rows = []
    for row in ontology.objects.DeliveryPlanRow.all():
        item = row.to_schema_object()
        item["sourceFile"] = source_file
        rows.append(item)
    return rows, list(ontology.warnings)


def derive_milestones(schema_rows: list[dict[str, str]]) -> dict[str, Any]:
    pods = sorted(
        {
            unit.strip()
            for row in schema_rows
            for unit in string_value(row.get("managementUnit")).replace("\uff0c", ",").split(",")
            if unit.strip()
        }
    )
    return {
        "project": {
            "ROOM_IMPLEMENTATION_DONE": max_end_date(schema_rows, ROOM_IMPLEMENTATION),
            "ARRIVAL": max_end_date(schema_rows, ARRIVAL),
            "POWER_ON": max_end_date(schema_rows, POWER_ON),
            "CLUSTER_DEBUG": max_end_date(schema_rows, CLUSTER_DEBUG),
        },
        "podPowerOn": {pod: max_end_date(schema_rows, POWER_ON, pod) for pod in pods},
    }


def build_report(
    source_bundle: dict[str, Any],
    converted_rows: list[dict[str, str]],
    schema_rows: list[dict[str, str]],
    warnings: list[str],
    milestones: dict[str, Any],
    auxiliary_rows: dict[str, list[dict[str, str]]] | None = None,
) -> dict[str, Any]:
    required_missing = {
        key: sum(1 for row in schema_rows if not string_value(row.get(key)))
        for key in ("rowKey", "projectKey", "localRowKey", "activityName")
    }
    source_missing = {
        key: sum(1 for row in converted_rows if not string_value(row.get(key)))
        for key in (FIELD_NAME, FIELD_START, FIELD_END, FIELD_UNIT, FIELD_ROW_ID)
    }
    names = Counter(row.get("activityName", "") for row in schema_rows)
    return {
        "status": "ok" if all(value == 0 for value in required_missing.values()) else "invalid",
        "paths": source_bundle["paths"],
        "sourceRows": {
            "issuedFull": len(source_bundle["issued"]),
            "activityDefinitions": len(source_bundle["activity_defs"]),
            "activityDependencies": len(source_bundle["dependencies"]),
            "baselineActivities": len(source_bundle["baseline"]),
            "equipmentRooms": len(source_bundle.get("equipment_rooms") or []),
            "arrivalItems": len(source_bundle.get("arrival_items") or []),
            "participants": len(source_bundle.get("participants") or []),
            "resourceTeams": len(source_bundle.get("resource_teams") or []),
            "batches": len(source_bundle.get("batches") or []),
            "scheduleRisks": len(source_bundle.get("schedule_risks") or []),
        },
        "converted": {
            "deliveryPlanRows": len(schema_rows),
            "uniqueActivityNames": len(names),
            "warnings": len(warnings),
        },
        "auxiliaryPlanRows": {key: len(value) for key, value in (auxiliary_rows or {}).items()},
        "requiredMissing": required_missing,
        "sourceMissing": source_missing,
        "milestonesDerived": milestones,
        "warningSamples": warnings[:8],
    }


def write_plan_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=PLAN_CSV_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def write_auxiliary_plan_csvs(output_dir: Path, rows_by_object_type: dict[str, list[dict[str, str]]]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for object_type, rows in rows_by_object_type.items():
        file_name = AUXILIARY_PLAN_CSV_FILES.get(object_type)
        if not file_name:
            continue
        write_object_csv(output_dir / file_name, object_type, rows)


def write_object_csv(path: Path, object_type: str, rows: list[dict[str, str]]) -> None:
    fieldnames = object_schema_fieldnames(object_type)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def object_schema_fieldnames(object_type: str) -> list[str]:
    schemas = json.loads((BASE_DIR / "schema" / "object-types.json").read_text(encoding="utf-8"))
    schema = schemas.get(object_type) or {}
    fields = []
    for fallback_name, property_schema in (schema.get("properties") or {}).items():
        if isinstance(property_schema, dict):
            fields.append(str(property_schema.get("apiName") or fallback_name))
    return fields


def update_milestone_csv(path: Path, project_key: str, milestones: dict[str, Any]) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = [dict(row) for row in reader]
    for field in object_schema_fieldnames("Milestone"):
        if field not in fieldnames:
            fieldnames.append(field)

    desired: dict[str, dict[str, str]] = {}
    desired.update(
        {
            make_milestone_key(project_key, "PROJECT", "GLOBAL", milestone_type): milestone_row(
                project_key=project_key,
                scope_type="PROJECT",
                scope_key="GLOBAL",
                milestone_type=milestone_type,
                milestone_name=name,
                anchor_date=anchor_date,
                description=description,
                direction_role=direction_role,
                dependencies=dependencies,
            )
            for milestone_type, name, anchor_date, description, direction_role, dependencies in [
                (
                    "ROOM_IMPLEMENTATION_DONE",
                    ROOM_IMPLEMENTATION,
                    milestones["project"].get("ROOM_IMPLEMENTATION_DONE", ""),
                    "Project global forward-start milestone derived from the refreshed plan rows.",
                    "FORWARD_START",
                    "",
                ),
                (
                    "ARRIVAL",
                    ARRIVAL,
                    milestones["project"].get("ARRIVAL", ""),
                    "Project global forward-start milestone derived from the refreshed plan rows.",
                    "FORWARD_START",
                    "",
                ),
                (
                    "POWER_ON",
                    POWER_ON,
                    milestones["project"].get("POWER_ON", ""),
                    "Backward target milestone; depends on room implementation and arrival.",
                    "BACKWARD_TARGET",
                    "ROOM_IMPLEMENTATION_DONE,ARRIVAL",
                ),
                (
                    "CLUSTER_DEBUG",
                    CLUSTER_DEBUG,
                    milestones["project"].get("CLUSTER_DEBUG", ""),
                    "Backward target milestone; serial after power-on.",
                    "BACKWARD_TARGET",
                    "POWER_ON",
                ),
            ]
        }
    )
    for pod, anchor_date in milestones.get("podPowerOn", {}).items():
        key = make_milestone_key(project_key, "POD", pod, "POWER_ON")
        desired[key] = milestone_row(
            project_key=project_key,
            scope_type="POD",
            scope_key=pod,
            milestone_type="POWER_ON",
            milestone_name=POWER_ON,
            anchor_date=anchor_date,
            description="Pod-level power-on milestone derived from refreshed plan rows.",
            direction_role="BACKWARD_TARGET",
            dependencies="ROOM_IMPLEMENTATION_DONE,ARRIVAL",
        )

    desired_keys = set(desired)
    output_rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in rows:
        key = string_value(row.get("milestoneKey"))
        if string_value(row.get("projectKey")) == project_key and (
            key in desired_keys
            or (string_value(row.get("scopeType")) == "POD" and string_value(row.get("milestoneType")) == "POWER_ON")
        ):
            if key in desired:
                output_rows.append({field: desired[key].get(field, "") for field in fieldnames})
                seen.add(key)
            continue
        output_rows.append(row)
    for key in sorted(desired_keys - seen):
        output_rows.append({field: desired[key].get(field, "") for field in fieldnames})
    for row in output_rows:
        if "anchorPriority" in fieldnames and not string_value(row.get("anchorPriority")):
            row["anchorPriority"] = "0"
        if "anchorKind" in fieldnames:
            row.setdefault("anchorKind", "")

    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(output_rows)


def sync_and_validate_dolt() -> dict[str, Any]:
    previous_overlay = os.environ.get("DOLT_WRITEBACK_OVERLAY")
    os.environ["DOLT_WRITEBACK_OVERLAY"] = "0"
    try:
        ensure_existing_dolt_columns_can_hold_refreshed_plan_rows()
        from dolt_mirror import sync_main_mirror, validate_main_mirror

        sync = sync_main_mirror(object_types=SYNC_OBJECT_TYPES)
        validate = validate_main_mirror(object_types=SYNC_OBJECT_TYPES)
        return {"sync": sync, "validate": validate}
    finally:
        if previous_overlay is None:
            os.environ.pop("DOLT_WRITEBACK_OVERLAY", None)
        else:
            os.environ["DOLT_WRITEBACK_OVERLAY"] = previous_overlay


def ensure_existing_dolt_columns_can_hold_refreshed_plan_rows() -> None:
    """Widen historical VARCHAR columns before inserting refreshed long plan facts."""
    from dolt_schema_sync import _get_dolt_engine, _sql_text

    engine = _get_dolt_engine()
    statements = [
        "ALTER TABLE `delivery_plan_row` MODIFY COLUMN `riskBelowStandard` TEXT NULL",
        "ALTER TABLE `delivery_plan_row` MODIFY COLUMN `equipmentList` TEXT NULL",
        "ALTER TABLE `delivery_plan_row` MODIFY COLUMN `dependencyActivities` TEXT NULL",
        "ALTER TABLE `equipment_room` MODIFY COLUMN `roomConditionJson` TEXT NULL",
        "ALTER TABLE `schedule_risk` MODIFY COLUMN `impactScopeJson` TEXT NULL",
    ]
    with engine.connect() as conn:
        for statement in statements:
            try:
                conn.execute(_sql_text(statement))
            except Exception as exc:
                message = str(exc).lower()
                if "doesn't exist" not in message and "unknown table" not in message:
                    raise
        conn.commit()


def milestone_row(
    *,
    project_key: str,
    scope_type: str,
    scope_key: str,
    milestone_type: str,
    milestone_name: str,
    anchor_date: str,
    description: str,
    direction_role: str,
    dependencies: str,
) -> dict[str, str]:
    return {
        "milestoneKey": make_milestone_key(project_key, scope_type, scope_key, milestone_type),
        "projectKey": project_key,
        "projectName": project_key,
        "scopeType": scope_type,
        "scopeKey": scope_key,
        "milestoneType": milestone_type,
        "milestoneName": milestone_name,
        "plannedDate": "",
        "actualDate": "",
        "anchorDate": anchor_date,
        "sourceFile": "Milestone.csv",
        "description": description,
        "directionRole": direction_role,
        "dependencyMilestoneTypes": dependencies,
        "anchorPriority": "0",
        "anchorKind": "",
    }


def make_milestone_key(project_key: str, scope_type: str, scope_key: str, milestone_type: str) -> str:
    return f"{project_key}::{scope_type}::{scope_key}::{milestone_type}"


def max_end_date(rows: list[dict[str, str]], contains: str, unit: str | None = None) -> str:
    dates = []
    for row in rows:
        if contains not in string_value(row.get("activityName")):
            continue
        if unit is not None and string_value(row.get("managementUnit")) != unit:
            continue
        parsed = parse_date(row.get("endDate"))
        if parsed is not None:
            dates.append(parsed)
    return max(dates).isoformat() if dates else ""


def load_env(path: Path) -> None:
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def normalize_activity_id(value: Any) -> str:
    text = string_value(value)
    return text[:-2] if text.endswith(".0") else text


def string_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def normalize_key(value: Any) -> str:
    text = string_value(value)
    normalized = re.sub(r"\s+", "_", text)
    return normalized or "unknown"


def date_value(value: Any) -> str:
    if value is None or string_value(value) == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = string_value(value).replace("/", "-")
    try:
        return datetime.fromisoformat(raw).date().isoformat()
    except ValueError:
        pass
    try:
        return datetime.fromisoformat(raw[:10]).date().isoformat()
    except ValueError:
        return raw


def parse_date(value: Any) -> date | None:
    raw = string_value(value).replace("/", "-")
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw[:10]).date()
    except ValueError:
        return None


if __name__ == "__main__":
    main()
