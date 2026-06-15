"""Load the 预案 风险列表 (21) and 假设表 (22) Excel sheets into Dolt.

Companion to load_contingency_excel_to_dolt.py. RiskItem / AssumptionItem are *writable
run-record* ObjectTypes (normally materialised by writeback Actions such as AdoptDerivedRisk),
not Dolt-backed reference assets — and per the proposal spec (README §不计入 §4.2 输出表) the
风险/假设 sheets are not delivery-plan output tables. They are loaded here for completeness so
the raw 售前公开风险 / 假设清单 facts live in Dolt alongside the other 22-sheet substrate.

Two wrinkles vs the equipment loader:
  * the RiskItem / AssumptionItem schemas declare almost no `sourceColumnName`, so the Chinese
    header→apiName mapping is spelled out explicitly below (RISK_HEADER_MAP / ASSUMPTION_HEADER_MAP);
  * the required system columns the sheets don't carry are synthesised: projectKey=京东 (JD 三期,
    matching the rest of the contingency substrate), planId / assessmentId = "" (plan-agnostic seed,
    like the decision_point loader), status = ACTIVE; riskName/assumptionName from the name column;
    AssumptionItem.assumptionId from the sheet's 序号.

Mechanism mirrors load_contingency_excel_to_dolt.py: schema-driven DDL (build_create_table_sql) +
ADD COLUMN for any schema column absent from the table (projectKey), DELETE + INSERT, DOLT_COMMIT.

Run (Dolt sql-server must be up; DOLT_DATABASE_URL from ontology/.env):
    cd ontology && python load_contingency_risks_assumptions_to_dolt.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ONTOLOGY_DIR = Path(__file__).resolve().parent
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))

try:  # load DOLT_DATABASE_URL et al. from ontology/.env when run standalone
    from dotenv import load_dotenv

    load_dotenv(ONTOLOGY_DIR / ".env")
except Exception:  # pragma: no cover - dotenv optional
    pass

import openpyxl  # noqa: E402

from dolt_schema_sync import (  # noqa: E402
    _resolve_engine,
    _sql_text,
    build_add_column_statements,
    build_create_table_sql,
    object_type_table_name,
    quote_ident,
    validate_sql_identifier,
)

EXCEL_DIR = ONTOLOGY_DIR / "contingency-data" / "预案 excel数据"
OBJECT_TYPES_PATH = ONTOLOGY_DIR / "schema" / "object-types.json"
# All contingency substrate is the JD 三期 (京东) project.
PROJECT_KEY = "京东"

# Chinese header → schema apiName. The schemas carry no sourceColumnName for these run-record
# fields, so the mapping is explicit. `状态`(处理中/已关闭) is the business lifecycle → `state`;
# the required system `status` is synthesised to ACTIVE.
RISK_HEADER_MAP = {
    "风险ID": "riskId",
    "风险点": "riskPoint",
    "描述": "description",
    "影响": "impact",
    "类型": "riskType",
    "等级": "severity",
    "来源": "source",
    "识别时间": "identifiedAt",
    "关闭时间": "closedAt",
    "应对策略": "mitigationPlan",
    "责任人": "owner",
    "状态": "state",
}
ASSUMPTION_HEADER_MAP = {
    "假设": "assumptionName",
    "来源": "source",
    "责任人": "owner",
    "兑现方式": "validationMethod",
    "关闭状态": "closeStatus",
}


def _load_object_schema(object_type: str) -> dict:
    schemas = json.loads(OBJECT_TYPES_PATH.read_text(encoding="utf-8"))
    schema = schemas.get(object_type)
    if not isinstance(schema, dict):
        raise SystemExit(f"ObjectType {object_type} not found in {OBJECT_TYPES_PATH}")
    return schema


def _schema_columns(schema: dict) -> list[str]:
    columns: list[str] = []
    for fallback, prop in (schema.get("properties") or {}).items():
        if not isinstance(prop, dict):
            continue
        column = str(prop.get("apiName") or fallback).strip()
        if column:
            validate_sql_identifier(column, kind="column")
            columns.append(column)
    return columns


def _excel_records(excel_path: Path, sheet: str, header_map: dict[str, str]) -> list[dict]:
    """Read a sheet into apiName-keyed records (mapped headers only, blank cells dropped)."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in raw_rows[0]]
    records: list[dict] = []
    for raw in raw_rows[1:]:
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        record: dict = {"_序号": raw[0]}  # keep the sheet's leading column for id synthesis
        for header, value in zip(headers, raw):
            api = header_map.get(header)
            if api and value is not None and str(value).strip() != "":
                record[api] = value
        records.append(record)
    return records


def _risk_rows() -> list[dict]:
    rows: list[dict] = []
    for i, rec in enumerate(
        _excel_records(EXCEL_DIR / "21_风险列表.xlsx", "风险列表", RISK_HEADER_MAP), start=1
    ):
        rec.pop("_序号", None)
        rec["riskId"] = str(rec.get("riskId") or "").strip() or f"R-{i:03d}"
        rec.setdefault("riskName", rec.get("riskPoint") or rec["riskId"])
        rec["projectKey"] = PROJECT_KEY
        rec.setdefault("planId", "")
        rec.setdefault("assessmentId", "")
        rec.setdefault("status", "ACTIVE")
        rec["sourceFile"] = "21_风险列表.xlsx"
        rows.append(rec)
    return rows


def _assumption_rows() -> list[dict]:
    rows: list[dict] = []
    for i, rec in enumerate(
        _excel_records(EXCEL_DIR / "22_假设表.xlsx", "假设表", ASSUMPTION_HEADER_MAP), start=1
    ):
        seq = rec.pop("_序号", None)
        try:
            seq_n = int(str(seq).strip())
        except (TypeError, ValueError):
            seq_n = i
        rec["assumptionId"] = f"ASSUM-{seq_n:03d}"
        rec.setdefault("assumptionName", rec["assumptionId"])
        rec["projectKey"] = PROJECT_KEY
        rec.setdefault("planId", "")
        rec.setdefault("assessmentId", "")
        rec.setdefault("status", "ACTIVE")
        rec["sourceFile"] = "22_假设表.xlsx"
        rows.append(rec)
    return rows


def _sql_value(value):
    # Only None -> SQL NULL (mirrors the decision_point loader). The required NOT-NULL planId /
    # assessmentId are intentionally "" (plan-agnostic seed), so "" must NOT be coerced to NULL.
    # Blank cells are already dropped in _excel_records, so no date column ever receives "" here.
    return None if value is None else value


def _fetch_column_names(conn, table_name: str) -> set[str]:
    result = conn.execute(
        _sql_text(
            "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = :t"
        ),
        {"t": table_name},
    )
    return {str(row[0]) for row in result.fetchall() if row[0] is not None}


def _load_one(conn, object_type: str, rows: list[dict]) -> tuple[str, int]:
    schema = _load_object_schema(object_type)
    table_name = object_type_table_name(object_type, schema)
    columns = _schema_columns(schema)
    if "sourceFile" not in columns:
        columns.append("sourceFile")

    conn.execute(_sql_text(build_create_table_sql(object_type, schema)))
    existing = _fetch_column_names(conn, table_name)
    add_statements, _added = build_add_column_statements(table_name, schema, existing)
    for statement in add_statements:
        conn.execute(_sql_text(statement))
    conn.commit()

    # sourceFile may not be in the schema (provenance helper column) — add it if missing.
    existing = _fetch_column_names(conn, table_name)
    if "sourceFile" not in existing:
        conn.execute(
            _sql_text(f"ALTER TABLE {quote_ident(table_name)} ADD COLUMN {quote_ident('sourceFile')} VARCHAR(255) NULL")
        )
        conn.commit()

    conn.execute(_sql_text(f"DELETE FROM {quote_ident(table_name)}"))
    column_sql = ", ".join(quote_ident(c) for c in columns)
    values_sql = ", ".join(f":{c}" for c in columns)
    insert_sql = f"INSERT INTO {quote_ident(table_name)} ({column_sql}) VALUES ({values_sql})"
    for row in rows:
        conn.execute(_sql_text(insert_sql), {c: _sql_value(row.get(c)) for c in columns})
    conn.commit()
    return table_name, len(rows)


def main() -> int:
    jobs = [
        ("RiskItem", _risk_rows()),
        ("AssumptionItem", _assumption_rows()),
    ]
    for object_type, rows in jobs:
        if not rows:
            raise SystemExit(f"No rows read for {object_type}")

    engine = _resolve_engine(None)
    if engine is None:
        raise SystemExit(
            "Dolt is not reachable. Set DOLT_DATABASE_URL (ontology/.env) and start the dolt sql-server."
        )

    loaded: list[tuple[str, int]] = []
    with engine.connect() as conn:
        for object_type, rows in jobs:
            loaded.append(_load_one(conn, object_type, rows))
        try:
            msg = "Load " + ", ".join(f"{t} ({n})" for t, n in loaded) + " from 21_风险列表/22_假设表"
            conn.execute(_sql_text("CALL dolt_commit('-A', '-m', :m)"), {"m": msg})
            conn.commit()
        except Exception as exc:  # nothing-to-commit is fine
            if "nothing to commit" not in str(exc).lower():
                raise

    for table_name, n in loaded:
        print(f"Loaded {n} rows into Dolt `{table_name}`.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
