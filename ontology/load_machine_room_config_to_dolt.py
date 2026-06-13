"""Load 机房机柜信息表.xlsx into Dolt as MachineRoomConfig (PoD × 机柜类型 展开).

The 预案「机房信息」chapter binds to MachineRoomConfig. The placeholder sheet
is shaped wide — one row per PoD with a column per 机柜类型
(计算柜/总线柜/参数面Leaf柜/业务面Leaf柜/管理面柜/样本面Leaf柜), each cell listing the cabinet
positions (e.g. "A18,A17,A16"). MachineRoomConfig's schema is the narrow asset shape
(machineRoomName / cabinetType / quantity / remark / location), so this loader **unpivots**: one
Dolt row per (PoD, 机柜类型) non-empty cell —
  * machineRoomName = 机房,  location = PoD,  cabinetType = 列名,
  * remark = the raw cabinet positions,  quantity = count of comma-separated positions,
  * power = NULL (the sheet carries no power), machineRoomId synthesised MRC-{PoD}-{code}.

Only the new Sheet1 header dialect is accepted. `预案版本号` is a load filter only:
blank and `草稿` rows are skipped, and the version value is not stored in the ontology row.

This fills MachineRoomConfig's existing schema without duplicating PodCabinetLayout's wide columns
(PodCabinetLayout stays a separate, real-source-backed object). Load is DELETE + INSERT + DOLT_COMMIT,
mirroring load_contingency_excel_to_dolt.py.

Run (Dolt sql-server up; DOLT_DATABASE_URL from ontology/.env):
    cd ontology && python load_machine_room_config_to_dolt.py
"""

from __future__ import annotations

import sys
import os
import re
from pathlib import Path

ONTOLOGY_DIR = Path(__file__).resolve().parent
if str(ONTOLOGY_DIR) not in sys.path:
    sys.path.insert(0, str(ONTOLOGY_DIR))

try:  # load DOLT_DATABASE_URL et al. from ontology/.env when run standalone
    from dotenv import load_dotenv

    load_dotenv(ONTOLOGY_DIR / ".env")
except Exception:  # pragma: no cover - dotenv optional
    pass

import json  # noqa: E402

import openpyxl  # noqa: E402

from dolt_schema_sync import (  # noqa: E402
    _resolve_engine,
    _sql_text,
    build_add_column_statements,
    build_create_table_sql,
    object_type_table_name,
    quote_ident,
    validate_sql_identifier,
)

OBJECT_TYPE = "MachineRoomConfig"
DEFAULT_EXCEL_PATH = Path(
    r"D:\Code\aida\temp\京东三期\早期介入\交付预案\输出结果\机房机柜信息表.xlsx"
)
EXCEL_PATH = Path(os.environ.get("MACHINE_ROOM_CONFIG_XLSX") or DEFAULT_EXCEL_PATH)
SHEET = "Sheet1"
OBJECT_TYPES_PATH = ONTOLOGY_DIR / "schema" / "object-types.json"

POD_HEADER = "PoD名称"
ROOM_HEADER = "机房名称"
VERSION_HEADER = "预案版本号"
# 机柜类型列 → machineRoomId 短码（PK 用 ASCII，便于追溯）。
CABINET_COLUMNS = {
    "计算柜": "COMP",
    "总线柜": "BUS",
    "参数面Leaf柜": "PLEAF",
    "业务面Leaf柜": "BLEAF",
    "管理面柜": "MGMT",
    "样本面Leaf柜": "SLEAF",
}
REQUIRED_HEADERS = (POD_HEADER, ROOM_HEADER, *CABINET_COLUMNS.keys(), VERSION_HEADER)


def _load_object_schema() -> dict:
    schemas = json.loads(OBJECT_TYPES_PATH.read_text(encoding="utf-8"))
    schema = schemas.get(OBJECT_TYPE)
    if not isinstance(schema, dict):
        raise SystemExit(f"ObjectType {OBJECT_TYPE} not found in {OBJECT_TYPES_PATH}")
    return schema


def _schema_columns(schema: dict) -> list[str]:
    columns: list[str] = []
    for fallback, prop in (schema.get("properties") or {}).items():
        if not isinstance(prop, dict):
            continue
        column = str(prop.get("apiName") or fallback).strip()
        if column:
            validate_sql_identifier(column, kind="column")
            columns.append(column)
    return columns


def _validate_headers(headers: list[str]) -> None:
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise SystemExit(
            f"Missing required headers in {EXCEL_PATH.name}: {', '.join(missing)}; got {headers}"
        )


def _is_non_draft_version(value: object) -> bool:
    version = str(value).strip() if value is not None else ""
    return bool(version) and version != "草稿"


def _split_tokens(value: str) -> list[str]:
    return [token.strip() for token in re.split(r"[,，、]", value) if token.strip()]


def _read_rows() -> list[dict]:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in raw_rows[0]]
    _validate_headers(headers)
    idx = {h: i for i, h in enumerate(headers)}
    pod_i = idx[POD_HEADER]
    room_i = idx[ROOM_HEADER]
    version_i = idx[VERSION_HEADER]

    rows: list[dict] = []
    for raw in raw_rows[1:]:
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        if version_i >= len(raw) or not _is_non_draft_version(raw[version_i]):
            continue
        pod = str(raw[pod_i]).strip() if raw[pod_i] is not None else ""
        room = str(raw[room_i]).strip() if raw[room_i] is not None else ""
        if not room:
            continue  # machineRoomName is required (NOT NULL)
        for cabinet, code in CABINET_COLUMNS.items():
            ci = idx.get(cabinet)
            if ci is None:
                continue
            value = raw[ci]
            if value is None or str(value).strip() == "":
                continue
            positions = str(value).strip()
            quantity = len(_split_tokens(positions))
            rows.append(
                {
                    "machineRoomId": f"MRC-{pod}-{code}" if pod else f"MRC-{room}-{code}",
                    "machineRoomName": room,
                    "status": "ACTIVE",
                    "cabinetType": cabinet,
                    "quantity": quantity,
                    "location": pod,
                    "remark": positions,
                    "power": None,  # sheet carries no power
                    "sourceFile": EXCEL_PATH.name,
                }
            )
    return rows


def _sql_value(value):
    # Only None -> SQL NULL; required columns always carry a value upstream, optional power/updatedAt
    # are None -> NULL. No date/datetime column ever receives "" here.
    return None if value is None else value


def _fetch_column_names(conn, table_name: str) -> set[str]:
    result = conn.execute(
        _sql_text(
            "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = :t"
        ),
        {"t": table_name},
    )
    return {str(row[0]) for row in result.fetchall() if row[0] is not None}


def main() -> int:
    schema = _load_object_schema()
    table_name = object_type_table_name(OBJECT_TYPE, schema)
    columns = _schema_columns(schema)
    rows = _read_rows()
    if not rows:
        raise SystemExit(f"No rows read from {EXCEL_PATH}")

    engine = _resolve_engine(None)
    if engine is None:
        raise SystemExit(
            "Dolt is not reachable. Set DOLT_DATABASE_URL (ontology/.env) and start the dolt sql-server."
        )

    with engine.connect() as conn:
        conn.execute(_sql_text(build_create_table_sql(OBJECT_TYPE, schema)))
        existing = _fetch_column_names(conn, table_name)
        add_statements, _added = build_add_column_statements(table_name, schema, existing)
        for statement in add_statements:
            conn.execute(_sql_text(statement))
        conn.commit()

        conn.execute(_sql_text(f"DELETE FROM {quote_ident(table_name)}"))
        column_sql = ", ".join(quote_ident(c) for c in columns)
        values_sql = ", ".join(f":{c}" for c in columns)
        insert_sql = f"INSERT INTO {quote_ident(table_name)} ({column_sql}) VALUES ({values_sql})"
        for row in rows:
            conn.execute(_sql_text(insert_sql), {c: _sql_value(row.get(c)) for c in columns})
        conn.commit()

        try:
            conn.execute(
                _sql_text("CALL dolt_commit('-A', '-m', :m)"),
                {"m": f"Load {table_name} from {EXCEL_PATH.name} ({len(rows)} rows, PoD×机柜类型 展开)"},
            )
            conn.commit()
        except Exception as exc:  # nothing-to-commit is fine
            if "nothing to commit" not in str(exc).lower():
                raise

    print(f"Loaded {len(rows)} rows into Dolt `{table_name}`:")
    for row in rows:
        print(f"  {row['machineRoomId']:22} {row['machineRoomName']:10} {row['cabinetType']:8} x{row['quantity']}  [{row['remark']}]")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
