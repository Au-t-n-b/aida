"""
交付预案 · 第 5 章（组网配置）+ 第 7 章（机房信息）· FastAPI 路由

对标 agent/sog_routes.py：APIRouter + Store 类。

路由前缀：/api/v1/projects/{project_id}/proposal

章节覆盖：
  - 5.1 网络平面配置（net-plane）
  - 5.2 网管服务器配置（net-mgmt）
  - 5.3 集群设备清单表（cluster-device-list）
  - 7.1 机房机柜信息（room-rack）

对齐：
  - 04-预案上下文.md §8（组网配置）、§10（机房信息）
  - 开发文档/第5章-组网配置/00-通用约定.md §6.1（路由一览）
  - 开发文档/第7章-机房信息/00-通用约定.md §5.1（路由一览）
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException, Query, Response

from .proposal_models import (
    ApiResponse,
    ClusterDeviceCreateReq,
    ClusterDeviceUpdateReq,
    DeviceRoleItem,
    ImportResult,
    NetMgmtCreateReq,
    NetMgmtUpdateReq,
    NetPlaneCreateReq,
    NetPlaneUpdateReq,
    ReleaseReq,
    ResponseMeta,
    RoomRackCreateReq,
    RoomRackUpdateReq,
    WarningItem,
)
from .proposal_store import ProposalStore

router = APIRouter(prefix="/api/v1/projects/{project_id}/proposal", tags=["proposal"])
_store = ProposalStore()


def _make_meta(project_id: str, warnings: list[WarningItem] | None = None) -> ResponseMeta:
    return ResponseMeta(
        project_id=project_id,
        proposal_version="draft",
        source_layer="draft",
        warnings=warnings or [],
    )


def _handle_store_error(exc: Exception) -> HTTPException:
    msg = str(exc)
    if msg.startswith("BR_PROPOSAL_"):
        code, detail = msg.split(":", 1)
        return HTTPException(status_code=422, detail={"code": code.strip(), "message": detail.strip()})
    return HTTPException(status_code=400, detail=msg)


# ─── 5.1 网络平面配置 ────────────────────────────────────────────────────────

@router.get("/chapters/5.1/net-plane")
def list_net_plane(project_id: str) -> ApiResponse:
    rows = _store.list_net_plane(project_id)
    return ApiResponse(
        data={"rows": [r.model_dump() for r in rows]},
        meta=_make_meta(project_id),
    )


@router.get("/chapters/5.1/net-plane/device-roles")
def list_device_roles(project_id: str) -> ApiResponse:
    """返回设备角色枚举值列表（供前端下拉选择）"""
    roles = _store.list_device_roles()
    return ApiResponse(
        data={"device_roles": [r.model_dump() for r in roles]},
        meta=_make_meta(project_id),
    )


@router.get("/chapters/5.1/net-plane/available-devices")
def list_available_devices(project_id: str) -> ApiResponse:
    """返回设备信息表可用设备（供前端选设备时自动填充厂家/型号/版本）"""
    devices = _store.list_available_devices()
    return ApiResponse(
        data={"devices": [d.model_dump() for d in devices]},
        meta=_make_meta(project_id),
    )


@router.post("/chapters/5.1/net-plane/export")
def export_net_plane(project_id: str, payload: dict[str, Any]) -> ApiResponse:
    """将当前网络平面配置保存到项目 data/delivery/共平面类型表.xlsx。"""
    from openpyxl import Workbook

    rows = payload.get("rows") or [row.model_dump() for row in _store.list_net_plane(project_id)]
    output_dir = Path(__file__).resolve().parent.parent / "data" / "delivery"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "共平面类型表.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "共平面类型"
    sheet.append(["设备角色", "设备型号", "设备厂家", "设备版本", "数量", "来源", "备注"])
    for row in rows:
        sheet.append([
            row.get("type", ""),
            row.get("model", ""),
            row.get("vendor", ""),
            row.get("ver", ""),
            row.get("qty", 0),
            row.get("source", ""),
            row.get("note") or "",
        ])
    for column, width in zip("ABCDEFG", [28, 24, 16, 24, 10, 18, 32]):
        sheet.column_dimensions[column].width = width
    workbook.save(output_path)

    return ApiResponse(
        data={"saved_path": str(output_path.relative_to(Path(__file__).resolve().parent.parent))},
        meta=_make_meta(project_id),
    )


@router.post("/chapters/5.1/net-plane", status_code=201)
def create_net_plane(project_id: str, req: NetPlaneCreateReq) -> ApiResponse:
    try:
        row = _store.create_net_plane(project_id, req.source_row_id)
        return ApiResponse(
            data={"row": row.model_dump()},
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.patch("/chapters/5.1/net-plane/{row_id}")
def update_net_plane(project_id: str, row_id: str, req: NetPlaneUpdateReq) -> ApiResponse:
    try:
        data = req.model_dump(exclude_unset=True)
        row, warnings = _store.update_net_plane(project_id, row_id, data)
        return ApiResponse(
            data={"row": row.model_dump(), "warnings": [w.model_dump() for w in warnings]},
            meta=_make_meta(project_id, warnings),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.delete("/chapters/5.1/net-plane/{row_id}")
def delete_net_plane(
    project_id: str,
    row_id: str,
    confirm: bool = Query(default=False, description="自动解析行需要 confirm=true"),
) -> ApiResponse:
    try:
        _store.delete_net_plane(project_id, row_id, confirm)
        return ApiResponse(
            data={"deleted": row_id},
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


# ─── 5.2 网管服务器配置 ──────────────────────────────────────────────────────

@router.get("/chapters/5.2/net-mgmt")
def list_net_mgmt(project_id: str) -> ApiResponse:
    rows = _store.list_net_mgmt(project_id)
    return ApiResponse(
        data={"rows": [r.model_dump() for r in rows]},
        meta=_make_meta(project_id),
    )


@router.post("/chapters/5.2/net-mgmt/initialize")
def initialize_net_mgmt(project_id: str) -> ApiResponse:
    rows = _store.initialize_net_mgmt(project_id)
    return ApiResponse(
        data={"rows": [r.model_dump() for r in rows]},
        meta=_make_meta(project_id),
    )


@router.post("/chapters/5.2/net-mgmt", status_code=201)
def create_net_mgmt(project_id: str, req: NetMgmtCreateReq) -> ApiResponse:
    try:
        row = _store.create_net_mgmt(project_id, req.model_dump())
        return ApiResponse(
            data=row.model_dump(),
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.patch("/chapters/5.2/net-mgmt/{row_id}")
def update_net_mgmt(project_id: str, row_id: str, req: NetMgmtUpdateReq) -> ApiResponse:
    try:
        row = _store.update_net_mgmt(project_id, row_id, req.model_dump(exclude_unset=True))
        return ApiResponse(
            data=row.model_dump(),
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.delete("/chapters/5.2/net-mgmt/{row_id}", status_code=204, response_class=Response)
def delete_net_mgmt(project_id: str, row_id: str):
    try:
        _store.delete_net_mgmt(project_id, row_id)
    except ValueError as e:
        raise _handle_store_error(e)


@router.post("/chapters/5.2/net-mgmt/export")
def export_net_mgmt(project_id: str, payload: dict[str, Any]) -> ApiResponse:
    """保存网管服务器配置表到 data/delivery。"""
    from openpyxl import Workbook

    rows = payload.get("rows") or [row.model_dump() for row in _store.list_net_mgmt(project_id)]
    output_dir = Path(__file__).resolve().parent.parent / "data" / "delivery"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "网管服务器配置表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "网管服务器配置"
    sheet.append(["服务器角色", "服务器型号", "数量"])
    for row in rows:
        sheet.append([row.get("server_role", ""), row.get("server_model", ""), row.get("quantity", 1)])
    for column, width in zip("ABC", [20, 28, 12]):
        sheet.column_dimensions[column].width = width
    workbook.save(output_path)
    return ApiResponse(
        data={"saved_path": str(output_path.relative_to(Path(__file__).resolve().parent.parent))},
        meta=_make_meta(project_id),
    )


# ─── 5.3 集群设备清单表 ──────────────────────────────────────────────────────

@router.get("/chapters/5.3/cluster-device-list")
def list_cluster_device(project_id: str) -> ApiResponse:
    rows = _store.list_cluster_device(project_id)
    return ApiResponse(
        data={"rows": [r.model_dump() for r in rows]},
        meta=_make_meta(project_id),
    )


@router.get("/chapters/5.3/net-plane-options")
def list_net_plane_options(project_id: str) -> ApiResponse:
    """返回 5.1 网络平面行（供 5.3 选择继承来源）"""
    rows = _store.list_net_plane(project_id)
    options = [
        {
            "row_id": r.row_id,
            "type": r.type,
            "vendor": r.vendor,
            "model": r.model,
            "qty": r.qty,
        }
        for r in rows
    ]
    return ApiResponse(
        data={"options": options},
        meta=_make_meta(project_id),
    )


@router.post("/chapters/5.3/cluster-device-list", status_code=201)
def create_cluster_device(project_id: str, req: ClusterDeviceCreateReq) -> ApiResponse:
    try:
        row = _store.create_cluster_device(project_id, req.source_net_plane_id)
        return ApiResponse(
            data=row.model_dump(),
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.patch("/chapters/5.3/cluster-device-list/{row_id}")
def update_cluster_device(project_id: str, row_id: str, req: ClusterDeviceUpdateReq) -> ApiResponse:
    try:
        data = req.model_dump(exclude_unset=True)
        row, warnings = _store.update_cluster_device(project_id, row_id, data)
        return ApiResponse(
            data=row.model_dump(),
            meta=_make_meta(project_id, warnings),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.delete("/chapters/5.3/cluster-device-list/{row_id}", status_code=200)
def delete_cluster_device(
    project_id: str,
    row_id: str,
    confirm: bool = Query(default=False, description="自动解析行需要 confirm=true"),
) -> None:
    try:
        _store.delete_cluster_device(project_id, row_id, confirm)
    except ValueError as e:
        raise _handle_store_error(e)


@router.post("/chapters/5.3/cluster-device-list/export")
def export_cluster_device(project_id: str, payload: dict[str, Any]) -> ApiResponse:
    """保存集群设备清单表到 data/delivery。"""
    from openpyxl import Workbook

    rows = payload.get("rows") or [row.model_dump() for row in _store.list_cluster_device(project_id)]
    output_dir = Path(__file__).resolve().parent.parent / "data" / "delivery"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "集群设备清单表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "集群设备清单"
    fields = [
        "cluster_id", "cluster_type", "super_pod_id", "storage_cluster_id", "zone_id",
        "ccae_cluster_id", "dme_cluster_id", "device_type", "vendor", "device_model",
        "device_purpose", "start_device_name", "end_device_name", "quantity",
    ]
    sheet.append([
        "集群ID", "集群类型", "超节点ID", "存储集群ID", "ZONE ID", "CCAE集群ID", "DME集群ID",
        "设备类型", "厂家", "设备型号", "设备用途", "起始设备命名", "截止设备命名", "数量",
    ])
    for row in rows:
        sheet.append([row.get(field) or "" for field in fields])
    workbook.save(output_path)
    return ApiResponse(
        data={"saved_path": str(output_path.relative_to(Path(__file__).resolve().parent.parent))},
        meta=_make_meta(project_id),
    )


# ─── 7.1 机房信息（PoD 机柜分配） ─────────────────────────────────────────────


@router.get("/chapters/7.1/device-info-table")
def list_device_info_table(project_id: str) -> ApiResponse:
    """获取 device_info_table 数据（mock，后续从数据中心 API 获取）"""
    items = _store.get_device_info_table()
    return ApiResponse(
        data={"items": items},
        meta=_make_meta(project_id),
    )


@router.get("/chapters/7.1/room-rack")
def list_room_rack(project_id: str) -> ApiResponse:
    rows = _store.list_room_rack(project_id)
    return ApiResponse(
        data={"rows": [r.model_dump() for r in rows]},
        meta=_make_meta(project_id),
    )


@router.post("/chapters/7.1/room-rack", status_code=201)
def create_room_rack(project_id: str, req: RoomRackCreateReq, source_row_id: str | None = Query(default=None)) -> ApiResponse:
    try:
        row = _store.create_room_rack(project_id, req.model_dump(), source_row_id)
        return ApiResponse(
            data=row.model_dump(),
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.post("/chapters/7.1/room-rack/export")
def export_room_rack(project_id: str, payload: dict[str, Any]) -> ApiResponse:
    """将当前机房信息保存到项目 data/delivery/机房信息表.xlsx。"""
    from openpyxl import Workbook

    rows = payload.get("rows") or [row.model_dump() for row in _store.list_room_rack(project_id)]
    output_dir = Path(__file__).resolve().parent.parent / "data" / "delivery"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "机房信息表.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "机房信息"
    fields = ["pod_name", "room_name", "compute", "bus", "param_leaf", "biz_leaf", "mgmt", "sample_leaf"]
    sheet.append(["PoD名称", "机房名称", "计算柜", "总线柜", "参数面Leaf柜", "业务面Leaf柜", "管理面柜", "样本面Leaf柜"])
    for row in rows:
        sheet.append([row.get(field, "") for field in fields])
    for column, width in zip("ABCDEFGH", [18, 20, 14, 14, 18, 18, 16, 18]):
        sheet.column_dimensions[column].width = width
    workbook.save(output_path)
    return ApiResponse(
        data={"saved_path": str(output_path.relative_to(Path(__file__).resolve().parent.parent))},
        meta=_make_meta(project_id),
    )


@router.patch("/chapters/7.1/room-rack/{row_id}")
def update_room_rack(project_id: str, row_id: str, req: RoomRackUpdateReq) -> ApiResponse:
    try:
        data = req.model_dump(exclude_unset=True)
        row = _store.update_room_rack(project_id, row_id, data)
        return ApiResponse(
            data=row.model_dump(),
            meta=_make_meta(project_id),
        )
    except ValueError as e:
        raise _handle_store_error(e)


@router.delete("/chapters/7.1/room-rack/{row_id}", status_code=204, response_class=Response)
def delete_room_rack(
    project_id: str,
    row_id: str,
    confirm: bool = Query(default=False, description="自动解析行需要 confirm=true"),
):
    try:
        _store.delete_room_rack(project_id, row_id, confirm)
    except ValueError as e:
        raise _handle_store_error(e)


# ─── Release ─────────────────────────────────────────────────────────────────

@router.post("/release")
def release(project_id: str, req: ReleaseReq) -> ApiResponse:
    chapters = req.chapters or ["5.1", "5.2", "5.3", "7.1"]
    return ApiResponse(
        data={
            "released_chapters": chapters,
            "proposal_version": "V1.0_mock",
            "status": "success",
        },
        meta=_make_meta(project_id),
    )


# ─── Mock 数据初始化（用于测试） ─────────────────────────────────────────────

@router.post("/mock/init")
def init_mock_data(project_id: str) -> ApiResponse:
    from .proposal_models import ClusterDeviceRow, NetMgmtRow, NetPlaneRow, RoomRackRow

    _store.net_plane_rows[project_id] = [
        NetPlaneRow(
            row_id="np-001",
            device_role="参数面（接入）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-002",
            device_role="样本面（接入）",
            vendor="华为",
            device_model="CE6881-48S6CQ",
            device_version="CE6881-48S6CQ V200R005C20",
            quantity=48,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-003",
            device_role="业务面（接入）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-004",
            device_role="存储面（接入）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-005",
            device_role="网管面（接入）",
            vendor="华为",
            device_model="6857E-48S6CQ",
            device_version="6857E-48S6CQ V200R020C10",
            quantity=12,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-006",
            device_role="参数面（汇聚）",
            vendor="华为",
            device_model="5855E-48T4S2Q",
            device_version="5855E-48T4S2Q V200R020C10",
            quantity=8,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-007",
            device_role="样本面（汇聚）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="HLD 待补",
        ),
        NetPlaneRow(
            row_id="np-008",
            device_role="业务面（汇聚）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-009",
            device_role="存储面（汇聚）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="",
        ),
        NetPlaneRow(
            row_id="np-010",
            device_role="网管面（汇聚）",
            vendor="",
            device_model="",
            device_version="",
            quantity=0,
            data_source="人工录入",
            note="",
        ),
    ]

    _store.net_mgmt_rows[project_id] = [
        NetMgmtRow(row_id="nm-001", server_role="NCE", server_model="通算服务器", quantity=2, data_source="自动解析"),
        NetMgmtRow(row_id="nm-002", server_role="CCAE", server_model="通算服务器", quantity=2, data_source="自动解析"),
        NetMgmtRow(row_id="nm-003", server_role="DME", server_model="通算服务器", quantity=1, data_source="自动解析"),
    ]

    _store.cluster_device_rows[project_id] = [
        ClusterDeviceRow(
            row_id="cd-001",
            cluster_id="CL-01",
            cluster_type="训练",
            super_pod_id="SP-01",
            storage_cluster_id="SC-01",
            zone_id="Z-01",
            ccae_cluster_id="CCAE-01",
            dme_cluster_id="DME-01",
            device_type="昇腾服务器",
            vendor="华为",
            device_model="Atlas 800T A2",
            device_purpose="训练计算节点",
            start_device_name="CL01-SP01-AT800-001",
            end_device_name="CL01-SP01-AT800-064",
            quantity=64,
            data_source="自动解析",
        ),
    ]

    _store.room_rack_rows[project_id] = [
        RoomRackRow(
            row_id="rr-001",
            pod_name="POD1",
            room_name="401",
            compute="A01,A02,A03,A04,A05,A06,A11,A12,A13,A14,A15,A16",
            bus="A07,A08,A09,A10",
            param_leaf="A17,A18",
            biz_leaf="A17,A18",
            mgmt="A17,A18",
            sample_leaf="",
            data_source="自动解析",
        ),
        RoomRackRow(
            row_id="rr-002",
            pod_name="POD2",
            room_name="401",
            compute="B01,B02,B03,B04,B05,B06,B11,B12,B13,B14,B15,B16",
            bus="B07,B08,B09,B10",
            param_leaf="B17,B18",
            biz_leaf="B17,B18",
            mgmt="B17,B18",
            sample_leaf="",
            data_source="自动解析",
        ),
    ]

    return ApiResponse(
        data={"status": "mock data initialized", "project_id": project_id},
        meta=_make_meta(project_id),
    )
