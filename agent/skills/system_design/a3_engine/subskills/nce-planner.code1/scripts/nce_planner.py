"""Offline NCE deployment planner."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import Alignment, Font

from nce_ip_rules import AllocatedRow, allocate_floating_row, allocate_node_rows
from nce_resource_parser import NceResourceRow, load_resource_dataframe, parse_nce_resources
from nce_topology_parser import (
    DEFAULT_SHEET,
    load_mlag_peer_names,
    load_topology_nce_devices,
    split_nce_devices,
)

MODULE_COLUMNS = [
    "设备名称",
    "网络平面",
    "接口名称",
    "IP地址",
    "掩码",
    "网关",
    "VLAN",
    "Bond模式",
    "目的网段",
    "目的掩码",
]
SHEET_NAME = "NCE部署方案"

FB_NODE_PLANES = (
    "NCEFB内部通信网络",
    "NCEFB北向网络",
    "NCEFB南向网络",
    "NCEFB BGP南向网络",
)
FI_NODE_PLANES = (
    "NCEFI北向网络",
    "NCEFI南向网络",
)
FB_FLOATING_PLANES = (
    ("NCEFB北向网络", "北向浮动IP", "北向通信网络"),
    ("NCEFB南向网络", "南向浮动IP", "南向通信网络"),
    ("NCEFB BGP南向网络", "BGP服务南向浮动IP", "BGP服务南向网络"),
)
FI_FLOATING_PLANES = (
    ("NCEFI北向网络", "北向浮动IP", "北向通信网络"),
    ("NCEFI南向网络", "南向浮动IP", "南向通信网络"),
)


@dataclass(frozen=True)
class PlannerInputs:
    topology_path: str
    resource_path: str
    sheet_topology: str = DEFAULT_SHEET
    resource_sheet_index: int = 0
    mlag_list_path: Optional[str] = None


@dataclass(frozen=True)
class PlannerResult:
    final_df: pd.DataFrame
    fb_node_rows: int
    fb_floating_rows: int
    fi_node_rows: int
    fi_floating_rows: int
    switch_mlag_flag: bool


def _row_to_dict(row: AllocatedRow) -> dict:
    return {
        "设备名称": row.device_name,
        "网络平面": row.plane_name,
        "接口名称": row.bond,
        "IP地址": row.ip,
        "掩码": row.mask,
        "网关": row.gateway,
        "VLAN": row.vlan,
        "Bond模式": row.bond_mode,
        "目的网段": row.dest_net,
        "目的掩码": row.dest_mask,
    }


def _title_row(title: str) -> dict:
    return {
        "设备名称": title,
        "网络平面": "网络平面",
        "接口名称": "接口名称",
        "IP地址": "IP地址",
        "掩码": "子网掩码",
        "网关": "网关",
        "VLAN": "VLAN",
        "Bond模式": "Bond模式",
        "目的网段": "目的网段",
        "目的掩码": "目的掩码",
    }


def _empty_row() -> dict:
    return {col: "" for col in MODULE_COLUMNS}


def _bond_mode_for_devices(devices: List[str], mlag_list_path: Optional[str]) -> tuple[str, bool]:
    if not mlag_list_path:
        return "mode1", False
    mlag_peers = load_mlag_peer_names(mlag_list_path)
    matched = any(device in mlag_peers for device in devices)
    return ("mode4(lacp)" if matched else "mode1"), matched


def _allocate_node_group(
    *,
    devices: List[str],
    plane_names: tuple[str, ...],
    resources: Dict[str, NceResourceRow],
    bond_mode: str,
) -> List[dict]:
    rows: List[dict] = []
    bond_index = 0
    if not devices:
        return rows
    for plane_name in plane_names:
        resource = resources.get(plane_name)
        if resource is None:
            continue
        allocated = allocate_node_rows(
            devices=devices,
            plane_name=resource.plane_name,
            ip_pool=resource.ip_pool,
            mask=resource.mask,
            vlan_raw=resource.vlan_raw,
            gateway_position=resource.gateway_position,
            gateway_address=resource.gateway_address,
            bond_name=f"bond{bond_index}",
            bond_mode=bond_mode,
        )
        rows.extend(_row_to_dict(row) for row in allocated)
        bond_index += 1
    return sorted(rows, key=lambda item: item["设备名称"])


def _allocate_floating_group(
    *,
    devices: List[str],
    floating_specs: tuple[tuple[str, str, str], ...],
    resources: Dict[str, NceResourceRow],
    bond_mode: str,
) -> List[dict]:
    rows: List[dict] = []
    bond_index = 0
    if not devices:
        return rows
    for source_plane, usage_name, display_plane in floating_specs:
        resource = resources.get(source_plane)
        if resource is None:
            continue
        allocated = allocate_floating_row(
            devices=devices,
            usage_name=usage_name,
            display_plane_name=display_plane,
            source_plane_name=resource.plane_name,
            ip_pool=resource.ip_pool,
            mask=resource.mask,
            vlan_raw=resource.vlan_raw,
            gateway_position=resource.gateway_position,
            gateway_address=resource.gateway_address,
            bond_name=f"bond{bond_index}",
            bond_mode=bond_mode,
        )
        rows.append(_row_to_dict(allocated))
        bond_index += 1
    return rows


def run_planner(inputs: PlannerInputs) -> PlannerResult:
    devices = load_topology_nce_devices(inputs.topology_path, inputs.sheet_topology)
    fb_devices, fi_devices = split_nce_devices(devices)
    if not fb_devices and not fi_devices:
        raise ValueError("未找到名称为 NCEFB 或 NCEFI 的设备，请检查端口连线表内容")

    resource_df = load_resource_dataframe(inputs.resource_path, inputs.resource_sheet_index)
    resources = parse_nce_resources(resource_df)
    bond_mode, switch_mlag_flag = _bond_mode_for_devices(devices, inputs.mlag_list_path)

    fb_nodes = _allocate_node_group(
        devices=fb_devices,
        plane_names=FB_NODE_PLANES,
        resources=resources,
        bond_mode=bond_mode,
    )
    fb_floating = _allocate_floating_group(
        devices=fb_devices,
        floating_specs=FB_FLOATING_PLANES,
        resources=resources,
        bond_mode=bond_mode,
    )
    fi_nodes = _allocate_node_group(
        devices=fi_devices,
        plane_names=FI_NODE_PLANES,
        resources=resources,
        bond_mode=bond_mode,
    )
    fi_floating = _allocate_floating_group(
        devices=fi_devices,
        floating_specs=FI_FLOATING_PLANES,
        resources=resources,
        bond_mode=bond_mode,
    )

    rows: List[dict] = []
    rows.extend(fb_nodes)
    rows.append(_empty_row())
    rows.append(_title_row("网络用途_NCEFB"))
    rows.extend(fb_floating)
    rows.append(_empty_row())
    rows.append(_title_row("设备名称"))
    rows.extend(fi_nodes)
    rows.append(_empty_row())
    rows.append(_title_row("网络用途_NCEFI"))
    rows.extend(fi_floating)

    final_df = pd.DataFrame(rows, columns=MODULE_COLUMNS).fillna("")
    if final_df.empty:
        raise ValueError("没有生成任何 NCE 部署规划结果")

    return PlannerResult(
        final_df=final_df,
        fb_node_rows=len(fb_nodes),
        fb_floating_rows=len(fb_floating),
        fi_node_rows=len(fi_nodes),
        fi_floating_rows=len(fi_floating),
        switch_mlag_flag=switch_mlag_flag,
    )


def write_nce_excel(result: PlannerResult, output_path: str) -> Path:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    result.final_df.to_excel(out, sheet_name=SHEET_NAME, index=False)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME]
    align_center = Alignment(horizontal="center", vertical="center")
    bold_font = Font(bold=True)
    target_names = {"网络用途_NCEFB", "设备名称", "网络用途_NCEFI"}

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align_center

    for row_idx in range(1, ws.max_row + 1):
        if ws.cell(row=row_idx, column=1).value in target_names:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).font = bold_font

    wb.save(out)
    return out


def validate_inputs(inputs: PlannerInputs) -> None:
    load_topology_nce_devices(inputs.topology_path, inputs.sheet_topology)
    resource_df = load_resource_dataframe(inputs.resource_path, inputs.resource_sheet_index)
    parse_nce_resources(resource_df)
    if inputs.mlag_list_path and not Path(inputs.mlag_list_path).is_file():
        raise ValueError(f"MLAG 列表文件不存在: {inputs.mlag_list_path}")
