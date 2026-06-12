"""从本地 ZTP LLD Excel 读取 cfg 数据 — 对齐 CfgDataReader（仅本地文件）。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from data_struct import CfgData, ElementInfo


class OfflineCfgDataReader:
    def __init__(self, conf_elements: dict, file_path: Path, sheet_name: str = "网络IP规划"):
        self.elements: dict[str, ElementInfo] = {}
        for key, value in conf_elements.items():
            self.elements[key] = ElementInfo(key, value)
        self.file_path = file_path
        self.sheet_name = sheet_name
        self.cfgDataList: list[CfgData] = []
        self.row_start = 0
        self.column_start = 0

    def load(self) -> list[CfgData]:
        book = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
        sheets = book.sheetnames
        if self.sheet_name not in sheets:
            for real in sheets:
                if real.strip() == self.sheet_name:
                    self.sheet_name = real
                    break
        sheet = book[self.sheet_name]
        self._get_data_pos(sheet)
        max_row = sheet.max_row
        self._generate_cfg_data_list(max_row, sheet)
        book.close()
        return self.cfgDataList

    def _get_data_pos(self, sheet) -> None:
        is_find = False
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and self._check_and_set_element(cell.value, cell.row, cell.column):
                    is_find = True
            if is_find:
                missing = self._element_not_find()
                if missing and missing != self.elements["level"].elementNameInExcel:
                    raise ValueError(f"表中缺少必要的配置元素或前面几行中有重复的表头名：{missing}")
                break

    def _element_not_find(self) -> str:
        names = []
        for ele in self.elements.values():
            if ele.rowNum == 0 or ele.columnNum == 0:
                names.append(ele.elementNameInExcel)
        return ", ".join(names)

    def _check_and_set_element(self, ele_name_in_excel, row, column) -> bool:
        for ele in self.elements.values():
            if ele_name_in_excel == ele.elementNameInExcel:
                if ele.rowNum != 0 or ele.columnNum != 0:
                    raise ValueError(f"表中存在重复的配置元素：{ele_name_in_excel}")
                ele.setPos(row, column)
                if self.row_start == 0 and self.column_start == 0:
                    self.row_start = row + 1
                    self.column_start = column
                return True
        return False

    def _get_columns_num(self) -> dict:
        return {name: self.elements[name].columnNum for name in self.elements}

    def _generate_cfg_data_list(self, max_row_num: int, sheet) -> None:
        columns = self._get_columns_num()

        for row in sheet.iter_rows(min_row=self.row_start, max_row=max_row_num, values_only=True):
            def gv(field_name):
                col_index = columns[field_name]
                if not col_index:
                    return None
                return row[col_index - 1]

            device_name = gv("deviceName")
            if device_name is None or str(device_name).strip() == "":
                continue

            gateway_raw = gv("gateway")
            gateway, sub_mask = None, None
            if gateway_raw:
                parts = str(gateway_raw).split("/")
                if len(parts) == 2:
                    gateway, sub_mask = parts[0], parts[1]

            room_name, rack_num, rack_loc = gv("roomName"), gv("rackNum"), gv("rackLocation")
            location = (
                f"{room_name}-{rack_num}-{rack_loc}"
                if all([room_name, rack_num, rack_loc])
                else None
            )

            cfg = CfgData(
                device_name,
                gv("esn"),
                location,
                gv("level"),
                gv("methIP"),
                gateway,
                sub_mask,
                gv("switchId"),
                gv("bgpAs"),
                gv("loopback0"),
                gv("loopback1"),
                gv("loopback2"),
                gv("loopback3"),
                gv("loopback4"),
                gv("loopback5"),
                gv("loopback6"),
            )
            self.cfgDataList.append(cfg)
