"""从项目信息收集表 ZTP配置 sheet 读取公共配置 — 对齐 generate_common_dict。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def generate_common_dict(file_path: Path, sheet_name: str = "ZTP配置") -> dict:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        for real in wb.sheetnames:
            if real.strip() == sheet_name:
                sheet_name = real
                break
    sheet = wb[sheet_name]
    common_dict: dict = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        var_name = row[4]
        var_value = row[2]
        if var_value is None:
            var_value = ""
        if var_name is None:
            continue
        if var_value in ("是", "否"):
            key = var_name
        else:
            key = "{{" + var_name + "}}"
            if key == "{{bmcUserPassword}}":
                key = "{{InerMgmtPassword}}"
        common_dict[key] = var_value

    wb.close()
    return common_dict
