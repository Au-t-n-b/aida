"""Offline LLD integrate — local files only, no EDM/DB."""

from __future__ import annotations

import warnings
from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from lld_staging import scan_latest_by_plane


def copy_sheet_with_format(src_wb, src_sheet_name: str, dst_wb, new_sheet_name: str):
    if src_sheet_name not in src_wb.sheetnames:
        return None
    ws = src_wb[src_sheet_name]
    if new_sheet_name in dst_wb.sheetnames:
        del dst_wb[new_sheet_name]
    new_ws = dst_wb.create_sheet(title=new_sheet_name[:31])
    for col_letter, dimension in ws.column_dimensions.items():
        new_ws.column_dimensions[col_letter].width = dimension.width
        new_ws.column_dimensions[col_letter].hidden = dimension.hidden
    for row_idx, dimension in ws.row_dimensions.items():
        new_ws.row_dimensions[row_idx].height = dimension.height
        new_ws.row_dimensions[row_idx].hidden = dimension.hidden
    # 样式按源 StyleArray 去重：每种样式只在目标工作簿注册一次（拿到 dst 的 _style 索引数组），
    # 之后每个单元格直接赋 _style（仅整型数组拷贝，省去逐 cell 的 font/border/fill add() 哈希）。
    # 大表（如 007「超平面端口互联」≈26 万单元格、仅 9 种样式）从 ~110s 降到 ~2s。
    # scratch 用独立临时 sheet 注册样式，避免污染真实数据单元格（注册的样式表是工作簿级，删 sheet 不失效）。
    scratch_ws = dst_wb.create_sheet(title="__style_scratch__")
    scratch = scratch_ws.cell(row=1, column=1)
    style_cache: dict = {}
    for row in ws.iter_rows():
        for cell in row:
            new_cell = new_ws[cell.coordinate]
            new_cell.value = cell.value
            if cell.has_style:
                key = tuple(cell._style)
                style_array = style_cache.get(key)
                if style_array is None:
                    scratch.font = copy(cell.font)
                    scratch.border = copy(cell.border)
                    scratch.fill = copy(cell.fill)
                    scratch.number_format = cell.number_format
                    scratch.protection = copy(cell.protection)
                    scratch.alignment = copy(cell.alignment)
                    style_array = copy(scratch._style)
                    style_cache[key] = style_array
                new_cell._style = copy(style_array)
    del dst_wb["__style_scratch__"]
    for merged_cell_range in ws.merged_cells.ranges:
        new_ws.merge_cells(str(merged_cell_range))
    return new_ws


def adjust_column_width(df: pd.DataFrame, writer: pd.ExcelWriter, sheet_name: str) -> None:
    if df.empty:
        df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        return
    trimmed = sheet_name[:31]
    df.to_excel(writer, sheet_name=trimmed, index=False)
    worksheet = writer.sheets[trimmed]
    header_widths = [len(str(col)) + 5 for col in df.columns]
    for idx, col in enumerate(df.columns):
        max_len = max(
            df[col].astype(str).map(len).max() if not df[col].empty else 0,
            header_widths[idx],
        )
        worksheet.column_dimensions[get_column_letter(idx + 1)].width = min(max_len + 2, 50)


def write_lld_data_from_file(file_path: Path, writer: pd.ExcelWriter, switch_dfs: List[pd.DataFrame]) -> None:
    excel_file = pd.ExcelFile(file_path)
    for sheet_name in excel_file.sheet_names:
        df = excel_file.parse(sheet_name, keep_default_na=False, na_values=[])
        if "交换机地址规划" in sheet_name:
            switch_dfs.append(df)
        else:
            adjust_column_width(df, writer, sheet_name)


def copy_simulation_workbook(src_path: Path, dst_wb) -> None:
    if not src_path.is_file():
        return
    src_wb = load_workbook(src_path, data_only=False)
    exclude_keywords = ["文档说明"]
    for sheet_name in src_wb.sheetnames:
        if any(kw in sheet_name for kw in exclude_keywords):
            continue
        trimmed = sheet_name.split(" for ")[0][:31]
        copy_sheet_with_format(src_wb, sheet_name, dst_wb, trimmed)


def copy_template_sheet(templates_dir: Path, template_file: str, sheet_name: str, dst_wb) -> bool:
    path = templates_dir / template_file
    if not path.is_file():
        warnings.warn(f"template missing, skip: {path}")
        return False
    src_wb = load_workbook(path)
    copy_sheet_with_format(src_wb, sheet_name, dst_wb, sheet_name[:31])
    return True


def general_config_from_resource(resource_path: Path, templates_dir: Path, dst_wb) -> None:
    template_path = templates_dir / "通用配置信息.xlsx"
    if not template_path.is_file():
        warnings.warn(f"skip general_config: {template_path} not found")
        return
    try:
        input_df = pd.read_excel(resource_path, sheet_name="ZTP配置", header=None)
    except Exception:
        warnings.warn("skip general_config: ZTP配置 sheet not found in resource")
        return
    input_dict = {}
    for b_val, c_val in zip(input_df.iloc[:, 1], input_df.iloc[:, 2]):
        if pd.isna(b_val):
            continue
        input_dict[str(b_val).strip().lower()] = c_val

    template_df = pd.read_excel(template_path, header=None)

    def map_value(row):
        a_val = row[template_df.columns[0]]
        if pd.isna(a_val):
            return row[template_df.columns[1]]
        key = str(a_val).strip().lower()
        matched = input_dict.get(key)
        return matched if matched is not None else row[template_df.columns[1]]

    template_df.iloc[:, 1] = template_df.apply(map_value, axis=1)
    src_wb = load_workbook(template_path)
    dst_sheet = copy_sheet_with_format(src_wb, "通用配置信息", dst_wb, "通用配置信息")
    if dst_sheet is None:
        return
    for r_idx, row in enumerate(template_df.itertuples(index=False), start=1):
        for c_idx, value in enumerate(row, start=1):
            dst_sheet.cell(row=r_idx, column=c_idx).value = value


def integrate_lld(
    project_name: str,
    output_path: Path,
    scan_dirs: Sequence[Path],
    templates_dir: Optional[Path] = None,
    simulation_files: Optional[Dict[str, Path]] = None,
    resource_path: Optional[Path] = None,
) -> Path:
    dirs: List[Path] = list(scan_dirs)
    latest_by_plane = scan_latest_by_plane(dirs)
    if not latest_by_plane:
        raise ValueError("未找到可融合的 A3*.xlsx 部分产物（请指定 --scan-dir）")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    writer = pd.ExcelWriter(buffer, engine="openpyxl")
    switch_dfs: List[pd.DataFrame] = []

    try:
        wb = writer.book
        if templates_dir and templates_dir.is_dir():
            copy_template_sheet(templates_dir, "版本修订记录.xlsx", "版本修订记录", wb)
            if resource_path:
                general_config_from_resource(resource_path, templates_dir, wb)

        if simulation_files:
            for _key, sim_path in simulation_files.items():
                if sim_path and sim_path.is_file():
                    copy_simulation_workbook(sim_path, wb)

        for plane_key in sorted(latest_by_plane.keys()):
            write_lld_data_from_file(latest_by_plane[plane_key], writer, switch_dfs)

        if switch_dfs:
            merged = pd.concat(switch_dfs, ignore_index=True)
            adjust_column_width(merged, writer, "交换机地址规划")
    finally:
        writer.close()

    output_path.write_bytes(buffer.getvalue())
    return output_path
