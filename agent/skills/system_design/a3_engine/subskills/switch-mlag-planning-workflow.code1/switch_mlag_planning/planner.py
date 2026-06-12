#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Standalone switch MLAG planner.

This module mirrors the core logic in ``a3_switch_mlag.py`` while using only
local Excel files. It intentionally avoids CPCIA_AGENT project services such as
DB/EDM upload, logging, and dialogue display so the skill can run independently.
"""

from __future__ import annotations

import io
import ipaddress
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import load_workbook


WEB_NETWORK_TYPE_CONFIG = {
    "存储面端口互联 | 样本面端口互联": "计算样本面",
    "计算管理面端口互联": "计算管理面",
    "计算业务面端口互联": "计算业务面",
    "计算管存面端口互联": "计算管存面",
    "参数面端口互联": "计算参数面",
    "样本面端口互联 | 数据面端口互联": "存储样本面",
    "存储业务面端口互联": "存储业务面",
    "存储管理面端口互联": "存储管理面",
    "计算带外管理面端口互联": "计算带外管理面",
    "网络带外管理面端口互联": "网络带外管理面",
    "灵衢带外管理面端口互联": "灵衢带外管理面",
    "存储带外管理面端口互联": "存储带外管理面",
    "带外管理面端口互联": "带外管理面",
    "超平面端口互联": "超平面",
    "网络业务地址规划": "其它",
}

OUTPUT_COLUMNS = [
    "网络平面",
    "本端设备",
    "本端接口",
    "本端接口IP地址",
    "本端接口掩码",
    "本端ETH-TRUNK",
    "本端VLAN",
    "对端设备",
    "对端接口",
    "对端接口IP地址",
    "对端接口掩码",
    "对端ETH-TRUNK",
    "对端VLAN",
    "PVID",
    "端口类型",
    "标签",
]

DEFAULT_TOPOLOGY_FILE = "建模仿真输出文档007-端口连线表.xlsx"
DEFAULT_RESOURCE_FILE = "项目信息收集表.xlsx"
DEFAULT_OUTPUT_FILE = "A3交换机MLAG规划.xlsx"
DEFAULT_TEMP_FILE = "交换机MLAG规划.xlsx"
OUTPUT_SHEET_NAME = "交换机MLAG规划"


@dataclass(frozen=True)
class NetResourceInfo:
    local_ip: str
    peer_ip: str
    mask: str


@dataclass(frozen=True)
class PlanningResult:
    result_df: pd.DataFrame
    output_file: Path
    temp_file: Path | None
    processed_sheets: tuple[str, ...]
    sheet_row_counts: dict[str, int]
    used_project_filenames: bool


def _clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return re.sub(r"\.0+$", "", text)


def _empty_output_df() -> pd.DataFrame:
    return pd.DataFrame(columns=OUTPUT_COLUMNS)


def _read_raw_sheet(excel_path: str | Path, sheet_name: str) -> pd.DataFrame:
    excel_path = Path(excel_path)
    candidates = _sheet_candidates(sheet_name)
    xls = pd.ExcelFile(excel_path)
    last_error: Exception | None = None
    for candidate in candidates:
        try:
            return pd.read_excel(excel_path, sheet_name=candidate, header=None)
        except ValueError as exc:
            last_error = exc
    raise ValueError(
        f"未在 {excel_path.name} 中找到 sheet：{sheet_name}；可用 sheet：{', '.join(xls.sheet_names)}"
    ) from last_error


def _sheet_candidates(sheet_name: str) -> list[str]:
    candidates: list[str] = []
    for candidate in [sheet_name, *sheet_name.split("|")]:
        candidate = candidate.strip()
        if candidate and candidate not in candidates:
            candidates.append(candidate)
    return candidates


def _find_header_row(df: pd.DataFrame, keyword: str = "设备命名") -> int:
    for idx, row in df.iterrows():
        if row.astype(str).str.contains(keyword, case=False, na=False).any():
            return int(idx)
    raise ValueError(f"未找到包含'{keyword}'的行，请检查表格内容")


def get_switch_data(topology_file: str | Path, sheet_name: str) -> tuple[pd.DataFrame, int]:
    """Read one 007 sheet and return leaf-leaf/spine-spine MLAG candidate links."""
    df_all = _read_raw_sheet(topology_file, sheet_name)
    header_row_idx = _find_header_row(df_all, "设备命名")
    df = df_all.iloc[header_row_idx + 1 :].reset_index(drop=True)
    df = df.dropna(how="all")
    if df.empty or len(df.columns) < 8:
        return pd.DataFrame(), 0

    result_df = df.rename(
        columns={
            df.columns[0]: "本端交换机",
            df.columns[1]: "本端端口",
            df.columns[2]: "本端接口带宽",
            df.columns[-4]: "对端带宽",
            df.columns[-2]: "对端端口",
            df.columns[-3]: "对端接口类型",
            df.columns[-1]: "对端交换机",
        }
    )
    result_df = result_df[
        result_df["本端交换机"].map(_clean_cell).str.contains("leaf|spine", case=False, na=False)
    ].drop_duplicates()

    local_leaf = result_df["本端交换机"].map(_clean_cell).str.contains("leaf", case=False, na=False)
    peer_leaf = result_df["对端交换机"].map(_clean_cell).str.contains("leaf", case=False, na=False)
    local_spine = result_df["本端交换机"].map(_clean_cell).str.contains("spine", case=False, na=False)
    peer_spine = result_df["对端交换机"].map(_clean_cell).str.contains("spine", case=False, na=False)

    result_df = result_df[(local_leaf & peer_leaf) | (local_spine & peer_spine)].copy()
    return result_df, int(result_df.shape[0])


def _find_column(columns: Iterable[object], pattern: str, *, required: bool = True) -> object | None:
    for col in columns:
        if re.search(pattern, str(col), re.IGNORECASE):
            return col
    if required:
        raise ValueError(f"未找到匹配列：{pattern}")
    return None


def _read_resource_sheet(resource_file: str | Path) -> pd.DataFrame:
    resource_file = Path(resource_file)
    xls = pd.ExcelFile(resource_file)
    candidates = ["网络资源需求表", "资源表"] + [
        name for name in xls.sheet_names if name not in {"网络资源需求表", "资源表"}
    ]
    for sheet in candidates:
        try:
            df = pd.read_excel(resource_file, sheet_name=sheet, header=0)
        except ValueError:
            continue
        df.columns = [_clean_cell(col) for col in df.columns]
        if any(str(col) == "网络平面" for col in df.columns):
            return df
    raise ValueError(f"{resource_file.name} 中未找到包含“网络平面”列的资源表")


def get_net_ni_resource_info(resource_file: str | Path, row_name: str = "MLAG") -> NetResourceInfo:
    """Read the MLAG row from the resource workbook."""
    df = _read_resource_sheet(resource_file)
    plane_col = _find_column(df.columns, r"^网络平面$")
    row = df[df[plane_col].map(_clean_cell) == row_name]
    if row.empty:
        raise ValueError(f"未找到{row_name}所在行。")

    ip_pool_col = _find_column(df.columns, r"^地址池\*?$|地址池")
    mask_col = _find_column(df.columns, r"最小规划掩码|掩码")
    first = row.iloc[0]
    ip_pool = _clean_cell(first[ip_pool_col])
    mask = _clean_cell(first[mask_col])
    mask_match = re.search(r"\d{1,2}", mask)
    if not mask_match:
        raise ValueError(f"最小规划掩码为空或非法：{mask}")

    pair = find_consecutive_pair_in_pool(ip_pool, int(mask_match.group(0)))
    if pair is not None:
        local_ip, peer_ip = pair
    else:
        local_ip, peer_ip = "", ""
    return NetResourceInfo(local_ip=local_ip, peer_ip=peer_ip, mask=mask_match.group(0))


def _concat_interface(prefix_value: object, port_value: object) -> str:
    return f"{_clean_cell(prefix_value)}{_clean_cell(port_value)}".strip()


def _extract_bandwidth_value(bandwidth_str: object) -> int:
    bandwidth = _clean_cell(bandwidth_str)
    if bandwidth == "GE":
        return 1
    num = "".join(filter(str.isdigit, bandwidth))
    if not num:
        raise ValueError(f"端口连线表中输入的带宽格式错误，请检查，可能错误值为：{bandwidth_str}")
    return int(num)


def allocate_connection(df: pd.DataFrame, net_resource_info: NetResourceInfo, network_type: str) -> pd.DataFrame:
    """Allocate DAD and PEER-LINK rows with the same rules as a3_switch_mlag.py."""
    result_df = _empty_output_df()
    if df.empty:
        return result_df

    df = df.copy()
    bandwidth = df["本端接口带宽"].drop_duplicates()
    df["本端接口"] = [
        _concat_interface(row.iloc[2], row.iloc[1]) for _, row in df.iterrows()
    ]
    df["对端接口"] = [
        _concat_interface(row.iloc[5], row.iloc[7]) for _, row in df.iterrows()
    ]

    local_ip = net_resource_info.local_ip
    peer_ip = net_resource_info.peer_ip
    mask = net_resource_info.mask

    if len(bandwidth) == 2:
        bandwidth_max = bandwidth.map(_extract_bandwidth_value).max()
        for _, row in df.iterrows():
            device1 = row["本端交换机"]
            device2 = row["对端交换机"]
            if str(bandwidth_max) in _clean_cell(row["本端接口带宽"]):
                current_trunk = 1
                label = "PEER-LINK"
            else:
                current_trunk = 0
                label = "DAD"
            result_df.loc[len(result_df)] = [
                network_type,
                device1,
                row["本端接口"],
                "",
                "",
                current_trunk,
                "NA",
                device2,
                row["对端接口"],
                "",
                "",
                current_trunk,
                "NA",
                "NA",
                "NA",
                label,
            ]
    elif len(bandwidth) == 1:
        port_counts = df.groupby("本端交换机").size()
        df["端口数量"] = df["本端交换机"].map(port_counts)
        df["端口序号"] = df.groupby("本端交换机").cumcount() + 1
        df = df.reset_index(drop=True)
        for _, row in df.iterrows():
            device1 = row["本端交换机"]
            device2 = row["对端交换机"]
            if (row["端口数量"] % 2 == 1 and row["端口序号"] < 2) or (
                row["端口数量"] % 2 == 0 and row["端口序号"] < 3
            ):
                current_trunk = 0
                label = "DAD"
            else:
                current_trunk = 1
                label = "PEER-LINK"
            result_df.loc[len(result_df)] = [
                network_type,
                device1,
                row["本端接口"],
                "",
                "",
                current_trunk,
                "NA",
                device2,
                row["对端接口"],
                "",
                "",
                current_trunk,
                "NA",
                "NA",
                "NA",
                label,
            ]
    else:
        return result_df

    result_df.loc[result_df["标签"] == "DAD", "本端接口IP地址"] = local_ip
    result_df.loc[result_df["标签"] == "DAD", "对端接口IP地址"] = peer_ip
    result_df.loc[result_df["标签"] == "DAD", "本端接口掩码"] = str(mask)
    result_df.loc[result_df["标签"] == "DAD", "对端接口掩码"] = str(mask)
    return result_df


def switch_connect(
    topology_file: str | Path,
    sheet_name: str,
    net_resource_info: NetResourceInfo,
) -> pd.DataFrame:
    if sheet_name not in WEB_NETWORK_TYPE_CONFIG:
        return _empty_output_df()
    web_network_type_name = WEB_NETWORK_TYPE_CONFIG[sheet_name]
    if web_network_type_name == "超平面":
        return _empty_output_df()

    switch_data, _ = get_switch_data(topology_file, sheet_name)
    result_df = allocate_connection(switch_data, net_resource_info, web_network_type_name)
    if "带外" in web_network_type_name and not result_df.empty:
        result_df = result_df[
            ~result_df["本端设备"].map(_clean_cell).str.contains("LEAF", case=False, na=False)
        ].copy()
    return result_df


def available_mlag_sheets(topology_file: str | Path, *, project_exact: bool = True) -> list[str]:
    xls = pd.ExcelFile(topology_file)
    if project_exact:
        return [sheet for sheet in xls.sheet_names if sheet in WEB_NETWORK_TYPE_CONFIG]

    workbook_sheets = set(xls.sheet_names)
    result: list[str] = []
    for config_sheet in WEB_NETWORK_TYPE_CONFIG:
        if any(candidate in workbook_sheets for candidate in _sheet_candidates(config_sheet)):
            result.append(config_sheet)
    return result


def resolve_config_sheet(sheet_name: str) -> str:
    if sheet_name in WEB_NETWORK_TYPE_CONFIG:
        return sheet_name
    for config_sheet in WEB_NETWORK_TYPE_CONFIG:
        if sheet_name in _sheet_candidates(config_sheet):
            return config_sheet
    raise ValueError(f"不支持的 sheet：{sheet_name}")


def _normalize_selected_sheets(sheets: Sequence[str]) -> list[str]:
    result: list[str] = []
    for sheet_name in sheets:
        config_sheet = resolve_config_sheet(sheet_name)
        if config_sheet not in result:
            result.append(config_sheet)
    return result


def _read_existing_result(existing_file: Path) -> pd.DataFrame:
    if not existing_file.exists():
        return _empty_output_df()
    df = pd.read_excel(existing_file, sheet_name=0, header=0)
    for col in OUTPUT_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    return df[OUTPUT_COLUMNS]


def build_project_output_paths(
    output_dir: str | Path,
    user_id: str,
    project_id: str,
) -> tuple[Path, Path]:
    output_dir = Path(output_dir)
    temp_file = output_dir / f"{user_id}_交换机MLAG规划.xlsx"
    output_file = output_dir / f"{user_id}_{project_id}_A3交换机MLAG规划.xlsx"
    return output_file, temp_file


def generate_switch_mlag(
    topology_file: str | Path = DEFAULT_TOPOLOGY_FILE,
    resource_file: str | Path = DEFAULT_RESOURCE_FILE,
    output_file: str | Path | None = None,
    *,
    sheets: Sequence[str] | None = None,
    merge_existing: bool = False,
    existing_file: str | Path | None = None,
    temp_file: str | Path | None = None,
    user_id: str | None = None,
    project_id: str | None = None,
    output_dir: str | Path = ".",
    project_exact_sheets: bool = True,
) -> PlanningResult:
    """Generate the switch MLAG workbook from local Excel inputs."""
    topology_file = Path(topology_file)
    resource_file = Path(resource_file)
    used_project_filenames = False
    if output_file is None and temp_file is None and user_id and project_id:
        output_file, temp_path = build_project_output_paths(output_dir, user_id, project_id)
        used_project_filenames = True
    else:
        output_file = Path(output_file or DEFAULT_OUTPUT_FILE)
        temp_path = Path(temp_file) if temp_file else Path(DEFAULT_TEMP_FILE)
    existing_path = Path(existing_file) if existing_file else temp_path

    selected_sheets = (
        _normalize_selected_sheets(sheets)
        if sheets
        else available_mlag_sheets(topology_file, project_exact=project_exact_sheets)
    )
    if not selected_sheets:
        raise ValueError("端口连线表中没有可用于交换机MLAG规划的 sheet")

    net_resource_info = get_net_ni_resource_info(resource_file, "MLAG")
    merged_df = _empty_output_df()
    if merge_existing:
        if existing_path and existing_path.exists():
            merged_df = _read_existing_result(existing_path)
        elif output_file.exists():
            merged_df = _read_existing_result(output_file)

    processed: list[str] = []
    sheet_row_counts: dict[str, int] = {}
    for sheet_name in selected_sheets:
        result_df = switch_connect(topology_file, sheet_name, net_resource_info)
        network_type = WEB_NETWORK_TYPE_CONFIG[sheet_name]
        if not merged_df.empty and "网络平面" in merged_df.columns:
            merged_df = merged_df[~merged_df["网络平面"].isin([network_type])]
        if not result_df.empty:
            merged_df = pd.concat([merged_df, result_df], axis=0, ignore_index=True)
        merged_df = standard_merge_df(merged_df)
        if temp_path:
            temp_path.parent.mkdir(parents=True, exist_ok=True)
            merged_df.to_excel(temp_path, sheet_name=OUTPUT_SHEET_NAME, index=False)
        processed.append(sheet_name)
        sheet_row_counts[sheet_name] = int(result_df.shape[0])

    if temp_path and temp_path.exists():
        final_source_df = _read_existing_result(temp_path)
    else:
        final_source_df = standard_merge_df(merged_df)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    try:
        merge_excel_regions(final_source_df, output_file)
    except Exception:
        final_source_df.to_excel(output_file, sheet_name=OUTPUT_SHEET_NAME, index=False)
    return PlanningResult(
        result_df=final_source_df,
        output_file=output_file,
        temp_file=temp_path,
        processed_sheets=tuple(processed),
        sheet_row_counts=sheet_row_counts,
        used_project_filenames=used_project_filenames,
    )


def _normalize_empty(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _to_clean_str(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\.0+$", "", str(value).strip())


def standard_merge_df(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in OUTPUT_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[OUTPUT_COLUMNS]
    cols = ["本端VLAN", "对端VLAN", "PVID", "端口类型"]
    df[cols] = df[cols].fillna("NA")
    for col in ["本端设备", "本端接口IP地址", "本端接口掩码", "本端ETH-TRUNK", "对端接口IP地址", "对端接口掩码", "对端ETH-TRUNK"]:
        df[col] = df[col].apply(_normalize_empty)
    df["本端接口掩码"] = df["本端接口掩码"].apply(_to_clean_str)
    df["本端ETH-TRUNK"] = df["本端ETH-TRUNK"].apply(_to_clean_str)
    df["对端接口掩码"] = df["对端接口掩码"].apply(_to_clean_str)
    df["对端ETH-TRUNK"] = df["对端ETH-TRUNK"].apply(_to_clean_str)
    return df


def merge_excel_regions(df: pd.DataFrame, output_file: str | Path) -> bytes:
    df = standard_merge_df(df)
    df_sorted = df.sort_values(
        by=["网络平面", "本端设备", "本端ETH-TRUNK", "标签"], kind="stable"
    ).reset_index(drop=True)

    out_io = io.BytesIO()
    with pd.ExcelWriter(out_io, engine="openpyxl") as writer:
        df_sorted.to_excel(writer, index=False, sheet_name=OUTPUT_SHEET_NAME)

    out_io.seek(0)
    wb = load_workbook(out_io)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    col_idx = {name: i + 1 for i, name in enumerate(header)}

    local_ip_col = col_idx["本端接口IP地址"]
    local_mask_col = col_idx["本端接口掩码"]
    local_eth_col = col_idx["本端ETH-TRUNK"]
    peer_ip_col = col_idx["对端接口IP地址"]
    peer_mask_col = col_idx["对端接口掩码"]
    peer_eth_col = col_idx["对端ETH-TRUNK"]

    grouped = df_sorted.groupby(["本端设备", "本端接口IP地址", "本端ETH-TRUNK"], sort=False)
    for _, group in grouped:
        if len(group) <= 1:
            continue
        start_row = int(group.index.min()) + 2
        end_row = int(group.index.max()) + 2
        ws.merge_cells(start_row=start_row, start_column=local_ip_col, end_row=end_row, end_column=local_ip_col)
        ws.merge_cells(start_row=start_row, start_column=local_mask_col, end_row=end_row, end_column=local_mask_col)
        ws.merge_cells(start_row=start_row, start_column=local_eth_col, end_row=end_row, end_column=local_eth_col)
        ws.merge_cells(start_row=start_row, start_column=peer_ip_col, end_row=end_row, end_column=peer_ip_col)
        ws.merge_cells(start_row=start_row, start_column=peer_mask_col, end_row=end_row, end_column=peer_mask_col)
        ws.merge_cells(start_row=start_row, start_column=peer_eth_col, end_row=end_row, end_column=peer_eth_col)

    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    result_io = io.BytesIO()
    wb.save(result_io)
    return result_io.getvalue()


def find_consecutive_pair_in_pool(pool_range: str, mask: int) -> Optional[Tuple[str, str]]:
    start_str, end_str = pool_range.split("-", 1)
    start = ipaddress.IPv4Address(start_str.strip())
    end = ipaddress.IPv4Address(end_str.strip())

    cur = int(start)
    last = int(end)
    while cur < last:
        ip1 = ipaddress.IPv4Address(cur)
        ip2 = ipaddress.IPv4Address(cur + 1)
        if int(ip2) > last:
            break
        net = ipaddress.ip_network(f"{ip1}/{mask}", strict=False)
        if ip2 in net:
            if (
                ip1 != net.network_address
                and ip1 != net.broadcast_address
                and ip2 != net.network_address
                and ip2 != net.broadcast_address
            ):
                return str(ip1), str(ip2)
        cur += 1
    return None
