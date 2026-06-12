"""最小冒烟测试：构造计算带外管理面端口连线表与项目信息收集表，跑通 L2/L3。"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from oob_interconnect import run_l2, run_l3  # noqa: E402


def _sheetnames(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _build_topology(path: Path) -> None:
    rows = [
        ["设备命名", "本端端口", "本端带宽", "_", "_", "对端端口", "对端类型", "对端带宽", "对端设备"],
        ["DWGL-LEAF-1", "GE0/1", "10G", "", "", "GE0/1", "ETH", "10G", "SPINE-1"],
        ["DWGL-LEAF-1", "GE0/2", "10G", "", "", "GE0/2", "ETH", "10G", "SPINE-2"],
        ["DWGL-LEAF-2", "GE0/1", "10G", "", "", "GE0/3", "ETH", "10G", "SPINE-1"],
        ["DWGL-LEAF-2", "GE0/2", "10G", "", "", "GE0/4", "ETH", "10G", "SPINE-2"],
        ["DWGL-LEAF-1", "MEth0/0", "1G", "", "", "MEth0/0", "ETH", "1G", "DWGL-LEAF-2"],
    ]
    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path) as xw:
        df.to_excel(xw, sheet_name="计算带外管理面端口互联", header=False, index=False)


def _build_resource(path: Path) -> None:
    df = pd.DataFrame(
        [
            {
                "网络平面": "计算带外管理面",
                "VLAN": 100,
                "内部网络设备互连地址段": "10.0.0.0-10.0.0.255",
                "网关位置*": "SPINE",
            }
        ]
    )
    with pd.ExcelWriter(path) as xw:
        df.to_excel(xw, sheet_name="项目信息", index=False)


def main() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        topo = tmp / "建模仿真输出文档007-端口连线表.xlsx"
        res = tmp / "项目信息收集表.xlsx"
        _build_topology(topo)
        _build_resource(res)

        out_l2 = tmp / "out_l2.xlsx"
        out_l3 = tmp / "out_l3.xlsx"
        out_merge = tmp / "out_merge.xlsx"

        df_l2 = run_l2(topo, "计算带外管理互联规划", res, out_l2)
        df_l3 = run_l3(topo, "计算带外管理互联规划", res, out_l3)
        pd.DataFrame(
            [
                {"网络平面": "旧平面", "本端设备": "OLD-LEAF", "标签": "OLD"},
                {"网络平面": "计算带外管理面", "本端设备": "STALE-LEAF", "标签": "STALE"},
            ]
        ).to_excel(out_merge, sheet_name="网络互联规划", index=False)
        df_merge = run_l2(topo, "计算带外管理互联规划", res, out_merge, merge_existing=True)

        assert not df_l2.empty, "L2 结果为空"
        assert not df_l3.empty, "L3 结果为空"
        assert set(["本端ETH-TRUNK", "本端VLAN", "标签"]).issubset(df_l2.columns)
        assert set(["本端接口IP地址", "对端接口IP地址", "本端接口掩码"]).issubset(df_l3.columns)
        assert (df_l2["端口类型"] == "trunk").all()
        assert (df_l3["本端接口IP地址"] == "").all()
        assert (df_l3["对端接口掩码"] == 30).all()
        assert (df_l2["标签"] == "INTER_LINK").all()
        assert (df_l3["标签"] == "INTER_LINK").all()
        assert (df_l2["网络平面"] == "计算带外管理面").all()
        assert (df_l3["网络平面"] == "计算带外管理面").all()
        assert "旧平面" in set(df_merge["网络平面"])
        assert "STALE-LEAF" not in set(df_merge["本端设备"])
        assert _sheetnames(out_l2) == ["网络互联规划"]
        assert _sheetnames(out_l3) == ["网络互联规划"]

        print("L2 rows:", len(df_l2), "L3 rows:", len(df_l3))
        print(df_l2.head().to_string())
        print(df_l3.head().to_string())
        print("smoke test OK")


if __name__ == "__main__":
    main()
