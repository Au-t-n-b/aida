#!/usr/bin/env python3
"""Unit tests for offline device naming (no sample xlsx required)."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

_SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))

from device_list_generator import generate_device_list_df
from lld_device_name_replacer import load_device_name_map, replace_lld_device_names
from ztp_device_name_replacer import load_source_target_mapping, replace_ztp_device_names


class TestDeviceNaming(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _write_location(self, path: Path) -> None:
        df = pd.DataFrame(
            {
                "设备名称": ["SW-A", "空挡板", "SW-B"],
                "所属机房": ["R1", "R1", "R1"],
                "所属机柜": ["C1", "C1", "C1"],
                "安装起始U位": ["1", "2", "3"],
            }
        )
        df.to_excel(path, sheet_name="设备位置信息", index=False)

    def test_generate_list(self) -> None:
        loc = self.root / "004设备位置.xlsx"
        self._write_location(loc)
        out = generate_device_list_df(loc)
        self.assertEqual(len(out), 2)
        self.assertListEqual(list(out.columns), [
            "编号", "设备名称", "所属机房", "所属机柜", "安装起始U位", "客户定义设备名称",
        ])
        self.assertEqual(int(out.iloc[0]["编号"]), 1)

    def test_replace_lld_dry_run(self) -> None:
        dev = self.root / "设备清单表.xlsx"
        pd.DataFrame(
            {
                "设备名称": ["SW-A"],
                "客户定义设备名称": ["CUSTOM-A"],
            }
        ).to_excel(dev, index=False)

        lld = self.root / "LLD设计.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "SW-A"
        wb.save(lld)

        out, stats = replace_lld_device_names(
            device_list_path=dev,
            lld_path=lld,
            out_dir=self.root / "output",
            dry_run=True,
            restrict_to_cwd=False,
        )
        self.assertIsNone(out)
        self.assertEqual(stats.cells_replaced, 1)

    def test_replace_ztp_plane_filter(self) -> None:
        mapping = self.root / "devicename-mapping.csv"
        pd.DataFrame({"source": ["S1", "S2"], "target": ["T1", "T2"]}).to_csv(
            mapping, index=False
        )
        ztp = self.root / "ZTP_LLD.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "网络IP规划"
        ws.cell(1, 1, "交换机名称")
        ws.cell(1, 2, "L1/L2平面")
        ws.cell(2, 1, "S1")
        ws.cell(2, 2, 1)
        ws.cell(3, 1, "S2")
        ws.cell(3, 2, 2)
        wb.save(ztp)

        m = load_source_target_mapping(mapping)
        self.assertEqual(len(m), 2)

        _, stats_l1 = replace_ztp_device_names(
            ztp_lld_path=ztp,
            mapping_path=mapping,
            plane=1,
            out_dir=self.root / "output",
            dry_run=True,
            restrict_to_cwd=False,
        )
        self.assertEqual(stats_l1.replaced, 1)
        self.assertEqual(stats_l1.filtered_rows, 1)

        _, stats_all = replace_ztp_device_names(
            ztp_lld_path=ztp,
            mapping_path=mapping,
            plane=None,
            out_dir=self.root / "output",
            dry_run=True,
            restrict_to_cwd=False,
        )
        self.assertEqual(stats_all.replaced, 2)


if __name__ == "__main__":
    unittest.main()
