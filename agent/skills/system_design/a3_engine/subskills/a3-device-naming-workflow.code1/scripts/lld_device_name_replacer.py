#!/usr/bin/env python3
"""Replace device names in full LLD workbook using 设备清单 mapping."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional, Tuple

import openpyxl
import pandas as pd

from naming_path_utils import new_run_dir, require_under_cwd, resolve_local_only


@dataclass
class LldReplaceStats:
    cells_scanned: int
    cells_replaced: int
    mapping_size: int


def load_device_name_map(
    device_list_path: Path,
    *,
    old_col: str = "设备名称",
    new_col: str = "客户定义设备名称",
    sheet_name: Optional[object] = "设备信息",
) -> Dict[str, str]:
    try:
        df = pd.read_excel(device_list_path, sheet_name=sheet_name, header=0, dtype=str)
    except ValueError:
        df = pd.read_excel(device_list_path, sheet_name=0, header=0, dtype=str)
    if old_col not in df.columns or new_col not in df.columns:
        raise ValueError(
            f"设备清单缺少列 {old_col!r}/{new_col!r}，实际: {list(df.columns)}"
        )
    df[old_col] = df[old_col].astype(str).str.strip()
    df[new_col] = df[new_col].astype(str).str.strip()
    df = df[df[new_col].notna() & (df[new_col] != "") & (df[new_col].str.lower() != "nan")]
    if df.empty:
        return {}
    return dict(zip(df[old_col].tolist(), df[new_col].tolist()))


def replace_lld_device_names(
    *,
    device_list_path: Path,
    lld_path: Path,
    out_dir: Path = Path("output"),
    dry_run: bool = False,
    old_col: str = "设备名称",
    new_col: str = "客户定义设备名称",
    restrict_to_cwd: bool = True,
) -> Tuple[Optional[Path], LldReplaceStats]:
    if restrict_to_cwd:
        device_list_path = require_under_cwd(device_list_path, "device-list")
        lld_path = require_under_cwd(lld_path, "lld")
    else:
        device_list_path = resolve_local_only(device_list_path)
        lld_path = resolve_local_only(lld_path)

    device_name_map = load_device_name_map(
        device_list_path, old_col=old_col, new_col=new_col
    )

    wb = openpyxl.load_workbook(lld_path, read_only=False, keep_vba=False)
    cells_scanned = 0
    cells_replaced = 0

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells_scanned += 1
                cell_value = str(cell.value).strip()
                if cell_value in device_name_map:
                    new_val = device_name_map[cell_value]
                    if pd.isna(new_val) or new_val is None:
                        continue
                    cells_replaced += 1
                    if not dry_run:
                        cell.value = new_val

    stats = LldReplaceStats(
        cells_scanned=cells_scanned,
        cells_replaced=cells_replaced,
        mapping_size=len(device_name_map),
    )

    if dry_run:
        return None, stats

    run_dir = new_run_dir(out_dir, restrict_to_cwd=restrict_to_cwd)
    out_path = run_dir / f"{lld_path.stem}_replaced{lld_path.suffix}"
    wb.save(out_path)
    return out_path, stats
