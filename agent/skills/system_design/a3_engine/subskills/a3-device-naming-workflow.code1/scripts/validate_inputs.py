#!/usr/bin/env python3
"""Validate inputs for offline device naming workflow. Exit 0/1/2."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List, Optional

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

try:
    import pandas as pd
except ImportError:
    print("ERROR: pandas required. pip install pandas openpyxl", file=sys.stderr)
    sys.exit(2)

from lld_device_name_replacer import load_device_name_map
from naming_path_utils import (
    autodetect_device_list_in_cwd,
    autodetect_lld_in_cwd,
    autodetect_location_in_cwd,
    autodetect_mapping_in_cwd,
    autodetect_ztp_lld_in_cwd,
    require_under_cwd,
)
from ztp_device_name_replacer import ZTP_SHEET, SWITCH_COL, PLANE_COL, load_source_target_mapping

MODES = [
    "generate-list",
    "replace-lld",
    "replace-ztp",
    "replace-ztp-l1",
    "replace-ztp-l2",
]


def validate_generate_list(path: Path, sheet: str) -> List[str]:
    errors: List[str] = []
    try:
        df = pd.read_excel(path, sheet_name=sheet, header=0)
    except Exception as e:
        return [f"location: cannot read: {e}"]
    for col in ["设备名称", "所属机房", "所属机柜", "安装起始U位"]:
        if col not in df.columns:
            errors.append(f"location: missing column {col!r}")
    if not errors:
        n = len(df[df["设备名称"].astype(str) != "空挡板"]) if "设备名称" in df.columns else 0
        print(f"location: OK sheet={sheet!r} rows~={n}")
    return errors


def validate_replace_lld(device_list: Path, lld: Path) -> List[str]:
    errors: List[str] = []
    try:
        m = load_device_name_map(device_list)
        print(f"device-list: OK mapping_entries={len(m)}")
    except Exception as e:
        errors.append(f"device-list: {e}")
    try:
        import openpyxl

        wb = openpyxl.load_workbook(lld, read_only=True)
        print(f"lld: OK sheets={len(wb.sheetnames)}")
        wb.close()
    except Exception as e:
        errors.append(f"lld: cannot open: {e}")
    return errors


def validate_replace_ztp(ztp: Path, mapping: Path, plane: Optional[int]) -> List[str]:
    errors: List[str] = []
    try:
        m = load_source_target_mapping(mapping)
        print(f"mapping: OK entries={len(m)}")
    except Exception as e:
        errors.append(f"mapping: {e}")
    try:
        import openpyxl

        wb = openpyxl.load_workbook(ztp, read_only=True)
        if ZTP_SHEET not in wb.sheetnames:
            errors.append(f"ztp-lld: missing sheet {ZTP_SHEET!r}")
        else:
            ws = wb[ZTP_SHEET]
            cols = [
                str(ws.cell(1, c).value).strip()
                for c in range(1, ws.max_column + 1)
                if ws.cell(1, c).value is not None
            ]
            if SWITCH_COL not in cols:
                errors.append(f"ztp-lld: missing column {SWITCH_COL!r}")
            if plane is not None and PLANE_COL not in cols:
                errors.append(f"ztp-lld: plane={plane} needs column {PLANE_COL!r}")
            print(f"ztp-lld: OK sheet={ZTP_SHEET!r} header_cols={len(cols)}")
        wb.close()
    except Exception as e:
        errors.append(f"ztp-lld: {e}")
    return errors


def main() -> int:
    p = argparse.ArgumentParser(description="Validate device naming inputs")
    p.add_argument("--mode", required=True, choices=MODES)
    p.add_argument("--location", type=Path, default=None)
    p.add_argument("--sheet", default="设备位置信息")
    p.add_argument("--device-list", type=Path, default=None)
    p.add_argument("--lld", type=Path, default=None)
    p.add_argument("--ztp-lld", type=Path, default=None)
    p.add_argument("--mapping", type=Path, default=None)
    args = p.parse_args()

    print(f"cwd: {Path.cwd().resolve()}")
    all_err: List[str] = []

    try:
        if args.mode == "generate-list":
            path = args.location or autodetect_location_in_cwd()
            path = require_under_cwd(path, "location")
            all_err.extend(validate_generate_list(path, args.sheet))

        elif args.mode == "replace-lld":
            dl = args.device_list or autodetect_device_list_in_cwd()
            lld = args.lld or autodetect_lld_in_cwd()
            dl = require_under_cwd(dl, "device-list")
            lld = require_under_cwd(lld, "lld")
            all_err.extend(validate_replace_lld(dl, lld))

        else:
            plane = None
            if args.mode == "replace-ztp-l1":
                plane = 1
            elif args.mode == "replace-ztp-l2":
                plane = 2
            ztp = args.ztp_lld or autodetect_ztp_lld_in_cwd()
            mp = args.mapping or autodetect_mapping_in_cwd()
            ztp = require_under_cwd(ztp, "ztp-lld")
            mp = require_under_cwd(mp, "mapping")
            all_err.extend(validate_replace_ztp(ztp, mp, plane))

    except SystemExit as e:
        print(str(e), file=sys.stderr)
        return 2

    if all_err:
        print("Validation failed:", file=sys.stderr)
        for e in all_err:
            print(f"  - {e}", file=sys.stderr)
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
