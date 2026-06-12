#!/usr/bin/env python3
"""ZTP_LLD 网络IP规划 sheet 交换机名称列替换 (source/target mapping)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional, Set, Tuple, Union

import openpyxl
import pandas as pd

from naming_path_utils import new_run_dir, require_under_cwd, resolve_local_only, ztp_output_suffix

ZTP_SHEET = "网络IP规划"
SWITCH_COL = "交换机名称"
PLANE_COL = "L1/L2平面"


@dataclass
class ZtpReplaceStats:
    plane_label: str
    mapping_size: int
    hit_sources: int
    replaced: int
    missed_sources: int
    filtered_rows: int


def parse_ztp_l12_plane_value(raw) -> Optional[int]:
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        v = int(raw)
        return v if v in (1, 2) else None
    s = str(raw).strip()
    if not s or s.lower() in ("none", "nan"):
        return None
    try:
        v = int(float(s))
    except (TypeError, ValueError):
        return None
    return v if v in (1, 2) else None


def load_source_target_mapping(mapping_path: Path) -> Dict[str, str]:
    name = mapping_path.name.lower()
    if name.endswith(".csv"):
        df = pd.read_csv(mapping_path)
    else:
        df = pd.read_excel(mapping_path)

    required_cols = {"source", "target"}
    if not required_cols.issubset(set(df.columns)):
        raise ValueError(
            f"映射文件缺少必要列 {required_cols}，实际列: {list(df.columns)}"
        )

    df = df[["source", "target"]].copy()
    df["source"] = df["source"].astype(str).str.strip()
    df["target"] = df["target"].astype(str).str.strip()

    dup_sources = df[df.duplicated(subset=["source"], keep=False)]["source"].tolist()
    if dup_sources:
        raise ValueError(f"映射文件 source 存在重复值: {dup_sources[:20]}")

    empty_targets = df[df["target"].fillna("").astype(str).str.strip() == ""]
    if not empty_targets.empty:
        raise ValueError(f"映射文件 target 存在空值，记录数: {len(empty_targets)}")

    mapping = dict(zip(df["source"].tolist(), df["target"].tolist()))
    if not mapping:
        raise ValueError("映射文件为空，未读取到任何 source/target 记录")
    return mapping


def _find_col_idx(ws, header_row: int, col_name: str) -> Optional[int]:
    for col_idx in range(1, ws.max_column + 1):
        if str(ws.cell(row=header_row, column=col_idx).value).strip() == col_name:
            return col_idx
    return None


def replace_ztp_device_names(
    *,
    ztp_lld_path: Path,
    mapping_path: Path,
    plane: Optional[int] = None,
    out_dir: Path = Path("output"),
    dry_run: bool = False,
    restrict_to_cwd: bool = True,
) -> Tuple[Optional[Path], ZtpReplaceStats]:
    if restrict_to_cwd:
        ztp_lld_path = require_under_cwd(ztp_lld_path, "ztp-lld")
        mapping_path = require_under_cwd(mapping_path, "mapping")
    else:
        ztp_lld_path = resolve_local_only(ztp_lld_path)
        mapping_path = resolve_local_only(mapping_path)

    mapping = load_source_target_mapping(mapping_path)
    wb = openpyxl.load_workbook(ztp_lld_path, read_only=False, keep_vba=False)

    if ZTP_SHEET not in wb.sheetnames:
        raise ValueError(
            f"ZTP设计文件缺少 sheet「{ZTP_SHEET}」，实际: {wb.sheetnames}"
        )

    ws = wb[ZTP_SHEET]
    header_row = 1
    target_col = _find_col_idx(ws, header_row, SWITCH_COL)
    if target_col is None:
        raise ValueError(f"ZTP设计文件未找到列「{SWITCH_COL}」")

    plane_col = None
    if plane is not None:
        plane_col = _find_col_idx(ws, header_row, PLANE_COL)
        if plane_col is None:
            raise ValueError(
                f"已指定平面过滤 plane={plane}，但未找到列「{PLANE_COL}」"
            )

    replaced = 0
    hit_sources: Set[str] = set()
    filtered_rows = 0

    for row_idx in range(2, ws.max_row + 1):
        if plane_col is not None:
            raw_plane = ws.cell(row=row_idx, column=plane_col).value
            row_plane = parse_ztp_l12_plane_value(raw_plane)
            if row_plane is None:
                continue
            if row_plane != plane:
                continue
            filtered_rows += 1

        cell = ws.cell(row=row_idx, column=target_col)
        if cell.value is None:
            continue
        original = str(cell.value).strip()
        if original in mapping:
            hit_sources.add(original)
            new_val = mapping[original]
            if new_val != original:
                replaced += 1
            if not dry_run:
                cell.value = new_val

    missed_sources = len(mapping) - len(hit_sources)
    plane_label = "ALL" if plane is None else str(plane)
    stats = ZtpReplaceStats(
        plane_label=plane_label,
        mapping_size=len(mapping),
        hit_sources=len(hit_sources),
        replaced=replaced,
        missed_sources=missed_sources,
        filtered_rows=filtered_rows if plane is not None else 0,
    )

    if dry_run:
        return None, stats

    run_dir = new_run_dir(out_dir, restrict_to_cwd=restrict_to_cwd)
    suffix = ztp_output_suffix(plane)
    out_path = run_dir / f"{ztp_lld_path.stem}{suffix}{ztp_lld_path.suffix}"
    wb.save(out_path)
    return out_path, stats
