"""
data_append · 勘测条目追加

意图: survey_work 专属

流程:
  1. HITL ChoiceCard 询问用户是否追加勘测条目（可选上传额外追加表 / 跳过）
  2. 若追加：
     a. 若用户上传了「追加工勘项表.xlsx」→ 读取并直接 append 到勘测结果表
     b. 否则：加载底表 → 按 generation_cooling 过滤分类=数据 → append_data_items()
  3. 返回追加数量 metrics
"""
from __future__ import annotations

import json
import os

from ...base import BaseStep, SkillContext, SkillState, StepResult, Emit, CheckResult
from ._intent_guard import should_skip

# 追加工勘项表的固定存储路径（相对 work_root）
EXTRA_ITEMS_REL = "ProjectData/Input/追加工勘项表.xlsx"

_APPEND_OPTIONS = [
    {
        "label": "追加勘测条目",
        "value": "append",
        "description": "可上传「追加工勘项表.xlsx」指定自定义条目，或从底表自动追加数据类条目",
    },
    {
        "label": "跳过，不追加",
        "value": "skip",
        "description": "保持当前勘测表内容不变，直接进入下一步",
    },
]

_HITL_CHOICE = {
    "id": "data_append_choice",
    "label": "是否追加勘测条目？",
    "options": _APPEND_OPTIONS,
}

# 阶段二：用户已选「追加」后，提供上传入口 + 确认（避免一点「追加」就直接跑完本步）
_HITL_APPEND_CONFIRM = {
    "id": "data_append_upload",
    "label": "追加工勘项表（可选）",
    "options": [
        {
            "label": "从底表追加数据类条目",
            "value": "confirm_base",
            "description": "不上传自定义表，自动从底表追加数据类勘测条目",
        },
    ],
    "upload_hint": "或上传「追加工勘项表.xlsx」指定自定义勘测条目",
    "upload_accept": ".xlsx,.xls",
}


def _get_survey_table(ctx: SkillContext) -> str | None:
    info_path = ctx.runtime_dir / "project_info.json"
    if info_path.exists():
        try:
            path = json.loads(info_path.read_text(encoding="utf-8")).get("survey_table_path", "")
            if path and os.path.exists(path):
                return path
        except Exception:
            pass
    tables = sorted(ctx.output_dir.glob("*全量勘测结果表*.xlsx")) if ctx.output_dir.exists() else []
    return str(tables[0]) if tables else None


def _get_generation_cooling(ctx: SkillContext) -> str:
    gc = ctx.project.get("generation_cooling", "")
    if gc:
        return gc
    info_path = ctx.runtime_dir / "project_info.json"
    if info_path.exists():
        try:
            return json.loads(info_path.read_text(encoding="utf-8")).get("generation_cooling", "")
        except Exception:
            pass
    return ""


def _already_appended(ctx: SkillContext) -> bool:
    """幂等标记：本 run 是否已追加过数据条目（防 full_restart 重放重复追加）。"""
    info_path = ctx.runtime_dir / "project_info.json"
    if info_path.exists():
        try:
            return bool(json.loads(info_path.read_text(encoding="utf-8")).get("data_append_done"))
        except Exception:
            pass
    return False


def _mark_appended(ctx: SkillContext, count: int) -> None:
    info_path = ctx.runtime_dir / "project_info.json"
    try:
        existing = json.loads(info_path.read_text(encoding="utf-8")) if info_path.exists() else {}
    except Exception:
        existing = {}
    existing["data_append_done"] = True
    existing["data_append_count"] = count
    ctx.runtime_dir.mkdir(parents=True, exist_ok=True)
    info_path.write_text(json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8")


class DataAppendStep(BaseStep):
    key = "data_append"
    name = "数据条目追加"
    artifacts_pattern = []

    def check_inputs(self, ctx: SkillContext) -> CheckResult:
        if should_skip(self.key, ctx.project):
            return {"ok": True, "missing": []}

        choice = ctx.project.get("data_append_choice")

        # 阶段一：ChoiceCard 追加 / 跳过
        if not choice:
            return {
                "ok": False,
                "missing": [],
                "need_inputs": [_HITL_CHOICE],
            }

        if choice == "skip":
            return {"ok": True, "missing": []}

        # 阶段二：已选追加 → 等待上传（可选）并确认，再执行 run
        if not ctx.project.get("data_append_confirmed"):
            return {
                "ok": False,
                "missing": [],
                "need_inputs": [_HITL_APPEND_CONFIRM],
            }

        return {"ok": True, "missing": []}

    def run(self, ctx: SkillContext, state: SkillState, emit: Emit) -> StepResult:
        if should_skip(self.key, ctx.project):
            return {}

        choice = ctx.project.get("data_append_choice", "skip")

        if choice == "skip":
            emit("[data_append] 跳过勘测条目追加")
            return {"metrics": {"data_append_skipped": True, "data_append_count": 0}}

        # 幂等：full_restart 重放本步时若已追加过则跳过，否则数据行会累积翻倍
        if _already_appended(ctx):
            emit("[data_append] ✓ 勘测条目已追加过（幂等跳过，避免重复行）")
            return {"metrics": {"data_append_idempotent_skip": True}}

        # choice == "append" —— 先检查有无用户上传的自定义追加表
        survey_table_path = _get_survey_table(ctx)
        if not survey_table_path:
            emit("[data_append] ⚠ 全量勘测结果表不存在，跳过追加")
            return {"metrics": {"data_append_skipped": True, "data_append_count": 0}}

        extra_xlsx = ctx.work_root / EXTRA_ITEMS_REL
        if extra_xlsx.is_file():
            count, total_rows = _append_from_custom_xlsx(extra_xlsx, survey_table_path, emit)
            _mark_appended(ctx, count)
            return {
                "metrics": {
                    "data_append_count": count,
                    "data_append_source": "custom_xlsx",
                    "total_rows_after_append": total_rows,
                }
            }

        # 无自定义上传表 → 从底表过滤数据类条目追加
        gen_cooling = _get_generation_cooling(ctx)
        if not gen_cooling:
            emit("[data_append] ⚠ generation_cooling 未知，跳过追加")
            return {"metrics": {"data_append_skipped": True, "data_append_count": 0}}

        from ..services.table_filter import (
            load_base_table,
            get_available_data_keywords,
            filter_items,
        )
        from ..services.survey_table_builder import append_data_items
        from ..path_config import get_base_table_path

        base_table_path = get_base_table_path()
        if not os.path.exists(base_table_path):
            emit("[data_append] ⚠ 底表文件不存在，跳过追加")
            return {"metrics": {"data_append_skipped": True, "data_append_count": 0}}

        emit(f"[data_append] 加载底表: {os.path.basename(base_table_path)}")
        items = load_base_table(base_table_path)

        keywords = get_available_data_keywords(items, gen_cooling)
        if not keywords:
            emit(f"[data_append] {gen_cooling} 下无数据类条目，跳过")
            return {"metrics": {"data_append_count": 0}}

        emit(f"[data_append] 数据类关键词: {keywords}")

        try:
            data_items = filter_items(
                items,
                generation_cooling=gen_cooling,
                category="数据",
                data_keywords=keywords,
            )
        except Exception as e:
            emit(f"[data_append] ⚠ 过滤失败: {e}，跳过")
            return {"metrics": {"data_append_skipped": True, "data_append_count": 0}}

        total_rows = append_data_items(survey_table_path, data_items)
        count = len(data_items)
        _mark_appended(ctx, count)
        emit(f"[data_append] ✓ 追加 {count} 条数据类勘测条目（表格总行数: {total_rows}）")

        return {
            "metrics": {
                "data_append_count": count,
                "data_append_source": "base_table",
                "total_rows_after_append": total_rows,
            }
        }


def _append_from_custom_xlsx(
    extra_xlsx,
    survey_table_path: str,
    emit,
) -> tuple[int, int]:
    """从用户上传的追加工勘项表读取条目并追加到勘测结果表。"""
    try:
        import openpyxl as _openpyxl
    except ImportError:
        emit("[data_append] ⚠ openpyxl 未安装，跳过自定义追加")
        return 0, 0

    emit(f"[data_append] 读取自定义追加工勘项表: {extra_xlsx.name}")
    wb = _openpyxl.load_workbook(str(extra_xlsx), read_only=True, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        wb.close()
        emit("[data_append] ⚠ 追加工勘项表为空，跳过")
        return 0, 0

    headers = [str(h).strip() if h else "" for h in header_row]
    col = {name: i for i, name in enumerate(headers) if name}

    required = {"细分场景", "勘测要素", "项目", "检查内容", "勘测方法"}
    missing_cols = required - set(col)
    if missing_cols:
        wb.close()
        emit(f"[data_append] ⚠ 追加表缺少必要列 {missing_cols}，跳过")
        return 0, 0

    custom_items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None and row[1] is None:
            continue

        def _g(key: str) -> str:
            idx = col.get(key)
            if idx is None or idx >= len(row):
                return ""
            return str(row[idx]).strip() if row[idx] is not None else ""

        item = {
            "细分场景": _g("细分场景"),
            "勘测要素": _g("勘测要素"),
            "项目": _g("项目"),
            "检查内容": _g("检查内容"),
            "勘测方法": _g("勘测方法"),
            "备注": _g("备注"),
        }
        if item["细分场景"] or item["项目"]:
            custom_items.append(item)

    wb.close()

    if not custom_items:
        emit("[data_append] ⚠ 追加工勘项表无有效数据行，跳过")
        return 0, 0

    from ..services.survey_table_builder import append_data_items as _append
    total_rows = _append(survey_table_path, custom_items)  # type: ignore[arg-type]
    emit(f"[data_append] ✓ 从自定义表追加 {len(custom_items)} 条勘测条目（表格总行数: {total_rows}）")
    return len(custom_items), total_rows
