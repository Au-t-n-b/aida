"""
assess · AI 五值评估

意图: survey_work / report_gen

流程:
  1. 找到全量勘测结果表
  2. evaluate_all() — LLM 批量评估每条勘测项（检查内容 + 最新检查结果 → 五值结论）
  3. get_assessment_statistics() — 统计各五值数量
  4. 返回 metrics（五值统计 + 评估总数）
"""
from __future__ import annotations

import json
import os
from typing import Any

from ...base import BaseStep, SkillContext, SkillState, StepResult, Emit, CheckResult
from ._intent_guard import should_skip


ASSESSMENT_LABELS = ("满足", "不满足", "不涉及", "未勘测", "无法识别")


def _get_survey_table(ctx: SkillContext) -> str | None:
    info_path = ctx.runtime_dir / "project_info.json"
    if info_path.exists():
        try:
            path = json.loads(info_path.read_text(encoding="utf-8")).get("survey_table_path", "")
            if path and os.path.exists(path):
                return path
        except Exception:
            pass
    tables = sorted(ctx.output_dir.glob("*全量勘测结果表*.xlsx")) if ctx.output_dir.exists() else []
    return str(tables[0]) if tables else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _pick_header(headers: list[str], *names: str) -> str | None:
    for name in names:
        if name in headers:
            return name
    return None


def _build_assessment_detail_groups(survey_table: str) -> dict[str, list[dict[str, str]]]:
    """按 AI 五值结果分组，供 SDUI NumberCard 点击查看明细。"""
    from openpyxl import load_workbook

    wb = load_workbook(survey_table, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {label: [] for label in ASSESSMENT_LABELS}

    headers = [_cell_text(v) for v in rows[0]]
    col = {name: i for i, name in enumerate(headers) if name}
    item_key = _pick_header(headers, "检查内容", "勘测项", "勘测要素", "项目")
    result_key = _pick_header(headers, "最新检查结果", "检查结果", "勘测结果")
    ai_key = _pick_header(headers, "AI评估结果", "评估结果")
    source_key = _pick_header(headers, "最新结果来源", "结果来源", "来源", "数据来源")
    time_key = _pick_header(headers, "最新勘测时间", "勘测时间", "上传时间", "更新时间")
    risk_key = next((h for h in headers if "风险" in h), None)

    groups: dict[str, list[dict[str, str]]] = {label: [] for label in ASSESSMENT_LABELS}
    if not ai_key:
        return groups

    def get(row: tuple[Any, ...], key: str | None) -> str:
        if not key:
            return ""
        idx = col.get(key)
        if idx is None or idx >= len(row):
            return ""
        return _cell_text(row[idx])

    for row in rows[1:]:
        label = get(row, ai_key) or "未勘测"
        if label not in groups:
            label = "无法识别"
        result = get(row, result_key)
        source = get(row, source_key)
        if not source:
            source = result_key if result else "未填写"
        groups[label].append({
            "item": get(row, item_key),
            "result": result,
            "source": source,
            "time": get(row, time_key),
            "risk": get(row, risk_key),
        })
    return groups


class AssessStep(BaseStep):
    key = "assess"
    name = "AI 五值评估"
    artifacts_pattern = ["ProjectData/Output/*全量勘测结果表*.xlsx"]

    def check_inputs(self, ctx: SkillContext) -> CheckResult:
        if should_skip(self.key, ctx.project):
            return {"ok": True, "missing": []}
        if _get_survey_table(ctx) is None:
            return {"ok": False, "missing": ["ProjectData/Output/*全量勘测结果表*.xlsx"]}
        return {"ok": True, "missing": []}

    def run(self, ctx: SkillContext, state: SkillState, emit: Emit) -> StepResult:
        if should_skip(self.key, ctx.project):
            return {}

        from ..services.assessment_engine import evaluate_all, get_assessment_statistics
        from ..services._llm_adapter import make_llm_adapter

        survey_table = _get_survey_table(ctx)
        if not survey_table:
            raise RuntimeError("assess: 全量勘测结果表不存在")

        emit(f"[assess] 开始 AI 五值评估: {os.path.basename(survey_table)}")
        llm = make_llm_adapter(ctx, step_key=self.key)

        results = evaluate_all(survey_table, llm)
        total = len(results)
        emit(f"[assess] 评估完成: {total} 条")

        stats = get_assessment_statistics(survey_table)
        detail_groups = _build_assessment_detail_groups(survey_table)
        emit(
            f"[assess] 统计: "
            f"满足={stats.get('满足', 0)} / 不满足={stats.get('不满足', 0)} / "
            f"不涉及={stats.get('不涉及', 0)} / 未勘测={stats.get('未勘测', 0)} / "
            f"无法识别={stats.get('无法识别', 0)}"
        )

        return {
            "metrics": {
                "assess_total": total,
                **{f"assess_{k}": v for k, v in stats.items()},
                "assess_detail_groups": detail_groups,
            },
            "artifacts": [ctx.rel(survey_table)],
        }
