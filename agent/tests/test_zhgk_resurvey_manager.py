from pathlib import Path

from openpyxl import Workbook, load_workbook

from agent.skills.zhgk.services.resurvey_manager import write_survey_results


def test_write_survey_results_records_source_and_survey_time(tmp_path: Path):
    path = tmp_path / "全量勘测结果表.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "检查内容", "最新检查结果", "AI评估结果"])
    ws.append([1, "冷通道封闭完整", "", ""])
    ws.append([2, "温湿度监控可用", "", ""])
    wb.save(path)

    stat = write_survey_results(
        str(path),
        {1: "已封闭"},
        1,
        source_label="第1次手动上传",
        survey_time="2026-06-14 17:35:00",
    )

    assert stat["matched"] == 1

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    latest_source_col = headers.index("最新结果来源") + 1
    latest_time_col = headers.index("最新勘测时间") + 1
    round_source_col = headers.index("第一轮结果来源") + 1
    round_time_col = headers.index("第一轮勘测时间") + 1

    assert ws.cell(2, latest_source_col).value == "第1次手动上传"
    assert ws.cell(2, latest_time_col).value == "2026-06-14 17:35:00"
    assert ws.cell(2, round_source_col).value == "第1次手动上传"
    assert ws.cell(2, round_time_col).value == "2026-06-14 17:35:00"
    assert ws.cell(3, latest_source_col).value is None
    wb.close()
