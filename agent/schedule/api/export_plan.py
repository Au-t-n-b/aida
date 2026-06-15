from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from agent.schedule.store import PlanVersionNotFound, PlanVersionStore, StoredPlanVersion
from agent.schedule.contracts.inputs import Activity, InputBundle
from agent.schedule.contracts.outputs import PlanResult, ScheduledActivity

SCHEDULE_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = SCHEDULE_ROOT / "project-data" / "12_输出文件"
DELIVERY_PLAN_FILE = OUTPUT_DIR / "交付计划表.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NAME = "sheet0"

DELIVERY_PLAN_HEADERS = (
    "ID",
    "TODO_ID",
    "SERIAL_NUMBER",
    "ACTIVITY_ID",
    "ACTIVITY_NAME",
    "TASK_RANK",
    "PARENT_ID",
    "PROJECT_ID",
    "INSTRUCTION",
    "FORMAT_INSTRUCTION",
    "SRC_AGENT",
    "TARGET_AGENT",
    "START_DATE",
    "END_DATE",
    "ACTUAL_START_DATE",
    "ACTUAL_END_DATE",
    "STATUS",
    "PROCESS",
    "PRINCIPAL",
    "PRINCIPAL_COMPANY",
    "SCENARIO",
    "CREATED_BY",
    "CREATION_DATE",
    "LAST_UPDATE_BY",
    "LAST_UPDATE_DATE",
    "group_id",
    "RAW_EQUIPMENT_LIST",
    "TASK_COMMENT",
    "MANAGEMENT_UNIT",
    "OWNER",
    "MANUAL_PROCESS",
    "REAL_MANAGEMENT_UNIT",
)


class ExportPlanQueryError(ValueError):
    """Raised when query params cannot identify a formal plan version."""


def resolve_plan_version_for_export(
    store: PlanVersionStore,
    *,
    plan_id: str | None = None,
    version: int | None = None,
) -> StoredPlanVersion:
    normalized_plan_id = plan_id.strip() if plan_id else None
    if normalized_plan_id and version is not None:
        return store.get_plan_version(normalized_plan_id, version)
    if version is not None:
        raise ExportPlanQueryError("version 必须与 plan_id 一起传入。")
    if normalized_plan_id:
        latest_version = _latest_version_for_plan(store.db_path, normalized_plan_id)
        return store.get_plan_version(normalized_plan_id, latest_version)

    latest_plan_id, latest_version = _latest_plan_key(store.db_path)
    return store.get_plan_version(latest_plan_id, latest_version)


def write_delivery_plan_xlsx(
    plan: PlanResult,
    inputs: InputBundle,
    output_path: Path | None = None,
) -> Path:
    target_path = output_path or DELIVERY_PLAN_FILE
    target_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(list(DELIVERY_PLAN_HEADERS))

    template_by_id = {activity.activity_id: activity for activity in inputs.activities}
    for serial_number, activity in enumerate(plan.activities, start=1):
        row = _activity_row(plan, inputs, activity, template_by_id.get(activity.activity_id), serial_number)
        sheet.append([row.get(header, "") for header in DELIVERY_PLAN_HEADERS])

    _format_sheet(sheet)
    workbook.save(target_path)
    return target_path


def _latest_version_for_plan(db_path: Path, plan_id: str) -> int:
    with sqlite3.connect(db_path) as conn:
        row = conn.execute(
            "SELECT MAX(version) AS version FROM plan_versions WHERE plan_id = ?",
            (plan_id,),
        ).fetchone()
    if row is None or row[0] is None:
        raise PlanVersionNotFound(f"{plan_id}@latest")
    return int(row[0])


def _latest_plan_key(db_path: Path) -> tuple[str, int]:
    with sqlite3.connect(db_path) as conn:
        row = conn.execute(
            """
            SELECT plan_id, version
            FROM plan_versions
            ORDER BY created_at DESC, version DESC, plan_id DESC
            LIMIT 1
            """
        ).fetchone()
    if row is None:
        raise PlanVersionNotFound("latest")
    return str(row[0]), int(row[1])


def _activity_row(
    plan: PlanResult,
    inputs: InputBundle,
    activity: ScheduledActivity,
    template: Activity | None,
    serial_number: int,
) -> dict[str, Any]:
    project = inputs.project
    management_unit = activity.scope_ref.ref_id or project.project_id
    principal = template.responsibility if template and template.responsibility else activity.team_id or ""
    owner = activity.team_id or principal
    return {
        "ID": f"{plan.plan_id}-v{plan.version}-{activity.instance_id}",
        "SERIAL_NUMBER": str(serial_number),
        "ACTIVITY_ID": activity.activity_id,
        "ACTIVITY_NAME": activity.activity_name,
        "TASK_RANK": activity.scope_ref.scope,
        "PROJECT_ID": project.project_id,
        "START_DATE": activity.start_date,
        "END_DATE": activity.end_date,
        "STATUS": "已排期",
        "PRINCIPAL": principal,
        "SCENARIO": project.scene or "",
        "group_id": management_unit,
        "TASK_COMMENT": template.note if template and template.note else "",
        "MANAGEMENT_UNIT": management_unit,
        "OWNER": owner,
    }


def _format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in sheet.iter_rows(min_row=2, min_col=13, max_col=16):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    widths = {
        "A": 36,
        "C": 12,
        "D": 14,
        "E": 28,
        "F": 12,
        "H": 18,
        "M": 14,
        "N": 14,
        "Q": 12,
        "S": 18,
        "U": 18,
        "Z": 18,
        "AB": 28,
        "AC": 20,
        "AD": 18,
    }
    for index, _header in enumerate(DELIVERY_PLAN_HEADERS, start=1):
        column = get_column_letter(index)
        sheet.column_dimensions[column].width = widths.get(column, 16)
    sheet.freeze_panes = "A2"
