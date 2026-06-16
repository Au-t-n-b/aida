from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, File, Form, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from agent.schedule.contracts.api import API_PREFIX, ChangeSet, ConflictDetail, ErrorResponse, ParseChangesResponse
from agent.schedule.contracts.inputs import ArrivalItem, Batch, Room
from agent.schedule.importer import DataImportError, load_input_bundle
from agent.schedule.settings import get_as_of_date

router = APIRouter(prefix=API_PREFIX, tags=["schedule"])

SCHEDULE_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = SCHEDULE_ROOT / "project-data" / "10_变更表模板"
TEMPLATE_FILE = TEMPLATE_DIR / "调整表模板.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ROOM_SHEET = "机房ready"
ARRIVAL_SHEET = "到货"
BATCH_SHEET = "批次目标"

ROOM_COLUMNS = ("机房id", "可布线", "可装设备", "可通液")
ARRIVAL_COLUMNS = ("管理单元", "设备类型", "型号", "数量", "单位", "到货日期")
BATCH_COLUMNS = ("batch_id", "上电目标", "上线目标")

_INVALID_DATE = object()


@router.get("/change-template")
def download_change_template():
    if not TEMPLATE_FILE.exists():
        return _import_error(["调整表模板文件未生成，请检查 02_项目数据/10_变更表模板。"])
    try:
        content = _prefilled_change_template()
    except (DataImportError, OSError, KeyError, ValueError) as exc:
        return _import_error([f"调整表模板预填失败：{exc}"])
    return StreamingResponse(
        BytesIO(content),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('调整表模板.xlsx')}"},
    )


@router.post(
    "/parse-changes",
    response_model=ParseChangesResponse,
    responses={422: {"model": ErrorResponse}},
)
def parse_changes(
    file: UploadFile = File(...),
    known_room_ids: str | None = Form(default=None),
    known_pod_ids: str | None = Form(default=None),
    known_batches: str | None = Form(default=None),
):
    warnings: list[str] = []
    fatal: list[str] = []
    room_ids = _parse_id_set(known_room_ids)
    pod_ids = _parse_id_set(known_pod_ids)
    batch_meta = _parse_batch_meta(known_batches)
    batch_ids = set(batch_meta)

    try:
        file.file.seek(0)
        workbook = load_workbook(BytesIO(file.file.read()), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises mixed parse exceptions.
        return _import_error([f"变更表无法读取：{exc}"])

    try:
        room_sheet = _required_sheet(workbook, ROOM_SHEET, ROOM_COLUMNS, fatal)
        arrival_sheet = _required_sheet(workbook, ARRIVAL_SHEET, ARRIVAL_COLUMNS, fatal)
        batch_sheet = _required_sheet(workbook, BATCH_SHEET, BATCH_COLUMNS, fatal)
        if fatal:
            return _import_error(fatal)

        try:
            source_arrivals = _current_project_arrivals_by_pod()
        except DataImportError as exc:
            return _import_error([f"当前项目到货数据不可用：{exc}"])
        rooms = _parse_room_rows(room_sheet, room_ids, warnings)
        arrivals = _parse_arrival_rows(arrival_sheet, pod_ids, source_arrivals, warnings)
        batches = _parse_batch_rows(batch_sheet, batch_ids, batch_meta, warnings)
    finally:
        workbook.close()

    return ParseChangesResponse(
        changes=ChangeSet(
            rooms=list(rooms.values()),
            arrivals=list(arrivals.values()),
            batches=list(batches.values()),
        ),
        warnings=warnings,
    )


def _prefilled_change_template() -> bytes:
    bundle = load_input_bundle()
    workbook = load_workbook(TEMPLATE_FILE)
    try:
        if ARRIVAL_SHEET not in workbook.sheetnames:
            raise KeyError(f"缺少sheet「{ARRIVAL_SHEET}」")
        _rewrite_arrival_sheet(workbook[ARRIVAL_SHEET], bundle.arrivals, include_dates=False)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def _rewrite_arrival_sheet(sheet, arrivals: list[ArrivalItem], *, include_dates: bool) -> None:
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    if sheet.max_column:
        sheet.delete_cols(1, sheet.max_column)

    sheet.append(list(ARRIVAL_COLUMNS))
    for arrival in arrivals:
        sheet.append(
            [
                arrival.pod_id,
                arrival.device_type,
                arrival.device_model,
                arrival.quantity,
                arrival.unit,
                arrival.arrival_date if include_dates else None,
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
    widths = {"A": 18, "B": 14, "C": 30, "D": 10, "E": 8, "F": 14}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet["F"]:
        cell.number_format = "yyyy-mm-dd"


def _current_project_arrivals_by_pod() -> dict[str, list[ArrivalItem]]:
    bundle = load_input_bundle()
    arrivals_by_pod: dict[str, list[ArrivalItem]] = {}
    for arrival in bundle.arrivals:
        arrivals_by_pod.setdefault(arrival.pod_id, []).append(arrival)
    return arrivals_by_pod


def _parse_room_rows(sheet, known_ids: set[str], warnings: list[str]) -> dict[str, Room]:
    headers = _headers(sheet)
    rooms: dict[str, Room] = {}
    for row_number, row in _data_rows(sheet):
        if not _has_any_value(row):
            continue
        room_id = _text(_cell(row, headers, "机房id"))
        if not room_id:
            warnings.append(f"「{ROOM_SHEET}」第{row_number}行缺少机房id，已跳过。")
            continue
        if known_ids and room_id not in known_ids:
            warnings.append(f"「{ROOM_SHEET}」第{row_number}行机房id「{room_id}」不在当前盘子中，已跳过。")
            continue

        fields: dict[str, date] = {}
        mapping = {
            "可布线": "cabling_ready_date",
            "可装设备": "install_ready_date",
            "可通液": "liquid_ready_date",
        }
        for column, field_name in mapping.items():
            parsed = _optional_date(_cell(row, headers, column), ROOM_SHEET, row_number, column, warnings)
            if parsed is not None and parsed is not _INVALID_DATE:
                fields[field_name] = parsed
        if fields:
            rooms[room_id] = Room(room_id=room_id, **fields)
    return rooms


def _parse_arrival_rows(
    sheet,
    known_ids: set[str],
    source_arrivals: dict[str, list[ArrivalItem]],
    warnings: list[str],
) -> dict[str, ArrivalItem]:
    headers = _headers(sheet)
    as_of_date = get_as_of_date()
    latest_by_pod: dict[str, date] = {}
    incomplete_pods: set[str] = set()

    for row_number, row in _data_rows(sheet):
        if not _has_any_value(row):
            continue
        pod_id = _text(_cell(row, headers, "管理单元"))
        if not pod_id:
            warnings.append(f"「{ARRIVAL_SHEET}」第{row_number}行缺少管理单元，已跳过。")
            continue
        if known_ids and pod_id not in known_ids:
            warnings.append(f"「{ARRIVAL_SHEET}」第{row_number}行管理单元「{pod_id}」不在当前盘子中，已跳过。")
            continue

        raw_date = _cell(row, headers, "到货日期")
        if not _text(raw_date):
            warnings.append(f"「{ARRIVAL_SHEET}」第{row_number}行《到货日期》未填，管理单元「{pod_id}」未聚合。")
            incomplete_pods.add(pod_id)
            continue
        arrival_date = _optional_date(raw_date, ARRIVAL_SHEET, row_number, "到货日期", warnings)
        if arrival_date is None or arrival_date is _INVALID_DATE:
            incomplete_pods.add(pod_id)
            continue
        current = latest_by_pod.get(pod_id)
        if current is None or arrival_date > current:
            latest_by_pod[pod_id] = arrival_date

    arrivals: dict[str, ArrivalItem] = {}
    for pod_id, arrival_date in latest_by_pod.items():
        if pod_id in incomplete_pods:
            continue
        source_items = source_arrivals.get(pod_id)
        if not source_items:
            arrivals[f"arrival-{pod_id}"] = _fallback_arrival(pod_id, arrival_date, as_of_date)
            continue
        for source in source_items:
            arrivals[source.arrival_id] = source.model_copy(
                update={
                    "arrival_date": arrival_date,
                    "arrival_status": _arrival_status(arrival_date, as_of_date),
                    "note": _merge_note(source.note, "变更表上传按管理单元聚合"),
                }
            )
    return arrivals


def _fallback_arrival(pod_id: str, arrival_date: date, as_of_date: date) -> ArrivalItem:
    return ArrivalItem(
        arrival_id=f"arrival-{pod_id}",
        pod_id=pod_id,
        device_type="设备到货齐套",
        device_model=None,
        unit="批",
        quantity=1,
        arrival_date=arrival_date,
        arrival_status=_arrival_status(arrival_date, as_of_date),
        note="变更表上传按管理单元聚合",
    )


def _arrival_status(arrival_date: date, as_of_date: date | None = None) -> str:
    return "已到货" if arrival_date <= (as_of_date or get_as_of_date()) else "在途"


def _merge_note(existing: str | None, extra: str) -> str:
    if not existing:
        return extra
    if extra in existing:
        return existing
    return f"{existing}；{extra}"


def _parse_batch_rows(
    sheet,
    known_ids: set[str],
    batch_meta: dict[str, dict[str, Any]],
    warnings: list[str],
) -> dict[str, Batch]:
    headers = _headers(sheet)
    batches: dict[str, Batch] = {}
    for row_number, row in _data_rows(sheet):
        if not _has_any_value(row):
            continue
        batch_id = _text(_cell(row, headers, "batch_id"))
        if not batch_id:
            warnings.append(f"「{BATCH_SHEET}」第{row_number}行缺少batch_id，已跳过。")
            continue
        if known_ids and batch_id not in known_ids:
            warnings.append(f"「{BATCH_SHEET}」第{row_number}行batch_id「{batch_id}」不在当前盘子中，已跳过。")
            continue

        power_on = _optional_date(_cell(row, headers, "上电目标"), BATCH_SHEET, row_number, "上电目标", warnings)
        online = _optional_date(_cell(row, headers, "上线目标"), BATCH_SHEET, row_number, "上线目标", warnings)
        fields: dict[str, date] = {}
        if power_on is not None and power_on is not _INVALID_DATE:
            fields["power_on_target_date"] = power_on
        if online is not None and online is not _INVALID_DATE:
            fields["online_target_date"] = online
        if not fields:
            continue

        meta = batch_meta.get(batch_id, {})
        batches[batch_id] = Batch(
            batch_id=batch_id,
            batch_name=str(meta.get("batch_name") or batch_id),
            pod_ids=[str(item) for item in meta.get("pod_ids", [])],
            **fields,
        )
    return batches


def _required_sheet(workbook, sheet_name: str, required_columns: tuple[str, ...], fatal: list[str]):
    if sheet_name not in workbook.sheetnames:
        fatal.append(f"变更表缺少sheet「{sheet_name}」。")
        return None
    sheet = workbook[sheet_name]
    headers = _headers(sheet)
    missing = [column for column in required_columns if column not in headers]
    if missing:
        fatal.append(f"变更表「{sheet_name}」缺少列：{'、'.join(missing)}。")
        return None
    return sheet


def _headers(sheet) -> dict[str, int]:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None) or ()
    headers: dict[str, int] = {}
    for index, value in enumerate(header_row):
        label = _text(value)
        if label:
            headers[label] = index
    return headers


def _data_rows(sheet):
    yield from enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2)


def _cell(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    index = headers[name]
    return row[index] if index < len(row) else None


def _optional_date(value: Any, sheet: str, row_number: int, column: str, warnings: list[str]) -> date | object | None:
    text = _text(value)
    if not text:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    warnings.append(f"「{sheet}」第{row_number}行《{column}》日期「{text}」无法识别，已跳过该单元格。")
    return _INVALID_DATE


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _has_any_value(row: tuple[Any, ...]) -> bool:
    return any(_text(value) for value in row)


def _parse_id_set(raw: str | None) -> set[str]:
    if not raw:
        return set()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        data = [part.strip() for part in raw.split(",")]
    if not isinstance(data, list):
        return set()
    return {str(item).strip() for item in data if str(item).strip()}


def _parse_batch_meta(raw: str | None) -> dict[str, dict[str, Any]]:
    if not raw:
        return {}
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return {}
    if isinstance(data, dict):
        rows = list(data.values())
    elif isinstance(data, list):
        rows = data
    else:
        return {}

    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        batch_id = str(row.get("batch_id") or "").strip()
        if not batch_id:
            continue
        pod_ids = row.get("pod_ids")
        result[batch_id] = {
            "batch_name": str(row.get("batch_name") or batch_id),
            "pod_ids": pod_ids if isinstance(pod_ids, list) else [],
        }
    return result


def _import_error(messages: list[str]) -> JSONResponse:
    error = ErrorResponse(
        code="IMPORT_ERROR",
        message="变更表不可用，请按固定模板补齐后重新上传。",
        conflicts=[ConflictDetail(constraint="变更表模板", detail=message) for message in messages],
    )
    return JSONResponse(status_code=422, content=error.model_dump(mode="json"))
