from __future__ import annotations

import json
import sys
from io import BytesIO
from datetime import date, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

SCHEDULE_ROOT = Path(__file__).resolve().parents[2]
REPO_ROOT = SCHEDULE_ROOT.parents[1]
sys.path.insert(0, str(REPO_ROOT))

from agent.schedule.api import export_plan as export_plan_module  # noqa: E402
from agent.schedule.app_main import create_app  # noqa: E402
from agent.schedule.importer import DataImportError  # noqa: E402
from agent.schedule.engine import select_shadow_recommendation  # noqa: E402
from agent.schedule.store import PlanVersionStore  # noqa: E402
from agent.schedule.contracts.api import (  # noqa: E402
    ADJUST_PATH,
    COMMIT_PATH,
    EXPORT_PLAN_PATH,
    GENERATE_PATH,
    PROJECT_DATA_PATH,
    AdjustResponse,
    CommitResponse,
    ErrorResponse,
    GenerateResponse,
)
from agent.schedule.contracts.common import TargetRef  # noqa: E402
from agent.schedule.contracts.inputs import (  # noqa: E402
    Activity,
    Anchor,
    Batch,
    DemandRequest,
    Dependency,
    InputBundle,
    Pod,
    Project,
    Room,
    Team,
)

GOLDEN_INPUT_BUNDLE = SCHEDULE_ROOT / "tests" / "fixtures" / "input_bundle.golden.json"


@pytest.fixture()
def store(tmp_path: Path) -> PlanVersionStore:
    return PlanVersionStore(tmp_path / "schedule.sqlite3")


@pytest.fixture()
def client(store: PlanVersionStore) -> TestClient:
    return TestClient(create_app(store=store))


def test_generate_saves_initial_version_and_returns_contract_response(
    client: TestClient,
    store: PlanVersionStore,
):
    bundle = _basic_bundle()

    response = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})

    assert response.status_code == 200
    parsed = GenerateResponse.model_validate(response.json())
    assert parsed.plan.plan_id == "api-demo-plan"
    assert parsed.plan.version == 1
    stored = store.get_plan_version(parsed.plan.plan_id, parsed.plan.version)
    assert stored.plan == parsed.plan
    assert stored.inputs == bundle


def test_project_data_returns_real_input_bundle_with_scenario_card_count(client: TestClient):
    response = client.get(PROJECT_DATA_PATH)

    assert response.status_code == 200
    parsed = InputBundle.model_validate(response.json())
    assert parsed.project.total_card_count == 3456
    assert len(parsed.rooms) == 3
    assert len(parsed.pods) == 9
    assert len(parsed.batches) == 3
    assert len(parsed.teams) == 2
    assert {room.room_id for room in parsed.rooms} == {"B2DH401", "B2DH402", "B2DH403"}


def test_project_data_query_card_count_overrides_scenario_table(client: TestClient):
    response = client.get(PROJECT_DATA_PATH, params={"total_card_count": 10001})

    assert response.status_code == 200
    parsed = InputBundle.model_validate(response.json())
    scale_activity = next(activity for activity in parsed.activities if activity.activity_id == "4.2")
    assert parsed.project.total_card_count == 10001
    assert scale_activity.standard_sla_days == 18
    assert scale_activity.minimum_sla_days == 15


def test_project_data_import_error_returns_error_response(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    def fail_import(*_args, **_kwargs):
        raise DataImportError("IMPORT_ERROR: 批次信息第2行《该批包含的PoD》PoD id 'P-MISSING' 对不上 05_机房机柜信息《PoD名称》")

    monkeypatch.setattr("agent.schedule.api.schedule.load_input_bundle", fail_import)

    response = client.get(PROJECT_DATA_PATH)

    assert response.status_code == 422
    error = ErrorResponse.model_validate(response.json())
    assert error.code == "IMPORT_ERROR"
    assert "P-MISSING" in error.message
    assert error.conflicts[0].constraint == "02_项目数据"


def test_export_plan_returns_xlsx_for_specified_plan_version(
    client: TestClient,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    output_path = tmp_path / "交付计划表.xlsx"
    monkeypatch.setattr(export_plan_module, "DELIVERY_PLAN_FILE", output_path)
    bundle = _basic_bundle()
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())

    response = client.get(
        EXPORT_PLAN_PATH,
        params={"plan_id": base.plan.plan_id, "version": base.plan.version},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(export_plan_module.XLSX_MEDIA_TYPE)
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]
    assert output_path.exists()

    rows = _workbook_rows(response.content)
    assert rows[0] == list(export_plan_module.DELIVERY_PLAN_HEADERS)
    data = [dict(zip(rows[0], row)) for row in rows[1:]]
    install = next(row for row in data if row["ACTIVITY_ID"] == "install")
    assert install["ACTIVITY_NAME"] == "安装 PoD"
    assert install["PROJECT_ID"] == "api-demo"
    assert install["START_DATE"] == date(2026, 1, 2)
    assert install["END_DATE"] == date(2026, 1, 3)
    assert install["STATUS"] == "已排期"
    assert install["MANAGEMENT_UNIT"] == "P1"


def test_export_plan_defaults_to_latest_formal_version(
    client: TestClient,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    output_path = tmp_path / "交付计划表.xlsx"
    monkeypatch.setattr(export_plan_module, "DELIVERY_PLAN_FILE", output_path)
    bundle = _basic_bundle()
    client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})

    response = client.get(EXPORT_PLAN_PATH)

    assert response.status_code == 200
    assert output_path.exists()


def test_generate_transmits_readiness_suggestions(client: TestClient):
    bundle = _readiness_bundle()

    response = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})

    assert response.status_code == 200
    parsed = GenerateResponse.model_validate(response.json())
    assert len(parsed.readiness_suggestions) == 1
    suggestion = parsed.readiness_suggestions[0]
    assert suggestion.batch_id == "B1"
    assert suggestion.suggested_room_ready == {"R1": date(2026, 1, 6)}
    assert suggestion.suggested_arrival == {"P1": date(2026, 1, 6)}


def test_generate_infeasible_returns_error_response(client: TestClient):
    bundle = _basic_bundle(
        anchors=[
            Anchor(
                anchor_id="too-early",
                target=TargetRef(kind="活动实例", ref_id="P1/install"),
                anchor_date=date(2026, 1, 1),
                anchor_source="指定活动",
            )
        ]
    )

    response = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})

    assert response.status_code == 422
    error = ErrorResponse.model_validate(response.json())
    assert error.code == "INFEASIBLE"
    assert error.conflicts


def test_adjust_merges_changes_by_id_and_commit_writes_incremented_version(
    client: TestClient,
    store: PlanVersionStore,
):
    bundle = _basic_bundle()
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())

    changed_room = Room(room_id="R1", install_ready_date=date(2026, 1, 10))
    adjusted = client.post(
        ADJUST_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "changes": {"rooms": [changed_room.model_dump(mode="json")]},
        },
    )

    assert adjusted.status_code == 200
    adjust_body = AdjustResponse.model_validate(adjusted.json())
    assert len(adjust_body.options) == 1
    option = adjust_body.options[0]
    assert option.option_id == "A"
    assert adjust_body.gap is not None
    assert adjust_body.gap.baseline_finish_date == option.plan.project_finish_date
    assert adjust_body.gap.target_date is None
    assert adjust_body.gap.gap_days == 0
    assert option.kpis.gap_days == 0
    activities = {activity.instance_id: activity for activity in option.plan.activities}
    assert activities["R1/ready"].end_date == date(2026, 1, 10)
    assert activities["P1/install"].start_date == date(2026, 1, 10)
    assert adjust_body.explanation.moved_activities

    committed = client.post(
        COMMIT_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "option_id": option.option_id,
        },
    )

    assert committed.status_code == 200
    commit_body = CommitResponse.model_validate(committed.json())
    assert commit_body.new_version == 2
    assert commit_body.plan.version == 2
    stored = store.get_plan_version(base.plan.plan_id, 2)
    assert stored.plan == commit_body.plan
    assert stored.inputs.rooms[0].install_ready_date == date(2026, 1, 10)


def test_commit_duration_overrides_recalculates_selected_option_and_saves_version(
    client: TestClient,
    store: PlanVersionStore,
):
    bundle = _basic_bundle()
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())

    changed_room = Room(room_id="R1", install_ready_date=date(2026, 1, 10))
    adjusted = client.post(
        ADJUST_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "changes": {"rooms": [changed_room.model_dump(mode="json")]},
        },
    )
    option = AdjustResponse.model_validate(adjusted.json()).options[0]

    committed = client.post(
        COMMIT_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "option_id": option.option_id,
            "duration_overrides": {"P1/install": 4},
        },
    )

    assert committed.status_code == 200
    commit_body = CommitResponse.model_validate(committed.json())
    activities = {activity.instance_id: activity for activity in commit_body.plan.activities}
    assert activities["P1/install"].start_date == date(2026, 1, 10)
    assert activities["P1/install"].end_date == date(2026, 1, 13)
    assert activities["P1/install"].actual_sla_days == 4
    stored = store.get_plan_version(base.plan.plan_id, 2)
    assert stored.plan == commit_body.plan


def test_commit_duration_override_below_limit_returns_422(client: TestClient):
    bundle = _basic_bundle()
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())
    adjusted = client.post(
        ADJUST_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "changes": {},
        },
    )
    option = AdjustResponse.model_validate(adjusted.json()).options[0]

    committed = client.post(
        COMMIT_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "option_id": option.option_id,
            "duration_overrides": {"P1/install": 0},
        },
    )

    assert committed.status_code == 422
    error = ErrorResponse.model_validate(committed.json())
    assert error.code == "INFEASIBLE"
    assert error.conflicts[0].constraint == "活动工期格式不合法"


def test_adjust_returns_abc_options_and_transmits_unmet(client: TestClient):
    bundle = _basic_bundle()
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())

    adjusted = client.post(
        ADJUST_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "changes": {
                "demands": [
                    DemandRequest(
                        demand_id="D-tighten-install",
                        target=TargetRef(kind="活动实例", ref_id="P1/install"),
                        direction="某日期前完成",
                        deadline=date(2026, 1, 1),
                    ).model_dump(mode="json")
                ]
            },
        },
    )

    assert adjusted.status_code == 200
    adjust_body = AdjustResponse.model_validate(adjusted.json())
    assert [option.option_id for option in adjust_body.options] == ["A", "B", "C"]
    assert [option.strategy for option in adjust_body.options] == ["均匀压缩", "集中压缩", "站货提拉"]
    assert adjust_body.gap is not None
    assert adjust_body.gap.baseline_finish_date == base.plan.project_finish_date
    assert adjust_body.gap.target_date == date(2026, 1, 1)
    assert adjust_body.gap.gap_days == 2
    assert all(option.kpis.gap_days == 2 for option in adjust_body.options)
    assert adjust_body.unmet
    assert adjust_body.unmet[0].gap_days == 2


def test_commit_any_adjust_option_keeps_full_initial_activity_set_for_real_data(
    client: TestClient,
    store: PlanVersionStore,
):
    bundle = InputBundle.model_validate(json.loads(GOLDEN_INPUT_BUNDLE.read_text(encoding="utf-8")))
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())
    base_activity_ids = {activity.instance_id for activity in base.plan.activities}
    target_id = base.plan.critical_path[-1]
    target_end = next(activity.end_date for activity in base.plan.activities if activity.instance_id == target_id)

    adjusted = client.post(
        ADJUST_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "changes": {
                "demands": [
                    DemandRequest(
                        demand_id="D-full-plan-regression",
                        target=TargetRef(kind="活动实例", ref_id=target_id),
                        direction="某日期前完成",
                        deadline=target_end - timedelta(days=2),
                    ).model_dump(mode="json")
                ]
            },
        },
    )

    assert adjusted.status_code == 200
    adjust_body = AdjustResponse.model_validate(adjusted.json())
    assert [option.option_id for option in adjust_body.options] == ["A", "B", "C"]
    assert adjust_body.gap is not None
    assert adjust_body.gap.baseline_finish_date == base.plan.project_finish_date
    assert adjust_body.gap.target_date == target_end - timedelta(days=2)
    assert adjust_body.gap.gap_days == 2
    assert len(base_activity_ids) == 396

    for option in adjust_body.options:
        committed = client.post(
            COMMIT_PATH,
            json={
                "plan_id": base.plan.plan_id,
                "base_version": base.plan.version,
                "option_id": option.option_id,
            },
        )

        assert committed.status_code == 200
        commit_body = CommitResponse.model_validate(committed.json())
        committed_activity_ids = {activity.instance_id for activity in commit_body.plan.activities}
        stored = store.get_plan_version(base.plan.plan_id, commit_body.new_version)
        assert committed_activity_ids == base_activity_ids
        assert {activity.instance_id for activity in stored.plan.activities} == base_activity_ids


def test_commit_records_shadow_recommendation_log(client: TestClient, store: PlanVersionStore):
    bundle = _basic_bundle()
    generated = client.post(GENERATE_PATH, json={"inputs": bundle.model_dump(mode="json")})
    base = GenerateResponse.model_validate(generated.json())

    adjusted = client.post(
        ADJUST_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "changes": {
                "demands": [
                    DemandRequest(
                        demand_id="D-shadow-log",
                        target=TargetRef(kind="活动实例", ref_id="P1/install"),
                        direction="某日期前完成",
                        deadline=date(2026, 1, 1),
                    ).model_dump(mode="json")
                ]
            },
        },
    )
    assert adjusted.status_code == 200
    adjust_body = AdjustResponse.model_validate(adjusted.json())
    shadow_option = select_shadow_recommendation(adjust_body.options)
    selected_option = next(
        (option for option in adjust_body.options if option.option_id != shadow_option.option_id),
        shadow_option,
    )

    committed = client.post(
        COMMIT_PATH,
        json={
            "plan_id": base.plan.plan_id,
            "base_version": base.plan.version,
            "option_id": selected_option.option_id,
        },
    )

    assert committed.status_code == 200
    commit_body = CommitResponse.model_validate(committed.json())
    logs = store.list_shadow_recommendation_logs(base.plan.plan_id, base.plan.version)
    assert len(logs) == 1
    log = logs[0]
    assert log.new_version == commit_body.new_version
    assert log.shadow_option_id == shadow_option.option_id
    assert log.selected_option_id == selected_option.option_id
    assert log.is_match == (shadow_option.option_id == selected_option.option_id)
    assert log.kpi_snapshot["shadow"] == shadow_option.kpis.model_dump(mode="json")
    assert log.kpi_snapshot["selected"] == selected_option.kpis.model_dump(mode="json")
    assert log.kpi_snapshot["duration_overrides_applied"] is False


def test_healthz(client: TestClient):
    response = client.get("/healthz")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def _workbook_rows(content: bytes) -> list[list[object]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook["sheet0"]
        rows: list[list[object]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([_normalize_excel_value(value) for value in row])
        return rows
    finally:
        workbook.close()


def _normalize_excel_value(value):
    if hasattr(value, "date"):
        return value.date()
    return value


def _activity(
    activity_id: str,
    name: str,
    scope: str,
    days: int | None = 1,
    *,
    activity_type: str = "普通",
    constraint_source: str | None = None,
) -> Activity:
    return Activity(
        activity_id=activity_id,
        activity_name=name,
        scope=scope,
        activity_type=activity_type,
        constraint_source=constraint_source,
        standard_sla_days=days,
        minimum_sla_days=1 if days else None,
    )


def _basic_bundle(*, anchors: list[Anchor] | None = None) -> InputBundle:
    return InputBundle(
        project=Project(project_id="api-demo", project_name="API demo", start_date=date(2026, 1, 1)),
        rooms=[Room(room_id="R1")],
        pods=[Pod(pod_id="P1", room_id="R1")],
        teams=[Team(team_id="T1", experience="丰富")],
        activities=[
            _activity("ready", "机房就位", "机房级", 1, activity_type="机房准备"),
            _activity("install", "安装 PoD", "PoD级", 2, constraint_source="人"),
        ],
        dependencies=[
            Dependency(from_activity_id="ready", to_activity_id="install", dep_type="FS"),
        ],
        batches=[Batch(batch_id="B1", batch_name="批次1", pod_ids=["P1"])],
        anchors=anchors or [],
    )


def _readiness_bundle() -> InputBundle:
    return InputBundle(
        project=Project(project_id="api-demo", project_name="API demo", start_date=date(2026, 1, 1)),
        rooms=[Room(room_id="R1")],
        pods=[Pod(pod_id="P1", room_id="R1")],
        teams=[Team(team_id="T1", experience="一般")],
        activities=[
            _activity("room_ready", "机房就位", "机房级", 7, activity_type="机房准备"),
            _activity("arrival", "设备到货", "PoD级", 1, activity_type="到货"),
            _activity("install", "机柜上架安装", "PoD级", 5, constraint_source="人"),
            _activity("cabling", "综合布线", "PoD级", 4, constraint_source="人"),
            _activity("power_on", "上电点亮", "批次级", 1, activity_type="里程碑"),
            _activity("online", "批次上线", "批次级", 1, activity_type="里程碑"),
        ],
        dependencies=[
            Dependency(from_activity_id="room_ready", to_activity_id="install", dep_type="FS"),
            Dependency(from_activity_id="arrival", to_activity_id="install", dep_type="FS"),
            Dependency(from_activity_id="install", to_activity_id="cabling", dep_type="FS"),
            Dependency(from_activity_id="cabling", to_activity_id="power_on", dep_type="FS"),
            Dependency(from_activity_id="power_on", to_activity_id="online", dep_type="FS"),
        ],
        batches=[Batch(batch_id="B1", batch_name="批次1", pod_ids=["P1"], online_target_date=date(2026, 1, 20))],
    )
