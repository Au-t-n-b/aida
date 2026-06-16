from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

SCHEDULE_ROOT = Path(__file__).resolve().parents[2]
REPO_ROOT = SCHEDULE_ROOT.parents[1]
sys.path.insert(0, str(REPO_ROOT))

from agent.schedule.app_main import create_app  # noqa: E402
from agent.schedule.contracts.api import ErrorResponse, PARSE_CHANGES_PATH, ParseChangesResponse  # noqa: E402


@pytest.fixture()
def client() -> TestClient:
    return TestClient(create_app())


def test_parse_changes_accepts_valid_template(client: TestClient):
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_known_form(),
        files=_upload(
            _workbook_bytes(
                room_rows=[["M1", "2026-07-01", None, "2026-07-03"]],
                arrival_rows=[
                    ["P-01", "计算柜", "A", 1, "柜", "2026-07-09"],
                    ["P-01", "网络", "B", 1, "台", "2026-07-10"],
                ],
                batch_rows=[["bt-1", "2026-08-01", "2026-08-10"]],
            )
        ),
    )

    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert parsed.warnings == []
    assert parsed.changes.rooms[0].room_id == "M1"
    assert str(parsed.changes.rooms[0].cabling_ready_date) == "2026-07-01"
    assert parsed.changes.arrivals[0].pod_id == "P-01"
    assert str(parsed.changes.arrivals[0].arrival_date) == "2026-07-10"
    assert parsed.changes.batches[0].batch_id == "bt-1"
    assert parsed.changes.batches[0].pod_ids == ["P-01", "P-02"]


def test_parse_changes_arrival_status_uses_configured_as_of_date(client: TestClient, monkeypatch):
    payload = _workbook_bytes(
        arrival_rows=[
            ["P-01", "计算柜", "A", 1, "柜", "2026-07-10"],
        ],
    )

    monkeypatch.setenv("SCHEDULE_AS_OF_DATE", "2026-07-09")
    response = client.post(PARSE_CHANGES_PATH, data=_known_form(), files=_upload(payload))
    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert parsed.changes.arrivals[0].arrival_status == "在途"

    monkeypatch.setenv("SCHEDULE_AS_OF_DATE", "2026-07-10")
    response = client.post(PARSE_CHANGES_PATH, data=_known_form(), files=_upload(payload))
    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert parsed.changes.arrivals[0].arrival_status == "已到货"


def test_parse_changes_returns_import_error_when_required_column_missing(client: TestClient):
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_known_form(),
        files=_upload(
            _workbook_bytes(
                arrival_headers=["管理单元"],
                arrival_rows=[["P-01"]],
            )
        ),
    )

    assert response.status_code == 422
    error = ErrorResponse.model_validate(response.json())
    assert error.code == "IMPORT_ERROR"
    assert any("缺少列：" in (conflict.detail or "") and "到货日期" in conflict.detail for conflict in error.conflicts)


def test_parse_changes_uses_real_project_arrival_rows_after_pod_max(client: TestClient):
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_known_form(pod_ids=["B2DH401-POD01"]),
        files=_upload(
            _workbook_bytes(
                arrival_rows=[
                    ["B2DH401-POD01", "灵衢线缆", "400G MPO线缆", 2688, "根", "2026-01-08"],
                    ["B2DH401-POD01", "计算柜", "Atlas 900 A3 SuperPoD计算柜", 12, "柜", "2026-01-10"],
                    ["B2DH401-POD01", "总线设备柜", "Atlas 900 A3 SuperPoD总线设备柜", 4, "柜", "2026-01-09"],
                    ["B2DH401-POD01", "网络", "CE9866-128DQ", 6, "台", "2026-01-07"],
                ],
            )
        ),
    )

    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    pod_arrivals = [arrival for arrival in parsed.changes.arrivals if arrival.pod_id == "B2DH401-POD01"]
    assert len(pod_arrivals) == 4
    assert {arrival.arrival_id for arrival in pod_arrivals} == {"1", "2", "3", "4"}
    assert {str(arrival.arrival_date) for arrival in pod_arrivals} == {"2026-01-10"}
    assert {arrival.device_type for arrival in pod_arrivals} == {"灵衢线缆", "计算柜", "总线设备柜", "网络"}


def test_parse_changes_stage2_sample_aggregates_jd_expected_dates(client: TestClient):
    sample_path = SCHEDULE_ROOT / "project-data" / "10_变更表模板" / "stage-2-调整表-样例.xlsx"
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_jd_known_form(),
        files=_upload(sample_path.read_bytes()),
    )

    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert parsed.warnings == []
    by_pod: dict[str, set[str]] = {}
    for arrival in parsed.changes.arrivals:
        by_pod.setdefault(arrival.pod_id, set()).add(str(arrival.arrival_date))
    assert by_pod["B2DH401-POD01"] == {"2026-01-10"}
    assert by_pod["B2DH403-POD09"] == {"2026-01-18"}
    for pod_id in ["B2DH402-POD05", "B2DH402-POD06", "B2DH402-POD07", "B2DH402-POD08"]:
        assert by_pod[pod_id] == {"2026-03-04"}


def test_parse_changes_warns_and_skips_invalid_date_cells(client: TestClient):
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_known_form(),
        files=_upload(
            _workbook_bytes(
                room_rows=[["M1", "不是日期", "2026-07-02", None]],
            )
        ),
    )

    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert any("日期" in warning and "无法识别" in warning for warning in parsed.warnings)
    assert parsed.changes.rooms[0].room_id == "M1"
    assert parsed.changes.rooms[0].cabling_ready_date is None
    assert str(parsed.changes.rooms[0].install_ready_date) == "2026-07-02"


def test_parse_changes_warns_and_skips_blank_arrival_date(client: TestClient):
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_known_form(),
        files=_upload(
            _workbook_bytes(
                arrival_rows=[
                    ["P-01", "计算柜", "A", 1, "柜", "2026-07-10"],
                    ["P-01", "网络", "B", 1, "台", None],
                ],
            )
        ),
    )

    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert any("到货日期" in warning and "未填" in warning for warning in parsed.warnings)
    assert parsed.changes.arrivals == []


def test_parse_changes_warns_and_skips_unknown_ids(client: TestClient):
    response = client.post(
        PARSE_CHANGES_PATH,
        data=_known_form(),
        files=_upload(
            _workbook_bytes(
                room_rows=[["MX", "2026-07-01", None, None]],
                arrival_rows=[["PX", "计算柜", "A", 1, "柜", "2026-07-10"]],
                batch_rows=[["bt-X", "2026-08-01", None]],
            )
        ),
    )

    assert response.status_code == 200
    parsed = ParseChangesResponse.model_validate(response.json())
    assert len(parsed.warnings) == 3
    assert parsed.changes.rooms == []
    assert parsed.changes.arrivals == []
    assert parsed.changes.batches == []


def test_change_template_download_returns_xlsx(client: TestClient):
    response = client.get("/api/v1/schedule/change-template")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    loaded = _load_workbook_bytes(response.content)
    try:
        assert loaded.sheetnames == ["机房ready", "到货", "批次目标"]
        arrival_sheet = loaded["到货"]
        assert [arrival_sheet.cell(1, column).value for column in range(1, 7)] == [
            "管理单元",
            "设备类型",
            "型号",
            "数量",
            "单位",
            "到货日期",
        ]
        assert arrival_sheet.max_row == 37
        assert arrival_sheet.cell(2, 1).value == "B2DH401-POD01"
        assert arrival_sheet.cell(2, 6).value is None
    finally:
        loaded.close()


def _known_form(pod_ids: list[str] | None = None) -> dict[str, str]:
    return {
        "known_room_ids": json.dumps(["M1"], ensure_ascii=False),
        "known_pod_ids": json.dumps(pod_ids or ["P-01"], ensure_ascii=False),
        "known_batches": json.dumps(
            [{"batch_id": "bt-1", "batch_name": "批次1", "pod_ids": ["P-01", "P-02"]}],
            ensure_ascii=False,
        ),
    }


def _jd_known_form() -> dict[str, str]:
    return {
        "known_room_ids": json.dumps(["B2DH401", "B2DH403", "B2DH402"], ensure_ascii=False),
        "known_pod_ids": json.dumps(
            [
                "B2DH401-POD01",
                "B2DH401-POD02",
                "B2DH401-POD03",
                "B2DH401-POD04",
                "B2DH403-POD09",
                "B2DH402-POD05",
                "B2DH402-POD06",
                "B2DH402-POD07",
                "B2DH402-POD08",
            ],
            ensure_ascii=False,
        ),
        "known_batches": json.dumps(
            [
                {
                    "batch_id": "批次1",
                    "batch_name": "批次1",
                    "pod_ids": ["B2DH401-POD01", "B2DH401-POD02", "B2DH401-POD03", "B2DH401-POD04"],
                },
                {"batch_id": "批次2", "batch_name": "批次2", "pod_ids": ["B2DH403-POD09"]},
                {
                    "batch_id": "批次3",
                    "batch_name": "批次3",
                    "pod_ids": ["B2DH402-POD05", "B2DH402-POD06", "B2DH402-POD07", "B2DH402-POD08"],
                },
            ],
            ensure_ascii=False,
        ),
    }


def _upload(content: bytes):
    return {
        "file": (
            "changes.xlsx",
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


def _workbook_bytes(
    *,
    room_headers: list[str] | None = None,
    room_rows: list[list[object | None]] | None = None,
    arrival_headers: list[str] | None = None,
    arrival_rows: list[list[object | None]] | None = None,
    batch_headers: list[str] | None = None,
    batch_rows: list[list[object | None]] | None = None,
) -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _sheet(workbook, "机房ready", room_headers or ["机房id", "可布线", "可装设备", "可通液"], room_rows or [])
    _sheet(
        workbook,
        "到货",
        arrival_headers or ["管理单元", "设备类型", "型号", "数量", "单位", "到货日期"],
        arrival_rows or [],
    )
    _sheet(workbook, "批次目标", batch_headers or ["batch_id", "上电目标", "上线目标"], batch_rows or [])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object | None]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def _load_workbook_bytes(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(BytesIO(content), read_only=True, data_only=True)
