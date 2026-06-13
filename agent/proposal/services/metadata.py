"""Metadata (预案版本信息表) service."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from agent.proposal.auth import ProposalSession, proposal_operator_display_name
from agent.proposal.draft_store import (
    format_display_datetime,
    load_chapter_02,
    load_chapter_81,
    load_chapter_82,
    load_chapter_83,
    load_chapter_84,
    load_manifest,
    load_version_info,
    physical_project_root,
    save_manifest,
    _now_iso,
)
from agent.proposal.errors import ProposalApiError
from agent.proposal.version_info_store import (
    DRAFT_VERSION_LABEL,
    append_snapshot,
    find_snapshot,
    import_legacy_version_info,
    list_snapshots,
    load_draft_row,
    save_draft_row,
    sync_xlsx,
)
from agent.constants.project_paths import proposal_paths

CHAPTER_LABELS: dict[str, str] = {
    "2": "2 设备配置信息",
    "8.1": "8.1 服务交付界面",
    "8.2": "8.2 服务内容",
    "8.3": "8.3 维保策略",
    "8.4": "8.4 维保SLA",
}

CHAPTER_LOADERS: dict[str, Any] = {
    "2": lambda pid, ver: load_chapter_02(pid, ver),
    "8.1": lambda pid, ver: load_chapter_81(pid, ver),
    "8.2": lambda pid, ver: load_chapter_82(pid, ver),
    "8.3": lambda pid, ver: load_chapter_83(pid, ver),
    "8.4": lambda pid, ver: load_chapter_84(pid, ver),
}


def _contract_basic_dir(project_id: str) -> Path:
    rel = proposal_paths("")["contract_project_basic_out"]
    return physical_project_root(project_id) / Path(rel)


def _read_json_basic(path: Path) -> dict[str, str] | None:
    if not path.with_suffix(".json").exists():
        return None
    data = json.loads(path.with_suffix(".json").read_text(encoding="utf-8"))
    if isinstance(data, list) and data:
        data = data[0]
    if not isinstance(data, dict):
        return None
    return {
        "projectId": str(
            data.get("项目ID")
            or data.get("projectId")
            or data.get("项目编码")
            or ""
        ),
        "projectName": str(
            data.get("项目名称") or data.get("projectName") or ""
        ),
    }


def _read_xlsx_basic(path: Path) -> dict[str, str] | None:
    xlsx = path.with_suffix(".xlsx")
    if not xlsx.exists():
        return None
    try:
        wb = load_workbook(xlsx, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
        wb.close()
    except (OSError, ValueError):
        return None
    if len(rows) < 2:
        return None
    headers = [str(h or "").strip() for h in rows[0]]
    values = rows[1]
    mapping = dict(zip(headers, values))
    project_id = str(mapping.get("项目ID") or mapping.get("项目编码") or "")
    project_name = str(mapping.get("项目名称") or "")
    if not project_id and not project_name:
        return None
    return {"projectId": project_id, "projectName": project_name}


def read_project_basic_info(project_id: str) -> tuple[dict[str, str], list[dict[str, str]]]:
    """Return (basic_info, dependencies)."""
    base_dir = _contract_basic_dir(project_id)
    info = _read_xlsx_basic(base_dir) or _read_json_basic(base_dir)
    dependencies: list[dict[str, str]] = []
    pid = (info.get("projectId") if info else None) or project_id
    pname = (info.get("projectName") if info else None) or f"{project_id}项目"
    if info and info.get("projectId") and info.get("projectName"):
        return {"projectId": pid, "projectName": pname}, dependencies

    dependencies.append(
        {
            "code": "SYS_DEPENDENCY",
            "message": "合同/输出结果/项目基础信息表 不可用，使用 projectId fallback",
        }
    )
    return {"projectId": pid, "projectName": pname}, dependencies


def _auto_document_summary(project_name: str) -> str:
    name = (project_name or "").strip()
    if not name:
        return "交付预案"
    if name.endswith("交付预案"):
        return name
    return f"{name}交付预案"


def _chapter_checksum(project_id: str, version: str, chapter_key: str) -> str:
    loader = CHAPTER_LOADERS[chapter_key]
    payload = loader(project_id, version)
    normalized = json.dumps(payload, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def compute_chapter_change_description(
    project_id: str,
    prev_version: str | None,
) -> str:
    changed: list[str] = []
    for key, label in CHAPTER_LABELS.items():
        draft_hash = _chapter_checksum(project_id, "draft", key)
        if prev_version:
            prev_hash = _chapter_checksum(project_id, prev_version, key)
            if draft_hash != prev_hash:
                changed.append(label)
        else:
            if draft_hash != loader_empty(key):
                changed.append(label)
    return "；".join(changed) if changed else "无章节变化"


def loader_empty(chapter_key: str) -> str:
    defaults = {
        "2": {"rows": []},
        "8.1": {"rows": []},
        "8.2": {"rows": []},
        "8.3": {"rows": []},
        "8.4": {"hardwareSupport": "", "serviceLevel": "", "rows": []},
    }
    normalized = json.dumps(defaults[chapter_key], sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _manual_to_manifest_records(manual: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "seq": raw.get("seq") or idx,
            "chapter": str(raw.get("chapter") or ""),
            "description": str(raw.get("changeDescription") or raw.get("description") or ""),
        }
        for idx, raw in enumerate(manual, start=1)
    ]


def _sync_manifest_change_records(project_id: str, manual: list[dict[str, Any]]) -> None:
    """草稿修改记录双写 manifest.json，兼容旧后端与前端 fallback。"""
    manifest = load_manifest(project_id)
    manifest["changeRecords"] = _manual_to_manifest_records(manual)
    save_manifest(project_id, manifest)


def _collect_manual_entries(
    project_id: str, draft: dict[str, Any]
) -> list[dict[str, Any]]:
    manual = list(draft.get("manualChangeLog") or [])
    if manual:
        return manual
    manifest = load_manifest(project_id)
    rows: list[dict[str, Any]] = []
    for idx, raw in enumerate(manifest.get("changeRecords") or [], start=1):
        chapter = str(raw.get("chapter") or "").strip()
        desc = str(raw.get("description") or raw.get("changeDescription") or "").strip()
        if not chapter and not desc:
            continue
        rows.append({"seq": raw.get("seq") or idx, "chapter": chapter, "changeDescription": desc})
    return rows


def _format_manual_change_description(manual: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for raw in manual:
        chapter = str(raw.get("chapter") or "").strip()
        desc = str(raw.get("changeDescription") or raw.get("description") or "").strip()
        if chapter and desc:
            parts.append(f"{chapter}：{desc}")
        elif chapter:
            parts.append(chapter)
        elif desc:
            parts.append(desc)
    return "；".join(parts)


def _serialize_row(row: dict[str, Any]) -> dict[str, Any]:
    created_at = row.get("createdAt")
    updated_at = row.get("updatedAt")
    return {
        "projectId": row.get("projectId") or "",
        "projectName": row.get("projectName") or "",
        "proposalVersion": row.get("proposalVersion") or DRAFT_VERSION_LABEL,
        "createdBy": row.get("createdBy"),
        "createdAt": format_display_datetime(created_at) if created_at else None,
        "updatedBy": row.get("updatedBy"),
        "updatedAt": format_display_datetime(updated_at) if updated_at else None,
        "changeDescription": row.get("changeDescription") or "",
        "documentSummary": row.get("documentSummary") or "",
    }


def _hydrate_draft_row(
    draft: dict[str, Any],
    basic: dict[str, str],
    operator: str,
    *,
    project_id: str = "",
    touch: bool = False,
    is_new: bool = False,
) -> dict[str, Any]:
    now = _now_iso()
    draft["projectId"] = basic["projectId"]
    draft["projectName"] = basic["projectName"]
    draft["proposalVersion"] = DRAFT_VERSION_LABEL
    draft.setdefault("manualChangeLog", [])

    if not draft.get("documentSummary"):
        draft["documentSummary"] = _auto_document_summary(basic["projectName"])

    manifest = load_manifest(project_id) if project_id else {}
    manifest_creator = manifest.get("createdBy")
    manifest_created_at = manifest.get("createdAt")
    manifest_updated_at = manifest.get("updatedAt")
    latest_release_updated_at = None
    latest_release_updated_by = None
    latest_release_created_at = None
    latest_release = manifest.get("latestReleaseVersion") if manifest else None
    if (
        project_id
        and latest_release
        and manifest.get("baseProposalVersion") == latest_release
        and manifest.get("dirty") is False
    ):
        try:
            latest_info = load_version_info(project_id, str(latest_release))
            latest_release_created_at = latest_info.get("createdAt")
            latest_release_updated_at = latest_info.get("updatedAt")
            latest_release_updated_by = latest_info.get("updatedBy")
        except Exception:
            latest_release_created_at = None
            latest_release_updated_at = None
            latest_release_updated_by = None
    if manifest_creator in (None, "", "frontend"):
        manifest_creator = None

    if is_new:
        draft["createdBy"] = manifest_creator or operator
        draft["createdAt"] = manifest_created_at or latest_release_created_at or now
        draft["updatedBy"] = latest_release_updated_by or operator
        draft["updatedAt"] = (
            latest_release_updated_at
            or manifest_updated_at
            or draft["createdAt"]
            or now
        )
    elif touch:
        if not draft.get("createdBy") or draft.get("createdBy") == "frontend":
            draft["createdBy"] = manifest_creator or operator
        if not draft.get("createdAt"):
            draft["createdAt"] = now
        draft["updatedBy"] = operator
        draft["updatedAt"] = now
    else:
        if (
            project_id
            and not touch
            and str(draft.get("proposalVersion") or "") == DRAFT_VERSION_LABEL
            and manifest.get("dirty") is False
            and manifest.get("baseProposalVersion") == manifest.get("latestReleaseVersion")
        ):
            aligned_updated_at = latest_release_updated_at or manifest_updated_at
            if aligned_updated_at:
                draft["updatedAt"] = aligned_updated_at
            if latest_release_updated_by:
                draft["updatedBy"] = latest_release_updated_by
            elif manifest.get("updatedBy"):
                draft["updatedBy"] = manifest.get("updatedBy")
        if (
            not draft.get("createdBy")
            or draft.get("createdBy") == "frontend"
            or not draft.get("createdAt")
        ):
            draft["createdBy"] = draft.get("createdBy") or manifest_creator or operator
            draft["createdAt"] = draft.get("createdAt") or now
        if not draft.get("updatedBy"):
            draft["updatedBy"] = operator
        if not draft.get("updatedAt"):
            draft["updatedAt"] = now

    return draft


def ensure_metadata_draft(
    project_id: str,
    operator: str,
    *,
    touch: bool = False,
) -> tuple[dict[str, Any], list[dict[str, str]]]:
    import_legacy_version_info(project_id)
    basic, deps = read_project_basic_info(project_id)
    draft = load_draft_row(project_id)

    if draft is None:
        draft = _hydrate_draft_row(
            {
                "changeDescription": "",
                "documentSummary": "",
                "manualChangeLog": [],
            },
            basic,
            operator,
            project_id=project_id,
            is_new=True,
        )
    else:
        draft = _hydrate_draft_row(
            draft, basic, operator, project_id=project_id, touch=touch
        )

    save_draft_row(project_id, draft)
    return _serialize_row(draft), deps


def get_metadata(
    project_id: str,
    version: str = "draft",
    operator: str | None = None,
) -> tuple[dict[str, Any], list[dict[str, str]]]:
    op = operator or proposal_operator_display_name()
    if version == "draft":
        return ensure_metadata_draft(project_id, op)

    import_legacy_version_info(project_id)
    snap = find_snapshot(project_id, version)
    if snap:
        return _serialize_row(snap), []

    from agent.proposal.draft_store import load_version_info

    legacy = load_version_info(project_id, version)
    if legacy.get("proposalVersion"):
        basic, deps = read_project_basic_info(project_id)
        legacy.setdefault("projectId", basic["projectId"])
        legacy.setdefault("projectName", basic["projectName"])
        if not legacy.get("documentSummary"):
            legacy["documentSummary"] = _auto_document_summary(basic["projectName"])
        return _serialize_row(legacy), deps

    raise ProposalApiError(404, "VERSION_NOT_FOUND", f"版本 {version} 不存在")


def touch_metadata_on_save(project_id: str, operator: str) -> dict[str, Any]:
    row, _ = ensure_metadata_draft(project_id, operator, touch=True)
    return row


def patch_metadata_draft(
    project_id: str,
    document_summary: str | None,
    session: ProposalSession,
) -> dict[str, Any]:
    operator = proposal_operator_display_name(session)
    ensure_metadata_draft(project_id, operator)
    draft_raw = load_draft_row(project_id) or {}

    if document_summary is not None:
        draft_raw["documentSummary"] = document_summary
    draft_raw["updatedBy"] = operator
    draft_raw["updatedAt"] = _now_iso()
    draft_raw["proposalVersion"] = DRAFT_VERSION_LABEL
    save_draft_row(project_id, draft_raw)
    return _serialize_row(draft_raw)


def save_manual_change_log(
    project_id: str,
    entries: list[dict[str, Any]],
    operator: str,
) -> list[dict[str, Any]]:
    basic, _ = read_project_basic_info(project_id)
    draft = load_draft_row(project_id) or {}
    draft = _hydrate_draft_row(
        draft, basic, operator, project_id=project_id, touch=True
    )

    manual: list[dict[str, Any]] = []
    for idx, raw in enumerate(entries, start=1):
        chapter = str(raw.get("chapter") or "").strip()
        desc = str(raw.get("changeDescription") or raw.get("description") or "").strip()
        if not chapter and not desc:
            continue
        manual.append(
            {
                "seq": idx,
                "chapter": chapter,
                "changeDescription": desc,
            }
        )

    draft["manualChangeLog"] = manual
    save_draft_row(project_id, draft)
    _sync_manifest_change_records(project_id, manual)
    return build_cumulative_change_log(project_id)


def promote_metadata_on_release(
    project_id: str,
    new_version: str,
    operator: str,
    *,
    prev_version: str | None,
) -> dict[str, Any]:
    basic, _ = read_project_basic_info(project_id)
    now = _now_iso()
    draft = load_draft_row(project_id) or {}
    draft = _hydrate_draft_row(
        draft, basic, operator, project_id=project_id, touch=True
    )

    manual_entries = _collect_manual_entries(project_id, draft)
    change_description = _format_manual_change_description(manual_entries)

    created_by = draft.get("createdBy") or operator
    created_at = draft.get("createdAt") or now

    snapshot = {
        "projectId": basic["projectId"],
        "projectName": basic["projectName"],
        "proposalVersion": new_version,
        "createdBy": created_by,
        "createdAt": created_at,
        "updatedBy": operator,
        "updatedAt": now,
        "changeDescription": change_description,
        "documentSummary": draft.get("documentSummary")
        or _auto_document_summary(basic["projectName"]),
        "changeRecords": _manual_to_manifest_records(manual_entries),
    }
    append_snapshot(project_id, snapshot)

    refreshed_draft = {
        **draft,
        "projectId": basic["projectId"],
        "projectName": basic["projectName"],
        "proposalVersion": DRAFT_VERSION_LABEL,
        "createdBy": created_by,
        "createdAt": created_at,
        "updatedBy": operator,
        "updatedAt": now,
        "changeDescription": "",
        "manualChangeLog": [],
    }
    save_draft_row(project_id, refreshed_draft)
    _sync_manifest_change_records(project_id, [])
    manifest = load_manifest(project_id)
    if manifest.get("createdBy") in (None, "", "frontend"):
        manifest["createdBy"] = created_by
        save_manifest(project_id, manifest)
    sync_xlsx(project_id)
    return _serialize_row(snapshot)


def build_cumulative_change_log(project_id: str) -> list[dict[str, Any]]:
    """累计修改记录：仅拼接各版本手工 changeRecords + 当前草稿 manualChangeLog（不含章节 diff 自动生成）。"""
    import_legacy_version_info(project_id)
    entries: list[dict[str, Any]] = []
    seq = 0

    for snap in list_snapshots(project_id):
        snap_records = snap.get("changeRecords") or []
        for raw in snap_records:
            chapter = str(raw.get("chapter") or "").strip()
            desc = str(raw.get("description") or raw.get("changeDescription") or "").strip()
            if not chapter and not desc:
                continue
            seq += 1
            entries.append(
                {
                    "seq": seq,
                    "chapter": chapter,
                    "changeDescription": desc,
                    "editable": False,
                    "source": "snapshot",
                }
            )

    draft = load_draft_row(project_id) or {}
    manual_rows = _collect_manual_entries(project_id, draft)
    for raw in manual_rows:
        chapter = str(raw.get("chapter") or "").strip()
        desc = str(raw.get("changeDescription") or "").strip()
        if not chapter and not desc:
            continue
        seq += 1
        entries.append(
            {
                "seq": seq,
                "chapter": chapter,
                "changeDescription": desc,
                "editable": True,
                "source": "manual",
            }
        )

    return entries
