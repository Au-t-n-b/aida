"""Chapter 2 设备配置信息 — business logic."""
from __future__ import annotations

import json
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from agent.integrations.occ_datamarket_client import (
    fetch_catalog_by_offering_name,
    fetch_lifecycle_row,
)
from agent.proposal.assemblers.device_boq_assembler import (
    assemble_device_info,
    patch_enrichment,
)
from agent.proposal.auth import proposal_operator_display_name
from agent.proposal.draft_store import (
    assert_draft_editable,
    load_chapter_02,
    manifest_activity_fields,
    product_basic_info_path,
    save_chapter_02_draft,
    touch_draft_manifest,
)
from agent.proposal.errors import ProposalApiError
from agent.proposal.models import DataSource, DeviceInfoRow, PatchDeviceInfoBody


def _debug_log(hypothesis_id: str, location: str, message: str, data: dict[str, Any]) -> None:
    # region agent log
    try:
        payload = {
            "sessionId": "5609cc",
            "runId": "pre-fix",
            "hypothesisId": hypothesis_id,
            "location": location,
            "message": message,
            "data": data,
            "timestamp": int(time.time() * 1000),
        }
        Path("debug-5609cc.log").open("a", encoding="utf-8").write(
            json.dumps(payload, ensure_ascii=False) + "\n"
        )
    except Exception:
        pass
    # endregion


def _row_from_dict(raw: dict[str, Any]) -> DeviceInfoRow:
    payload = dict(raw)
    payload.pop("categoryName", None)
    payload.pop("pbiProductName", None)
    return DeviceInfoRow.model_validate(payload)


def _load_all_rows(project_id: str, version: str) -> list[dict[str, Any]]:
    payload = load_chapter_02(project_id, version)
    return list(payload.get("rows") or [])


def _candidate_product_basic_files(project_id: str) -> list[Path]:
    base = product_basic_info_path(project_id)
    candidates = [
        base.with_suffix(".xlsx"),
        base / "产品基本信息表.xlsx",
        base / "records.xlsx",
    ]
    try:
        from agent.config import BUSINESS_ROOT

        candidates.extend(
            [
                BUSINESS_ROOT / "组织资产" / "产品基本信息表.xlsx",
                BUSINESS_ROOT / "组织资产" / "产品基本信息表" / "产品基本信息表.xlsx",
            ]
        )
    except Exception:
        pass
    return candidates


def _load_u_height_by_model(project_id: str) -> tuple[dict[str, float], bool]:
    path = next((p for p in _candidate_product_basic_files(project_id) if p.exists()), None)
    if path is None:
        return {}, True

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except (OSError, ValueError):
        return {}, True
    if not rows:
        return {}, True

    headers = [str(v or "").strip() for v in rows[0]]
    model_idx = next(
        (
            idx
            for idx, h in enumerate(headers)
            if h in {"设备型号", "型号", "产品型号", "deviceModel", "device_model"}
        ),
        -1,
    )
    u_idx = next(
        (
            idx
            for idx, h in enumerate(headers)
            if h in {"设备U高", "U高", "uHeight", "deviceUHeight", "device_u_height"}
        ),
        -1,
    )
    if model_idx < 0 or u_idx < 0:
        return {}, True

    result: dict[str, float] = {}
    for row in rows[1:]:
        model = str(row[model_idx] or "").strip() if model_idx < len(row) else ""
        raw_u = row[u_idx] if u_idx < len(row) else None
        if not model or raw_u in (None, ""):
            continue
        try:
            result[model] = float(str(raw_u).replace("U", "").replace("u", "").strip())
        except ValueError:
            continue
    return result, False


def _row_needs_enrichment(row: dict[str, Any]) -> bool:
    product_name = str(row.get("pbiProductName") or row.get("deviceModel") or "")
    if not product_name:
        return False
    return not any(
        row.get(field)
        for field in (
            "productCode",
            "lifecycleStatus",
            "gaActualDate",
            "gaPlanDate",
            "eomActualDate",
            "eomPlanDate",
            "eosActualDate",
            "eosPlanDate",
        )
    )


def _enrich_rows(
    project_id: str,
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    dependencies: list[dict[str, str]] = []

    u_height_by_model, u_missing = _load_u_height_by_model(project_id)
    if u_missing:
        dependencies.append({
            "code": "SYS_DEPENDENCY",
            "message": "组织资产/产品基本信息表.xlsx 不可用，设备U高未补全",
        })
    for row in rows:
        model = str(row.get("deviceModel") or "")
        if model in u_height_by_model:
            row["deviceUHeight"] = u_height_by_model[model]

    product_names: set[str] = set()
    for row in rows:
        name = str(row.get("pbiProductName") or row.get("deviceModel") or "").strip()
        if name:
            product_names.add(name)

    catalog_cache: dict[str, dict[str, Any] | None] = {}
    if product_names:
        workers = min(6, max(1, len(product_names)))
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {
                name: pool.submit(fetch_catalog_by_offering_name, name)
                for name in product_names
            }
            catalog_cache = {name: fut.result() for name, fut in futures.items()}

    catalog_hits = sum(1 for v in catalog_cache.values() if v)
    missed_names = sorted(name for name in product_names if not catalog_cache.get(name))
    if product_names and catalog_hits == 0:
        dependencies.append({
            "code": "SYS_DEPENDENCY",
            "message": (
                "§17 PBI 全量销售目录未命中：请确认 offering_name 与目录产品名称一致，"
                "或检查 sort_no=1 拆分后的名称"
            ),
        })
    elif missed_names:
        dependencies.append({
            "code": "SYS_DEPENDENCY",
            "message": f"§17 未命中产品名称：{', '.join(missed_names)}",
        })

    lifecycle_cache: dict[str, dict[str, Any] | None] = {}
    offering_nos = {
        str(catalog.get("offering_no") or catalog.get("offeringNo") or "").strip()
        for catalog in catalog_cache.values()
        if catalog
    }
    offering_nos.discard("")
    _debug_log(
        "H3",
        "agent/proposal/services/device_info.py:202",
        "device enrichment catalog lookup stats",
        {
            "projectId": project_id,
            "productNameCount": len(product_names),
            "catalogHits": catalog_hits,
            "missedNames": missed_names[:10],
            "offeringNoCount": len(offering_nos),
        },
    )
    if offering_nos:
        with ThreadPoolExecutor(max_workers=min(6, len(offering_nos))) as pool:
            futures = {
                no: pool.submit(fetch_lifecycle_row, "产品", no) for no in offering_nos
            }
            lifecycle_cache = {no: fut.result() for no, fut in futures.items()}

    for row in rows:
        product_name = str(row.get("pbiProductName") or row.get("deviceModel") or "")
        catalog = catalog_cache.get(product_name)
        if not catalog:
            continue
        offering_no = str(
            catalog.get("offering_no") or catalog.get("offeringNo") or ""
        ).strip()
        lifecycle = lifecycle_cache.get(offering_no) if offering_no else None
        patch_enrichment(row, {"catalog": catalog, "lifecycle": lifecycle})

    # Local fallback: when PBI/OCC is unreachable, keep chapter-2 display usable.
    # Reuse parsed partCode as productCode so frontend table is not blank.
    fallback_count = 0
    for row in rows:
        if row.get("productCode"):
            continue
        part_code = str(row.get("partCode") or "").strip()
        if part_code:
            row["productCode"] = part_code
            fallback_count += 1

    if fallback_count > 0:
        dependencies.append({
            "code": "SYS_DEPENDENCY",
            "message": f"PBI 未命中，已使用 partCode 回填 productCode（{fallback_count} 行）",
        })

    return rows, dependencies


def list_rows(
    project_id: str,
    *,
    version: str = "draft",
) -> tuple[list[DeviceInfoRow], list[dict[str, str]]]:
    dependencies: list[dict[str, str]] = []
    if version != "draft":
        payload = load_chapter_02(project_id, version)
        if not payload.get("rows"):
            raise ProposalApiError(404, "NOT_FOUND", f"版本 {version} 不存在")
        rows = list(payload.get("rows") or [])
    else:
        payload = load_chapter_02(project_id, "draft")
        rows = list(payload.get("rows") or [])
    _debug_log(
        "H4",
        "agent/proposal/services/device_info.py:264",
        "device rows listed",
        {"projectId": project_id, "version": version, "rowCount": len(rows)},
    )
    return [_row_from_dict(r) for r in rows], dependencies


def enrich_draft_rows(
    project_id: str,
    *,
    version: str = "draft",
) -> tuple[list[DeviceInfoRow], list[dict[str, str]]]:
    """补全 draft：§17 product_name → offering_no + GA/EOM/EOS。"""
    assert_draft_editable(project_id, version)
    rows = _load_all_rows(project_id, "draft")
    if not rows:
        raise ProposalApiError(422, "BR_PROPOSAL_PARSE_EMPTY", "设备信息表为空，请先解析设备 BOQ")

    rows, dependencies = _enrich_rows(project_id, rows)
    save_chapter_02_draft(project_id, {"rows": rows})
    touch_draft_manifest(
        project_id,
        updated_by=proposal_operator_display_name(),
        mark_dirty=False,
    )
    return [_row_from_dict(r) for r in rows], dependencies


def parse_device_boq(
    project_id: str,
    *,
    force: bool = False,
    enrich: bool = False,
    version: str = "draft",
) -> tuple[list[DeviceInfoRow], list[dict[str, str]]]:
    assert_draft_editable(project_id, version)

    existing = _load_all_rows(project_id, "draft")
    manual_rows = [r for r in existing if r.get("dataSource") == DataSource.MANUAL.value]

    assembled = assemble_device_info(project_id)
    if force:
        merged = assembled
    else:
        merged = assembled + manual_rows

    if not merged:
        raise ProposalApiError(
            422,
            "BR_PROPOSAL_PARSE_EMPTY",
            "未找到设备 BOQ 解析结果",
        )

    dependencies: list[dict[str, str]] = []
    if enrich:
        merged, dependencies = _enrich_rows(project_id, merged)
    save_chapter_02_draft(project_id, {"rows": merged})
    touch_draft_manifest(
        project_id,
        updated_by=proposal_operator_display_name(),
        mark_dirty=True,
    )
    return [_row_from_dict(r) for r in merged], dependencies


def patch_row(
    project_id: str,
    row_id: str,
    body: PatchDeviceInfoBody,
    *,
    version: str = "draft",
) -> tuple[DeviceInfoRow, dict[str, str | None]]:
    assert_draft_editable(project_id, version)

    rows = _load_all_rows(project_id, "draft")
    target: dict[str, Any] | None = None
    for row in rows:
        if row.get("rowId") == row_id:
            target = row
            break
    if target is None:
        raise ProposalApiError(404, "NOT_FOUND", f"行 {row_id} 不存在")

    updates = body.model_dump(by_alias=True, exclude_unset=True)
    for key, value in updates.items():
        target[key] = value
    target["dataSource"] = DataSource.MANUAL.value

    save_chapter_02_draft(project_id, {"rows": rows})
    manifest = touch_draft_manifest(
        project_id,
        updated_by=proposal_operator_display_name(),
        mark_dirty=True,
    )
    return _row_from_dict(target), manifest_activity_fields(manifest)
