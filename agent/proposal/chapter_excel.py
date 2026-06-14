"""Export chapter JSON payloads to sibling XLSX files."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from agent.proposal.chapter_registry import ChapterSpec

# camelCase API keys → 表头（通用行导出）
_FIELD_LABELS: dict[str, str] = {
    "deviceModel": "设备型号",
    "productCode": "产品编码",
    "quantity": "数量",
    "version": "版本",
    "lifecycleStatus": "生命周期状态",
    "gaActualDate": "GA实际时间",
    "gaPlanDate": "GA计划时间",
    "eomActualDate": "EOM实际时间",
    "eomPlanDate": "EOM计划时间",
    "eosActualDate": "EOS实际时间",
    "eosPlanDate": "EOS计划时间",
    "deviceUHeight": "设备U高",
    "dataSource": "来源",
    "proposalVersion": "预案版本号",
    "productPartCategory": "产品部件类别",
    "partCode": "部件编码",
    "hardwareSubtype": "硬件子类",
    "deviceRole": "设备角色",
    "serviceMajor": "服务大类",
    "serviceItem": "服务细项",
    "deliveryChannel": "交付界面",
    "serviceName": "服务名称",
    "serviceContent": "服务内容",
    "unit": "单位",
    "rowLevel": "层级",
    "offeringId": "Offering ID",
    "offeringName": "名称",
    "seq": "序号",
    "productModel": "产品型号",
    "warrantyPolicy": "保修策略",
    "maintenancePolicy": "维保策略",
    "maintStartDate": "维保开始时间",
    "maintEndDate": "维保结束时间",
    "productEosDate": "产品EOS时间",
    "overEos": "是否超EOS服务",
    "overEosApproval": "超EOS审批结论",
    "severityLevel": "问题等级",
    "coveragePeriod": "服务覆盖时间",
    "responseTime": "响应时间",
    "restoreTime": "回复时间",
    "resolveTime": "解决时间",
    "type": "类型",
    "isHW": "华为软件",
    "software": "软件型号",
    "ver": "版本",
    "source": "来源",
    "remark": "备注",
    "vendor": "厂商",
    "name": "名称",
    "qty": "数量",
    "note": "备注",
    "background": "客户项目背景",
    "goal": "客户项目目标",
    "scope": "客户项目范围",
    "planSummary": "客户项目计划",
    "plan": "客户项目计划",
    "chapter": "章节",
    "changeDescription": "修改描述",
    "projectId": "项目ID",
    "projectName": "项目名称",
    "documentSummary": "文档摘要",
    "createdBy": "创建人",
    "createdAt": "创建时间",
    "updatedBy": "最后修改人",
    "updatedAt": "修改时间",
    "serverRole": "服务器角色",
    "serverModel": "服务器型号",
    "podName": "PoD名称",
    "roomName": "机房名称",
    "paramLeaf": "参数面Leaf柜",
    "bizLeaf": "业务面Leaf柜",
    "sampleLeaf": "样本面Leaf柜",
    "clusterId": "集群ID",
    "clusterType": "集群类型",
    "superPodId": "超节点ID",
    "storageClusterId": "存储集群ID",
    "zoneId": "ZONE ID",
    "ccaeClusterId": "CCAE集群ID",
    "dmeClusterId": "DME集群ID",
    "deviceType": "设备类型",
    "deviceModel": "设备型号",
    "devicePurpose": "设备用途",
    "startDeviceName": "起始设备命名",
    "endDeviceName": "截止设备命名",
    "stack": "技术栈",
    "cat": "活动分类（一级）",
    "act": "活动（二级）",
    "gts": "GTS",
    "hw": "华为云",
    "partner": "伙伴",
    "customer": "客户",
    "scheme": "验收方案",
    "standard": "验收标准",
    "milestone": "验收里程碑",
    "doc": "验收文档",
    "payment": "回款条款",
    "paymentMilestone": "回款里程碑",
    "l1": "一级分类",
    "l2": "二级分类",
    "l3": "用例名称",
    "purpose": "测试目的",
    "topology": "测试组网",
    "pre": "预置条件",
    "steps": "测试步骤",
    "expects": "预期结果",
    "result": "测试结果",
    "selected": "纳入验收",
    "progress": "进度",
    "progressTone": "进度色",
}

_META_SKIP = frozenset({"manualChangeLog", "changeRecords"})
_KEY_BY_LABEL: dict[str, str] = {v: k for k, v in _FIELD_LABELS.items()}

# Chapter templates for header-only initialization when rows are empty.
_CHAPTER_TEMPLATE_KEYS: dict[str, list[str]] = {
    "1": ["key", "label", "value", "proposalVersion"],
    "2": [
        "deviceModel",
        "productCode",
        "quantity",
        "version",
        "lifecycleStatus",
        "gaActualDate",
        "gaPlanDate",
        "eomActualDate",
        "eomPlanDate",
        "eosActualDate",
        "eosPlanDate",
        "deviceUHeight",
        "dataSource",
        "proposalVersion",
        "productPartCategory",
        "partCode",
        "hardwareSubtype",
        "deviceRole",
    ],
    "8.1": ["serviceMajor", "serviceItem", "deliveryChannel", "dataSource", "proposalVersion"],
    "8.2": [
        "serviceName",
        "serviceContent",
        "quantity",
        "unit",
        "rowLevel",
        "dataSource",
        "proposalVersion",
    ],
    "8.3": [
        "seq",
        "productModel",
        "warrantyPolicy",
        "maintenancePolicy",
        "maintStartDate",
        "maintEndDate",
        "productEosDate",
        "overEos",
        "overEosApproval",
        "dataSource",
        "proposalVersion",
    ],
    "8.4": [
        "severityLevel",
        "coveragePeriod",
        "responseTime",
        "restoreTime",
        "resolveTime",
        "serviceItem",
        "dataSource",
        "proposalVersion",
    ],
    "4": [
        "type",
        "isHW",
        "software",
        "ver",
        "vendor",
        "qty",
        "source",
        "remark",
        "proposalVersion",
    ],
    "5.1": ["type", "model", "vendor", "ver", "qty", "source", "note", "proposalVersion"],
    "5.2": ["serverRole", "serverModel", "quantity", "dataSource", "proposalVersion"],
    "5.3": [
        "clusterId",
        "clusterType",
        "superPodId",
        "storageClusterId",
        "zoneId",
        "ccaeClusterId",
        "dmeClusterId",
        "deviceType",
        "vendor",
        "deviceModel",
        "devicePurpose",
        "startDeviceName",
        "endDeviceName",
        "quantity",
        "dataSource",
        "proposalVersion",
    ],
    "6": ["name", "desc", "owner", "state", "date", "proposalVersion"],
    "7": [
        "podName",
        "roomName",
        "compute",
        "bus",
        "paramLeaf",
        "bizLeaf",
        "mgmt",
        "sampleLeaf",
        "dataSource",
        "proposalVersion",
    ],
    "9": ["stack", "cat", "act", "gts", "hw", "partner", "customer", "proposalVersion"],
    "10": [
        "name",
        "start",
        "end",
        "actualStart",
        "actualEnd",
        "owner",
        "unit",
        "status",
        "progress",
        "progressTone",
        "proposalVersion",
    ],
    "11": [
        "cat",
        "scheme",
        "standard",
        "milestone",
        "doc",
        "payment",
        "paymentMilestone",
        "proposalVersion",
    ],
    "12": [
        "id",
        "l1",
        "l2",
        "l3",
        "purpose",
        "topology",
        "pre",
        "steps",
        "expects",
        "remark",
        "result",
        "selected",
        "proposalVersion",
    ],
}


def _label(key: str) -> str:
    return _FIELD_LABELS.get(key, key)


def _write_kv_sheet(ws, pairs: list[tuple[str, Any]]) -> None:
    ws.append(["字段", "内容"])
    for k, v in pairs:
        ws.append([_label(k), _cell_value(v)])


def _collect_row_keys(rows: list[dict[str, Any]]) -> list[str]:
    keys: list[str] = []
    seen: set[str] = set()
    skip = {"rowId"}
    for row in rows:
        if not isinstance(row, dict):
            continue
        for k in row:
            if k in skip or k in seen:
                continue
            seen.add(k)
            keys.append(k)
    return keys


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_rows_sheet(
    ws,
    rows: list[dict[str, Any]],
    *,
    title: str | None = None,
    template_keys: list[str] | None = None,
) -> None:
    if title:
        ws.title = title[:31]
    if not rows:
        if template_keys:
            ws.append([_label(k) for k in template_keys])
        else:
            ws.append(["（无数据）"])
        return
    keys = list(template_keys) if template_keys else _collect_row_keys(rows)
    if not keys:
        ws.append(["（无数据）"])
        return
    ws.append([_label(k) for k in keys])
    for row in rows:
        if not isinstance(row, dict):
            continue
        ws.append([_cell_value(row.get(k, "")) for k in keys])


def export_chapter_xlsx(path: Path, spec: ChapterSpec | None, payload: dict[str, Any]) -> None:
    """Write XLSX next to JSON; `path` ends with .xlsx."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    key = spec.key if spec else (payload.get("chapterKey") or "chapter")
    title = (spec.label if spec else payload.get("chapterTitle") or key)[:31]
    ws.title = title or "Sheet1"

    rows = payload.get("rows")
    fields = payload.get("fields")
    extras = payload.get("extras") or {}

    if key == "meta" or (not rows and not fields and payload.get("projectId")):
        pairs: list[tuple[str, Any]] = []
        for k, v in payload.items():
            if k in _META_SKIP or k in ("chapterKey", "chapterTitle", "rows", "fields", "extras"):
                continue
            pairs.append((k, v))
        _write_kv_sheet(ws, pairs)
        manual = payload.get("manualChangeLog") or []
        if manual:
            ws2 = wb.create_sheet("修改记录")
            _write_rows_sheet(ws2, manual, title="修改记录")
    elif isinstance(fields, dict) and fields:
        _write_kv_sheet(ws, list(fields.items()))
        if isinstance(rows, list) and rows:
            ws2 = wb.create_sheet("明细")
            _write_rows_sheet(ws2, rows)
    elif isinstance(rows, list):
        template_keys = _CHAPTER_TEMPLATE_KEYS.get(str(key))
        _write_rows_sheet(ws, rows, title=title, template_keys=template_keys)
    elif _CHAPTER_TEMPLATE_KEYS.get(str(key)):
        _write_rows_sheet(
            ws,
            [],
            title=title,
            template_keys=_CHAPTER_TEMPLATE_KEYS.get(str(key)),
        )
    elif isinstance(extras, dict) and extras:
        _write_kv_sheet(ws, list(extras.items()))
    else:
        pairs = [(k, v) for k, v in payload.items() if k not in ("chapterKey", "chapterTitle")]
        _write_kv_sheet(ws, pairs)

    wb.save(path)


def export_xlsx_for_json_file(json_path: Path) -> Path | None:
    """Read JSON and write sibling .xlsx; returns xlsx path or None if skipped."""
    if json_path.name in {"manifest.json", "version-info.json"}:
        return None
    if not json_path.exists():
        return None
    from agent.proposal.draft_store import load_json

    payload = load_json(json_path, {})
    if not isinstance(payload, dict):
        return None
    from agent.proposal.chapter_registry import CHAPTER_BY_JSON

    spec = CHAPTER_BY_JSON.get(json_path.name)
    xlsx_path = json_path.with_suffix(".xlsx")
    export_chapter_xlsx(xlsx_path, spec, payload)
    return xlsx_path


def read_rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    """Read first-sheet tabular rows from chapter xlsx."""
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    header = [str(c or "").strip() for c in rows[0]]
    if not any(header):
        return []

    result: list[dict[str, Any]] = []
    for values in rows[1:]:
        row: dict[str, Any] = {}
        non_empty = False
        for idx, label in enumerate(header):
            if not label:
                continue
            key = _KEY_BY_LABEL.get(label, label)
            value = values[idx] if idx < len(values) else None
            if value not in (None, ""):
                non_empty = True
            row[key] = value
        if non_empty:
            result.append(row)
    return result


def read_header_keys_from_xlsx(path: Path) -> list[str]:
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True, min_row=1, max_row=1)
        header_row = next(rows, None)
    finally:
        wb.close()
    if not header_row:
        return []
    keys: list[str] = []
    for cell in header_row:
        label = str(cell or "").strip()
        if not label:
            continue
        keys.append(_KEY_BY_LABEL.get(label, label))
    return keys


def chapter_template_keys(chapter_key: str) -> list[str]:
    return list(_CHAPTER_TEMPLATE_KEYS.get(str(chapter_key), []))
