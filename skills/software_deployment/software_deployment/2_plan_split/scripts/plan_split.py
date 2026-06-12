"""拆分调测计划（部署调测 Agent 规则，代码随 skill 内置）。"""

from __future__ import annotations

import copy
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from paths import load_second_to_third, skill_root_from_cwd

from . import constants as C

_TRAIN_INFER_ZH = {"推理": C.INFER, "训练": C.TRAIN, "训推": C.TRAIN_INFER}
_POD_ZH = {"单pod": C.SINGLE_POD, "多pod": C.MULTI_PODS}
_COOLING_ZH = {"风冷": "air_cooling", "液冷": C.LIQUID_COOLING}

# 作业管理 Excel 可能用泛称，按设备类型补全为 second_to_third 的 key
_ACTIVITY_BY_DEVICE: dict[str, dict[str, str]] = {
    C.DEVICE_SERVER: {
        "设备硬装初始化": C.SERVER_DEVICE_HARDWARE_INIT,
        "计算子系统部署调测": "计算子系统部署调测-智算服务器",
        "计算子系统验证": C.COMPUTE_SUBSYSTEM_TEST_SERVER,
    },
    C.DEVICE_LQ_SWITCH: {
        "设备硬装初始化": "设备硬装初始化-灵衢网络",
        "计算子系统部署调测": "计算子系统部署调测-灵衢网络",
        "计算子系统验证": "计算子系统验证-灵衢网络",
    },
    C.DEVICE_STORAGE: {
        "设备硬装初始化": "设备硬装初始化-存储服务器",
        "存储子系统部署调测": "存储子系统部署调测",
        "存储子系统验证": "存储子系统验证",
    },
    C.DEVICE_U_SERVER: {
        "设备硬装初始化": "设备硬装初始化-通算服务器",
        "计算子系统部署调测": "计算子系统部署调测",
        "计算子系统验证": "计算子系统验证",
    },
    C.DEVICE_SWITCH: {
        "基础网络配置": "基础网络配置",
        "网络子系统部署调测": "网络子系统部署调测",
        "网络子系统验证": "网络子系统验证",
    },
}


def _norm_enum(raw: str, zh_map: dict[str, str], allowed: set[str], default: str) -> str:
    val = str(raw or "").strip()
    if val in allowed:
        return val
    zh = zh_map.get(val) or zh_map.get(val.lower())
    if zh:
        return zh
    low = val.lower().replace("-", "_")
    if low in allowed:
        return low
    return default


def _scene_fields(scene: dict[str, Any]) -> dict[str, str]:
    cooling_raw = scene.get("cooling") or scene.get("coolingScene") or "air_cooling"
    if str(cooling_raw).strip() in _COOLING_ZH:
        cooling = _COOLING_ZH[str(cooling_raw).strip()]
    else:
        cooling = str(cooling_raw or "air_cooling")

    ti_raw = scene.get("train_infer_scene") or scene.get("trainInferScene") or "infer"
    train_infer = _norm_enum(
        str(ti_raw),
        _TRAIN_INFER_ZH,
        {C.TRAIN, C.INFER, C.TRAIN_INFER},
        C.INFER,
    )

    pod_raw = scene.get("pod_info") or scene.get("podInfo") or scene.get("pod_scene") or C.SINGLE_POD
    pod_info = _norm_enum(
        str(pod_raw),
        _POD_ZH,
        {C.SINGLE_POD, C.MULTI_PODS},
        C.SINGLE_POD,
    )

    return {
        "cooling": cooling,
        "product_specification": str(scene.get("productSpecification") or scene.get("product_spec") or ""),
        "train_infer_scene": train_infer,
        "pod_info": pod_info,
    }


def generate_task_map(scene: dict[str, Any], *, skill_root: Path | None = None) -> dict[str, list[str]]:
    """CPCIA ``ThirdTaskHandler.generate_task_map``，基础表来自 ``data/rules/second_to_third.json``。"""
    _ = skill_root
    fields = _scene_fields(scene)
    task_map: dict[str, list[str]] = copy.deepcopy(load_second_to_third(skill_root))
    if fields["cooling"] == C.LIQUID_COOLING:
        task_map.setdefault(C.SERVER_DEVICE_HARDWARE_INIT, []).extend(C.server_liquid_cooling)
    if fields["train_infer_scene"] == C.TRAIN:
        task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).extend(C.SERVER_SINGLE_TRAIN)
        task_map.setdefault(C.CLUSTER_SYSTEM_INTEGRATION_TEST, []).extend(C.SERVER_CLUSTER_TRAIN)
        if fields["pod_info"] == C.MULTI_PODS:
            task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).append(C.CLUSTER_MODELS_TEST_SINGLE_POD)
    elif fields["train_infer_scene"] == C.INFER:
        task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).extend(C.SERVER_SINGLE_INFER)
        task_map.setdefault(C.CLUSTER_SYSTEM_INTEGRATION_TEST, []).extend(C.SERVER_CLUSTER_INFER)
        if fields["pod_info"] == C.MULTI_PODS:
            task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).append(C.CLUSTER_INFER_SINGLE_POD)
    elif fields["train_infer_scene"] == C.TRAIN_INFER:
        task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).extend(C.SERVER_SINGLE_TRAIN_INFER)
        task_map.setdefault(C.CLUSTER_SYSTEM_INTEGRATION_TEST, []).extend(C.SERVER_CLUSTER_TRAIN_INFER)
        if fields["pod_info"] == C.MULTI_PODS:
            task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).append(C.CLUSTER_MODELS_TEST_SINGLE_POD)
            task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).append(C.CLUSTER_INFER_SINGLE_POD)
    if fields["pod_info"] == C.MULTI_PODS:
        task_map.setdefault(C.COMPUTE_SUBSYSTEM_TEST_SERVER, []).append(C.HCCL_TEST_SINGLE_POD)
    return task_map


def _device_list_stub(device_type: str) -> str:
    m = {
        "SERVER": "Atlas 900 A3 SuperPoD",
        "LQ_SWITCH": "LingQu",
        "STORAGE": "OceanDisk1600T",
        "SWITCH": "CE6865E",
        "U_SERVER": "U_SERVER",
        C.DEVICE_SERVER: "Atlas 900 A3 SuperPoD",
        C.DEVICE_LQ_SWITCH: "LingQu 630 V1",
        C.DEVICE_SWITCH: "CE8875",
        C.DEVICE_STORAGE: "OceanDisk1600T",
        C.DEVICE_U_SERVER: "U_SERVER",
    }
    return m.get(device_type, "Atlas 900 A3 SuperPoD")


def get_device_type_from_model(model: str) -> str:
    """CPCIA ``DeviceModel.get_device_type``。"""
    text = str(model or "")
    token = text.split(" ")[0]
    if "Atlas" in text or "智算" in text:
        return C.DEVICE_SERVER
    if "Ocean" in text or "存储" in text:
        return C.DEVICE_STORAGE
    if "LingQu" in text or "灵衢" in text or "LQS" in text:
        return C.DEVICE_LQ_SWITCH
    if "CE" in token or "Eudemon" in token or "XH" in token or "S5736" in text:
        return C.DEVICE_SWITCH
    return C.DEVICE_U_SERVER


def _parse_device_quantity_string(text: str) -> list[dict[str, Any]]:
    """解析 ``Atlas 900 A3 48台, CE8875 16台`` 形式。"""
    out: list[dict[str, Any]] = []
    for part in re.split(r"[,，;；]", str(text or "")):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(.+?)\s*(\d+)\s*(台|个|套)?$", part)
        if m:
            out.append({"device_model": m.group(1).strip(), "quantity": int(m.group(2)), "unit": m.group(3) or "台"})
        else:
            out.append({"device_model": part, "quantity": 1, "unit": "台"})
    return out


def _format_device_quantity_list(device_list: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for item in device_list:
        model = str(item.get("device_model") or item.get("model") or "").strip()
        qty = item.get("quantity", 0)
        unit = str(item.get("unit") or "台")
        if model:
            parts.append(f"{model} {qty}{unit}")
    return ", ".join(parts)


def _extract_raw_equipment_list(row: dict[str, Any], device_text: str) -> dict[str, Any]:
    raw = row.get("raw_equipment_list") or row.get("rawEquipmentList")
    if isinstance(raw, str) and raw.strip():
        try:
            raw = json.loads(raw)
        except Exception:
            raw = {}
    if isinstance(raw, dict) and raw.get("device_list"):
        return raw

    device_list = _parse_device_quantity_string(device_text)
    if device_list:
        return {"device_list": device_list}
    return {"device_list": []}


def _device_quantity(raw_equipment_list: dict[str, Any]) -> int:
    device_list = raw_equipment_list.get("device_list") if isinstance(raw_equipment_list, dict) else []
    if not isinstance(device_list, list):
        return 0
    return sum(int(item.get("quantity") or 0) for item in device_list if isinstance(item, dict))


def _has_device_list(second: dict[str, Any]) -> bool:
    raw = second.get("raw_equipment_list") or {}
    device_list = raw.get("device_list") if isinstance(raw, dict) else []
    if isinstance(device_list, list) and device_list:
        return _device_quantity(raw) > 0 or any(
            str(item.get("device_model") or item.get("model") or "").strip()
            for item in device_list
            if isinstance(item, dict)
        )
    dtql = str(second.get("device_type_quantity_list") or "").strip()
    return bool(dtql)


def resolve_activity_name(name: str, device_list_str: str, task_map: dict[str, list[str]]) -> str:
    """将泛称二级活动补全为 ``second_to_third.json`` 中的 key。"""
    name = str(name or "").strip()
    if not name:
        return name
    if name in task_map:
        return name
    dt = get_device_type_from_model(device_list_str)
    mapped = _ACTIVITY_BY_DEVICE.get(dt, {}).get(name)
    if mapped and mapped in task_map:
        return mapped
    return name


def should_split_secondary_task(
    second: dict[str, Any],
    activity_name: str,
    third_activities: list[str],
) -> bool:
    """CPCIA ``ThirdTaskHandler.should_split_secondary_task``。"""
    if (
        not _has_device_list(second)
        and activity_name not in C.NO_BELONG_DEVICE_TASKS
    ):
        return False
    if not third_activities:
        return False
    return True


def parse_pod_id(raw: Any) -> int | None:
    """将 PoD/POD01/1/管理单元后缀等解析为整数 pod_id。"""
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        n = int(raw)
        return n if n >= 0 else None
    s = str(raw).strip()
    if not s:
        return None
    if s.isdigit():
        return int(s)
    m = re.search(r"POD\s*0*(\d+)", s, re.IGNORECASE)
    if m:
        return int(m.group(1))
    m2 = re.search(r"(\d+)$", s)
    if m2:
        return int(m2.group(1))
    return None


def pod_id_from_management_unit(management_unit: str) -> int | None:
    """Agent 默认规则：管理单元最后一个横杠后的段作为实际 POD 编号。"""
    mu = str(management_unit or "").strip()
    if not mu:
        return None
    suffix = mu.rsplit("-", 1)[-1] if "-" in mu else mu
    return parse_pod_id(suffix)


def _accumulate_pod_map_row(
    out: dict[str, set[int]],
    *,
    management_unit: str,
    pod_raw: Any = None,
) -> None:
    mu = str(management_unit or "").strip()
    if not mu:
        return
    pod_int = parse_pod_id(pod_raw) if pod_raw is not None and str(pod_raw).strip() else None
    if pod_int is None:
        pod_int = pod_id_from_management_unit(mu)
    if pod_int is None:
        return
    out.setdefault(mu, set()).add(pod_int)


def load_pod_map(skill_root: Path | None = None) -> dict[str, list[int]]:
    """读取 ``pod_map`` 槽位 xlsx/json：management_unit → [pod_id, ...]。

    支持两种表结构：
    - 专用映射表：``管理单元`` + ``pod_id`` / ``实际POD编号``
    - 到货表（多行重复管理单元）：``管理单元`` + ``PoD`` 列，按管理单元去重聚合
    """
    root = skill_root or skill_root_from_cwd()
    rt = root / "runtime"
    import sys

    if str(rt) not in sys.path:
        sys.path.insert(0, str(rt))
    from paths import resolve_slot  # noqa: WPS433

    slot = resolve_slot("pod_map", skill_root=root)
    primary = str(slot.get("primary") or "").strip()
    if not primary:
        return {}
    path = Path(primary)
    if not path.is_file():
        return {}

    out: dict[str, list[int]] = {}
    if path.suffix.lower() == ".json":
        raw = json.loads(path.read_text(encoding="utf-8"))
        rows = raw if isinstance(raw, list) else raw.get("rows") or raw.get("map") or []
        if isinstance(rows, dict):
            for k, v in rows.items():
                out[str(k).strip()] = [int(v)] if isinstance(v, (int, str)) and str(v).isdigit() else list(v)
        elif isinstance(rows, list):
            tmp: dict[str, set[int]] = {}
            for row in rows:
                if not isinstance(row, dict):
                    continue
                mu = str(row.get("management_unit") or row.get("managementUnit") or row.get("管理单元") or "").strip()
                pod = (
                    row.get("pod_id")
                    or row.get("podId")
                    or row.get("real_pod")
                    or row.get("POD_ID")
                    or row.get("PoD")
                    or row.get("POD")
                    or row.get("实际POD编号")
                )
                _accumulate_pod_map_row(tmp, management_unit=mu, pod_raw=pod)
            out = {k: sorted(v) for k, v in tmp.items()}
        return out

    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return {}

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return out
        header = [str(c or "").strip() for c in rows[0]]
        norm = [h.replace(" ", "").lower() for h in header]

        def _idx(*names: str) -> int | None:
            for name in names:
                key = name.replace(" ", "").lower()
                for i, h in enumerate(norm):
                    if h == key or key in h:
                        return i
            return None

        i_mu = _idx("管理单元", "management_unit", "managementUnit")
        i_pod = _idx(
            "pod_id",
            "podId",
            "real_pod",
            "POD_ID",
            "Pod",
            "PoD",
            "POD",
            "实际POD编号",
            "实际pod编号",
        )
        if i_mu is None:
            return out
        tmp: dict[str, set[int]] = {}
        for row in rows[1:]:
            cells = list(row or [])
            mu = str(cells[i_mu] if i_mu < len(cells) else "").strip()
            pod_raw = cells[i_pod] if i_pod is not None and i_pod < len(cells) else None
            _accumulate_pod_map_row(tmp, management_unit=mu, pod_raw=pod_raw)
        return {k: sorted(v) for k, v in tmp.items()}
    finally:
        wb.close()


def resolve_pod_ids(
    management_unit: str | list[str],
    *,
    scene: dict[str, Any],
    pod_map: dict[str, list[int]] | None = None,
) -> list[int]:
    fields = _scene_fields(scene)
    units: list[str] = []
    if isinstance(management_unit, list):
        units = [str(u).strip() for u in management_unit if str(u).strip()]
    elif management_unit:
        units = [p.strip() for p in re.split(r"[,，]", str(management_unit)) if p.strip()]

    if fields["pod_info"] == C.MULTI_PODS and units:
        pod_ids: list[int] = []
        for unit in units:
            mapped = (pod_map or {}).get(unit)
            if mapped:
                pod_ids.extend(mapped)
                continue
            parsed = pod_id_from_management_unit(unit)
            if parsed is not None:
                pod_ids.append(parsed)
                continue
            try:
                pod_ids.append(int(unit))
            except ValueError:
                pass
        if pod_ids:
            return sorted(set(pod_ids))

    return [0]


def second_row_to_internal(row: dict[str, Any], *, project_id: str) -> dict[str, Any]:
    sn = str(row.get("id") or row.get("serial_number") or "").strip() or "L2-0001"
    device_text = str(
        row.get("device_type_quantity_list")
        or row.get("deviceList")
        or row.get("device_list")
        or row.get("deviceTypeQuantityList")
        or ""
    ).strip()
    raw_equipment_list = _extract_raw_equipment_list(row, device_text)
    if not device_text and raw_equipment_list.get("device_list"):
        device_text = _format_device_quantity_list(raw_equipment_list["device_list"])
    if not device_text:
        device_text = _device_list_stub(str(row.get("deviceType") or get_device_type_from_model("")))

    mu = row.get("managementUnit") or row.get("management_unit") or ""
    if isinstance(mu, list):
        management_unit = ",".join(str(x).strip() for x in mu if str(x).strip())
    else:
        management_unit = str(mu or "").strip()

    return {
        "project_id": project_id,
        "serial_number": sn,
        "activity_name": str(row.get("activityName") or row.get("name") or row.get("activity_name") or "").strip(),
        "management_unit": management_unit,
        "start_date": row.get("startDate") or row.get("start_date") or "",
        "end_date": row.get("endDate") or row.get("end_date") or "",
        "principal": row.get("principal") or "",
        "raw_equipment_list": raw_equipment_list,
        "device_type_quantity_list": device_text,
    }


def build_third_task(
    second: dict[str, Any],
    third_name: str,
    *,
    scene: dict[str, Any],
    project_id: str,
    pod_map: dict[str, list[int]] | None = None,
) -> dict[str, Any]:
    """CPCIA ``generate_third_level_activity`` 字段构造。"""
    second_sn = second["serial_number"]
    pod_ids = resolve_pod_ids(second.get("management_unit") or "", scene=scene, pod_map=pod_map)
    serial_number = hashlib.md5(f"{third_name}{second_sn}".encode("utf-8")).hexdigest()
    principal = "远程工程师" if third_name in C.AGENT_ACTIVITIES else "现场测试工程师"
    task_type = "handwork" if "现场" in principal else "agent"
    device_list = second.get("device_type_quantity_list", "")
    device_quantity = _device_quantity(second.get("raw_equipment_list") or {})
    return {
        "project_id": project_id,
        "serial_number": serial_number,
        "start_date": second.get("start_date"),
        "end_date": second.get("end_date"),
        "device_list": device_list,
        "device_quantity": device_quantity if device_quantity else "",
        "finish_num": 0 if device_quantity else "",
        "second_task_sn": second_sn,
        "second_instruction": second.get("activity_name"),
        "third_instruction": third_name,
        "third_task_status": C.TASK_STATUS_PENDING_DISPATCHED,
        "update_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "principal": principal,
        "task_type": task_type,
        "pod_ids": json.dumps(pod_ids),
        "management_unit": second.get("management_unit") or "",
    }


def third_internal_to_nanobot(params: dict[str, Any], *, device_type: str) -> dict[str, Any]:
    try:
        pod_ids = json.loads(params.get("pod_ids") or "[0]")
    except Exception:
        pod_ids = [0]
    dq = params.get("device_quantity")
    return {
        "id": params.get("serial_number"),
        "secondId": params.get("second_task_sn"),
        "secondActivityName": params.get("second_instruction"),
        "managementUnit": params.get("management_unit") or "",
        "thirdActivityName": params.get("third_instruction"),
        "principal": params.get("principal"),
        "status": "待派发",
        "taskType": params.get("task_type") or "agent",
        "deviceType": device_type,
        "deviceList": params.get("device_list") or "",
        "deviceQuantity": dq if dq != "" else None,
        "finishNum": params.get("finish_num") if dq != "" else None,
        "podIds": pod_ids if isinstance(pod_ids, list) else [0],
        "startDate": params.get("start_date"),
        "endDate": params.get("end_date"),
    }


def split_plan(
    second_tasks: list[dict[str, Any]],
    *,
    scene: dict[str, Any],
    project_id: str = "nanobot-local",
    skill_root: Path | None = None,
) -> list[dict[str, Any]]:
    root = skill_root or skill_root_from_cwd()
    task_map = generate_task_map(scene, skill_root=root)
    pod_map = load_pod_map(root)
    out: list[dict[str, Any]] = []
    for row in second_tasks:
        internal = second_row_to_internal(row, project_id=project_id)
        activity = resolve_activity_name(
            internal.get("activity_name") or "",
            internal.get("device_type_quantity_list") or "",
            task_map,
        )
        internal["activity_name"] = activity
        third_names = task_map.get(activity, [])
        if not should_split_secondary_task(internal, activity, third_names):
            continue
        for third_name in third_names:
            params = build_third_task(
                internal,
                third_name,
                scene=scene,
                project_id=project_id,
                pod_map=pod_map,
            )
            dt = get_device_type_from_model(params.get("device_list") or "")
            out.append(third_internal_to_nanobot(params, device_type=dt))
    return out


def split_plan_report(
    second_tasks: list[dict[str, Any]],
    *,
    scene: dict[str, Any],
    project_id: str = "nanobot-local",
    skill_root: Path | None = None,
) -> dict[str, Any]:
    """拆分统计：便于 UI 展示与 Agent 结果对账。"""
    root = skill_root or skill_root_from_cwd()
    task_map = generate_task_map(scene, skill_root=root)
    pod_map = load_pod_map(root)
    skipped_no_device: list[str] = []
    skipped_no_map: list[str] = []
    third_count_by_second: dict[str, int] = {}
    split_rows: list[dict[str, Any]] = []
    skipped_second = 0

    for row in second_tasks:
        internal = second_row_to_internal(row, project_id=project_id)
        activity = resolve_activity_name(
            internal.get("activity_name") or "",
            internal.get("device_type_quantity_list") or "",
            task_map,
        )
        internal["activity_name"] = activity
        third_names = task_map.get(activity, [])
        if activity and activity not in task_map:
            skipped_no_map.append(activity)
        if not should_split_secondary_task(internal, activity, third_names):
            if activity and activity not in C.NO_BELONG_DEVICE_TASKS and not _has_device_list(internal):
                skipped_no_device.append(activity)
            skipped_second += 1
            continue
        third_count_by_second[activity] = third_count_by_second.get(activity, 0) + len(third_names)
        for third_name in third_names:
            params = build_third_task(
                internal,
                third_name,
                scene=scene,
                project_id=project_id,
                pod_map=pod_map,
            )
            dt = get_device_type_from_model(params.get("device_list") or "")
            split_rows.append(third_internal_to_nanobot(params, device_type=dt))

    fields = _scene_fields(scene)
    tree = build_plan_tree(second_tasks, split_rows, scene=scene, skill_root=root)
    return {
        "thirdCount": len(split_rows),
        "secondCount": len(second_tasks),
        "skippedSecondCount": skipped_second,
        "scene": fields,
        "thirdCountBySecondActivity": third_count_by_second,
        "skippedNoDevice": sorted(set(skipped_no_device)),
        "unmappedSecondActivities": sorted(set(skipped_no_map)),
        "podMapLoaded": bool(pod_map),
        "tasks": split_rows,
        "tree": tree,
    }


def build_plan_tree(
    second_tasks: list[dict[str, Any]],
    third_tasks: list[dict[str, Any]],
    *,
    scene: dict[str, Any] | None = None,
    skill_root: Path | None = None,
) -> list[dict[str, Any]]:
    """对齐 Agent ``format_task_info_list``：二级父行 + ``subActivities`` 三级子行。"""
    scene = scene or {}
    root = skill_root or skill_root_from_cwd()
    pod_map = load_pod_map(root)
    third_by_second: dict[str, list[dict[str, Any]]] = {}
    for t in third_tasks:
        sid = str(t.get("secondId") or "").strip()
        if sid:
            third_by_second.setdefault(sid, []).append(t)

    tree: list[dict[str, Any]] = []
    for row in second_tasks:
        sid = str(row.get("id") or row.get("serial_number") or "").strip()
        internal = second_row_to_internal(row, project_id="")
        pod_ids = resolve_pod_ids(internal.get("management_unit") or "", scene=scene, pod_map=pod_map)
        parent = {
            "level": 2,
            "serialNumber": sid,
            "activityName": internal.get("activity_name") or "",
            "startDate": internal.get("start_date") or "",
            "endDate": internal.get("end_date") or "",
            "actualStartDate": row.get("actualStartDate") or row.get("actual_start_date") or "--",
            "actualEndDate": row.get("actualEndDate") or row.get("actual_end_date") or "--",
            "podIds": pod_ids,
            "deviceList": internal.get("device_type_quantity_list") or "",
            "taskType": "",
            "principal": row.get("principal") or internal.get("principal") or "",
            "status": row.get("status") or "已接收",
            "deviceQuantity": None,
            "finishNum": None,
            "progress": None,
            "subActivities": [],
        }
        for t in third_by_second.get(sid, []):
            parent["subActivities"].append(
                {
                    "level": 3,
                    "serialNumber": t.get("id"),
                    "activityName": t.get("thirdActivityName"),
                    "secondId": sid,
                    "secondActivityName": t.get("secondActivityName"),
                    "startDate": t.get("startDate"),
                    "endDate": t.get("endDate"),
                    "actualStartDate": "--",
                    "actualEndDate": "--",
                    "podIds": [],
                    "deviceList": "",
                    "taskType": t.get("taskType"),
                    "principal": t.get("principal"),
                    "status": t.get("status"),
                    "deviceQuantity": t.get("deviceQuantity"),
                    "finishNum": t.get("finishNum"),
                    "progress": 0 if t.get("finishNum") is not None else None,
                }
            )
        tree.append(parent)
    return tree


def flatten_plan_tree(tree: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """拍平为 Agent UI 表格行（父行 + 子行）。"""
    rows: list[dict[str, Any]] = []
    for parent in tree:
        rows.append(
            {
                "level": 2,
                "activityName": parent.get("activityName"),
                "startDate": parent.get("startDate"),
                "endDate": parent.get("endDate"),
                "actualStartDate": parent.get("actualStartDate"),
                "actualEndDate": parent.get("actualEndDate"),
                "podIds": parent.get("podIds"),
                "deviceList": parent.get("deviceList"),
                "taskType": parent.get("taskType"),
                "principal": parent.get("principal"),
                "status": parent.get("status"),
                "deviceQuantity": parent.get("deviceQuantity"),
                "finishNum": parent.get("finishNum"),
                "progress": parent.get("progress"),
            }
        )
        for sub in parent.get("subActivities") or []:
            rows.append(dict(sub))
    return rows
