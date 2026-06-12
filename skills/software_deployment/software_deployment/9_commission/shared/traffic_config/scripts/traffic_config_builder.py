# -*- coding: utf-8 -*-
"""打流测试 trafficTestConfigs 运行时生成（对齐 Agent TrafficTestConfigGenerator）。

Agent 不从 execute.json 静态读取 configs，而是：
  1. 读参数页签「灵衢总线打流测试」（测试类型 / 芯片类型 / 是否跨节点 / 端口 / 数据大小 / 迭代次数）
  2. 按 POD 分组 deviceIds
  3. 用 CloudOps《服务器信息》BMC ↔ 管理网 IP 映射填 sourceIp/targetIp
  4. 将 flow 参数写入每条 config 的 port / dataSize / cycleNum（对齐 Agent TrafficTestParamsSet）
"""
from __future__ import annotations

import copy
import json
from pathlib import Path
from typing import Any

_CONFIG_TEMPLATE: dict[str, Any] = {
    "port": "",
    "testMethod": "portTrafficTest",
    "sourceIp": "",
    "sourceChipType": "npu",
    "sourceChipId": "0",
    "sourceDie": "0",
    "targetIp": "",
    "targetChipType": "npu",
    "targetChipId": "0",
    "targetDie": "0",
    "dataSize": "",
    "cycleNum": "",
    "sourceDeviceId": "",
    "targetDeviceId": "",
    "testType": "HCCS",
}

_LABEL_MAP = {
    "NPU - NPU": "npu",
    "CPU - NPU": "cpu",
    "是": "cross_node",
    "否": "intra_node",
}

# Agent TrafficTestParamsSet 字段；CloudOps CreateTask 校验：
#   - dataSize 带 M/G 后缀会 9003；HCCS 场景留空由执行机默认
#   - cycleNum 可填数字字符串（如 40）
_AGENT_FLOW_DEFAULTS: dict[str, str] = {
    "port": "",
    "dataSize": "",
    "cycleNum": "40",
}


def _skill_root(skill_dir: str | Path) -> Path:
    return Path(skill_dir).resolve()


def _load_json(path: Path) -> Any:
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _product_specification(skill_dir: str | Path) -> str:
    scene = _load_json(_skill_root(skill_dir) / "ProjectData" / "plan" / "RunTime" / "scene.json")
    if isinstance(scene, dict):
        ps = str(scene.get("productSpecification") or scene.get("product_spec") or "").strip()
        if ps in ("A2", "A3"):
            return ps
    return "A3"


def _load_traffic_sheet_defaults(skill_dir: str | Path) -> dict[str, str]:
    """读 runtime/params_template_sheets.json 中「灵衢总线打流测试」首行有效数据。"""
    root = _skill_root(skill_dir)
    data = _load_json(root / "runtime" / "params_template_sheets.json")
    rows = data.get("灵衢总线打流测试") if isinstance(data, dict) else None
    base = {
        "test_type": "HCCS",
        "chip_type": "npu",
        "cross_node": "cross_node",
        **_AGENT_FLOW_DEFAULTS,
    }
    if not isinstance(rows, list) or len(rows) < 2:
        return base
    header = [str(c or "").strip() for c in rows[0]]
    values = rows[1]
    row = {header[i]: str(values[i] or "").strip() for i in range(min(len(header), len(values)))}
    test_type = row.get("测试类型", "HCCS") or "HCCS"
    chip_type = _LABEL_MAP.get(row.get("芯片类型", "NPU - NPU"), "npu")
    cross = _LABEL_MAP.get(row.get("是否跨节点", "是"), "cross_node")
    if test_type.upper() == "ROCE":
        chip_type, cross = "npu", "cross_node"
    elif test_type.upper() == "HCCS" and chip_type == "cpu":
        cross = "intra_node"
    flow = dict(_AGENT_FLOW_DEFAULTS)
    for cn, key in (("端口", "port"), ("数据大小", "dataSize"), ("迭代次数", "cycleNum")):
        val = row.get(cn, "").strip()
        if val:
            flow[key] = val
    return {
        "test_type": test_type.upper(),
        "chip_type": chip_type,
        "cross_node": cross,
        **flow,
    }


def _normalize_data_size(val: str) -> str:
    """CloudOps CreateTask 不接受 256M/512M 后缀，仅保留数值部分。"""
    v = (val or "").strip()
    if not v:
        return ""
    upper = v.upper()
    for suffix in ("GB", "MB", "KB", "G", "M", "K"):
        if upper.endswith(suffix):
            return v[: -len(suffix)].strip()
    return v


def _apply_flow_params(tmpl: dict[str, Any], params: dict[str, str]) -> dict[str, Any]:
    row = copy.deepcopy(tmpl)
    for key in ("port", "dataSize", "cycleNum"):
        val = params.get(key, _AGENT_FLOW_DEFAULTS.get(key, ""))
        if key == "dataSize":
            val = _normalize_data_size(str(val))
        row[key] = val
    return row


def load_bmc_to_management_map(skill_dir: str | Path) -> dict[str, str]:
    """BMC/设备ID → 管理网 IPv4（与 Agent load_lld_bmc_map 一致，数据源为 CloudOps 完整配置）。"""
    full = (
        _skill_root(skill_dir)
        / "ProjectData"
        / "plan"
        / "Output"
        / "CloudOps完整配置文件.xlsx"
    )
    if not full.is_file():
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}
    mapping: dict[str, str] = {}
    wb = load_workbook(str(full), read_only=True, data_only=True)
    try:
        sheet = "服务器信息" if "服务器信息" in wb.sheetnames else None
        if not sheet:
            return {}
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {}
        header = [str(c or "").strip() for c in rows[0]]
        idx = {n: i for i, n in enumerate(header) if n}

        def col(*names: str) -> int:
            for n in names:
                if n in idx:
                    return idx[n]
            return -1

        bmc_i = col("设备ID*（主键）", "设备ID")
        mgmt_i = col("管理网-IPV4地址", "管理网IP", "管理网-IP地址")
        if bmc_i < 0:
            bmc_i = 0
        for r in rows[1:]:
            vals = list(r)
            if bmc_i >= len(vals):
                continue
            bmc = str(vals[bmc_i] or "").strip()
            if not bmc:
                continue
            mgmt = ""
            if 0 <= mgmt_i < len(vals):
                mgmt = str(vals[mgmt_i] or "").strip()
            mapping[bmc] = mgmt or bmc
    finally:
        wb.close()
    return mapping


def _mgmt_ip(bmc_ip: str, bmc_map: dict[str, str]) -> str:
    return bmc_map.get(bmc_ip) or bmc_ip


def _traverse_npu(tmpl: dict[str, Any], die: list[int], out: list[dict[str, Any]]) -> None:
    for npu_id in range(8):
        row = copy.deepcopy(tmpl)
        row["sourceChipId"] = str(npu_id)
        row["targetChipId"] = str(npu_id)
        for die_id in die:
            row["sourceDie"] = str(die_id)
            row["targetDie"] = str(die_id)
            out.append(copy.deepcopy(row))


def _generate_cross_node_npu(
    ips: list[str], bmc_map: dict[str, str], test_type: str, flow_params: dict[str, str]
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    if len(ips) < 2:
        return out
    tmpl = _apply_flow_params(_CONFIG_TEMPLATE, flow_params)
    tmpl["testType"] = test_type
    tmpl["sourceChipType"] = tmpl["targetChipType"] = "npu"
    pairs = [(ips[i], ips[len(ips) - 1 - i]) for i in range(len(ips) // 2)]
    die = [0, 1]
    for src_bmc, dst_bmc in pairs:
        tmpl["sourceIp"] = _mgmt_ip(src_bmc, bmc_map)
        tmpl["sourceDeviceId"] = src_bmc
        tmpl["targetIp"] = _mgmt_ip(dst_bmc, bmc_map)
        tmpl["targetDeviceId"] = dst_bmc
        _traverse_npu(tmpl, die, out)
    return out


def _generate_intra_node_npu(
    ips: list[str], bmc_map: dict[str, str], test_type: str, product: str, flow_params: dict[str, str]
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    tmpl = _apply_flow_params(_CONFIG_TEMPLATE, flow_params)
    tmpl["testType"] = test_type
    tmpl["sourceChipType"] = tmpl["targetChipType"] = "npu"
    for ip in ips:
        tmpl["sourceIp"] = tmpl["targetIp"] = _mgmt_ip(ip, bmc_map)
        tmpl["sourceDeviceId"] = tmpl["targetDeviceId"] = ip
        if product == "A2":
            for npu_src, npu_dst in ((0, 7), (1, 6), (2, 5), (3, 4)):
                row = copy.deepcopy(tmpl)
                row["sourceChipId"] = str(npu_src)
                row["targetChipId"] = str(npu_dst)
                row["sourceDie"] = row["targetDie"] = "0"
                out.append(row)
        else:
            for npu_id in range(8):
                row = copy.deepcopy(tmpl)
                row["sourceChipId"] = str(npu_id)
                row["targetChipId"] = str(npu_id)
                row["sourceDie"] = "0"
                row["targetDie"] = "1"
                out.append(row)
    return out


def _generate_intra_node_cpu(
    ips: list[str], bmc_map: dict[str, str], test_type: str, product: str, flow_params: dict[str, str]
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    tmpl = _apply_flow_params(_CONFIG_TEMPLATE, flow_params)
    tmpl["testType"] = test_type
    tmpl["sourceChipType"] = "cpu"
    tmpl["targetChipType"] = "npu"
    die = [0] if product == "A2" else [0, 1]
    for ip in ips:
        tmpl["sourceIp"] = tmpl["targetIp"] = _mgmt_ip(ip, bmc_map)
        tmpl["sourceDeviceId"] = tmpl["targetDeviceId"] = ip
        for npu_id in range(8):
            row = copy.deepcopy(tmpl)
            row["targetChipId"] = str(npu_id)
            for die_id in die:
                row["targetDie"] = str(die_id)
                out.append(copy.deepcopy(row))
    return out


def build_traffic_test_configs(
    skill_dir: str | Path,
    *,
    devices: list[Any],
    ips: list[str],
) -> list[dict[str, Any]]:
    """按 POD 分组生成 trafficTestConfigs。"""
    params = _load_traffic_sheet_defaults(skill_dir)
    product = _product_specification(skill_dir)
    bmc_map = load_bmc_to_management_map(skill_dir)
    test_type = params["test_type"]
    chip = params["chip_type"]
    mode = params["cross_node"]
    flow_params = {k: params[k] for k in ("port", "dataSize", "cycleNum")}

    ip_set = set(ips)
    by_pod: dict[int, list[str]] = {}
    for d in devices:
        if getattr(d, "ip", None) not in ip_set:
            continue
        pod = int(getattr(d, "pod_id", 0) or 0)
        by_pod.setdefault(pod, []).append(d.ip)

    configs: list[dict[str, Any]] = []
    gen_name = f"generate_{mode}_{chip}_config"
    for _pod, pod_ips in sorted(by_pod.items()):
        pod_ips = sorted(set(pod_ips))
        if gen_name == "generate_cross_node_npu_config":
            configs.extend(_generate_cross_node_npu(pod_ips, bmc_map, test_type, flow_params))
        elif gen_name == "generate_intra_node_npu_config":
            configs.extend(_generate_intra_node_npu(pod_ips, bmc_map, test_type, product, flow_params))
        elif gen_name == "generate_intra_node_cpu_config":
            configs.extend(_generate_intra_node_cpu(pod_ips, bmc_map, test_type, product, flow_params))
    return configs


def enrich_traffic_execute_body(
    skill_dir: str | Path,
    body: dict[str, Any],
    *,
    devices: list[Any],
    ips: list[str],
) -> dict[str, Any]:
    if "trafficTestConfigs" not in body:
        return body
    existing = body.get("trafficTestConfigs")
    if isinstance(existing, list) and existing:
        return body
    body["trafficTestConfigs"] = build_traffic_test_configs(skill_dir, devices=devices, ips=ips)
    return body
