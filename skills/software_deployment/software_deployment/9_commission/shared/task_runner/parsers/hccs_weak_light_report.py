# -*- coding: utf-8 -*-
"""灵衢光链路检查 · 解析 report.zip 中「灵衢光链路检查报告」xlsx → result.json 结构。"""
from __future__ import annotations

import io
import zipfile
from pathlib import PurePosixPath
from typing import Any

REPORT_PREFIX = "灵衢光链路检查报告"
SUMMARY_SHEET_COUNT = 3


def _find_report_name(names: list[str]) -> str:
    for name in names:
        base = PurePosixPath(name).name
        if base.startswith(REPORT_PREFIX) and base.endswith(".xlsx"):
            return name
    return ""


def _sheet_row_count(ws) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    data_rows = 0
    for row in rows[1:]:
        if any(cell is not None and str(cell).strip() for cell in row):
            data_rows += 1
    return data_rows


def parse_report(
    report_bytes: bytes,
    artifact: Any = None,
) -> dict[str, Any]:
    """解析 6 Sheet 报告；摘要取前 3 Sheet 行数，合并采集统计。"""
    from openpyxl import load_workbook

    collect_summary = {"total": 0, "success": 0, "failed": 0}
    devices: list[dict[str, Any]] = []
    if artifact is not None:
        cs = getattr(artifact, "collect_summary", None) or {}
        if isinstance(cs, dict):
            collect_summary = {
                "total": int(cs.get("total") or 0),
                "success": int(cs.get("success") or 0),
                "failed": int(cs.get("failed") or 0),
            }
        devs = getattr(artifact, "devices", None) or []
        if isinstance(devs, list):
            devices = list(devs)

    result: dict[str, Any] = {
        "collectSummary": collect_summary,
        "devices": devices,
        "sheetCounts": {},
        "reportFile": "",
    }

    if not report_bytes:
        return result

    report_name = ""
    xlsx_bytes = b""
    with zipfile.ZipFile(io.BytesIO(report_bytes)) as zf:
        report_name = _find_report_name(zf.namelist())
        if report_name:
            xlsx_bytes = zf.read(report_name)

    if not xlsx_bytes:
        return result

    result["reportFile"] = PurePosixPath(report_name).name

    wb = load_workbook(filename=io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
        for i, name in enumerate(sheet_names):
            if i >= SUMMARY_SHEET_COUNT:
                break
            ws = wb[name]
            result["sheetCounts"][name] = _sheet_row_count(ws)
    finally:
        wb.close()

    return result
