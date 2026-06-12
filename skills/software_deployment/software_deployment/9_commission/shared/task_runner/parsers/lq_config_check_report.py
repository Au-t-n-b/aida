# -*- coding: utf-8 -*-
"""灵衢配置检查 · 解析 xlsx Sheet「配置检查结果」→ result.json + Markdown。"""
from __future__ import annotations

import io
import math
from typing import Any

SHEET_NAME = "配置检查结果"


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    return str(v).strip()


def _num(v: Any) -> int | float:
    if v is None or v == "":
        return 0
    try:
        n = float(v)
    except Exception:
        return 0
    return int(n) if n.is_integer() else n


def _read_sheet(report_bytes: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(report_bytes), read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return [], []
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        header_idx = 0
        for idx, row in enumerate(rows):
            values = [_cell_str(c) for c in row]
            if "检查结果" in values or "IP" in values:
                header_idx = idx
                break
        columns = [_cell_str(c) for c in rows[header_idx]]
        columns = [c for c in columns if c]
        records: list[dict[str, Any]] = []
        for row in rows[header_idx + 1 :]:
            vals = list(row)
            if not any(_cell_str(v) for v in vals):
                continue
            records.append(
                {
                    columns[i]: vals[i] if i < len(vals) else ""
                    for i in range(len(columns))
                }
            )
        return columns, records
    finally:
        wb.close()


def _summary(columns: list[str], records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if "检查结果" not in columns:
        return []
    grouped: dict[str, dict[str, Any]] = {}
    sum_columns = columns[2:-2] if len(columns) > 2 else []
    if not sum_columns and len(columns) > 2:
        sum_columns = columns[2:]
    for row in records:
        key = _cell_str(row.get("检查结果")) or "未知"
        item = grouped.setdefault(key, {"检查结果": key, "IP数量": 0})
        item["IP数量"] = int(item.get("IP数量") or 0) + 1
        for col in sum_columns:
            item[col] = _num(item.get(col)) + _num(row.get(col))
    return list(grouped.values())


def _devices_from_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    devices: list[dict[str, Any]] = []
    for row in records:
        ip = _cell_str(row.get("IP") or row.get("nodeIp") or row.get("设备IP"))
        check_result = _cell_str(row.get("检查结果") or row.get("checkResult"))
        if not ip:
            continue
        devices.append(
            {
                "ip": ip,
                "checkResult": check_result,
                "result": "Fail" if check_result == "检查失败" else "Pass",
            }
        )
    return devices


def _markdown(summary: list[dict[str, Any]]) -> str:
    if not summary:
        return ""
    keys = ["检查结果", "IP数量"]
    for row in summary:
        for key in row:
            if key not in keys:
                keys.append(key)
    lines = [
        "| " + " | ".join(keys) + " |",
        "| " + " | ".join("---" for _ in keys) + " |",
    ]
    for row in summary:
        lines.append("| " + " | ".join(_cell_str(row.get(k)) for k in keys) + " |")
    return "\n".join(lines)


def parse_report(report_bytes: bytes, artifact: Any = None) -> dict[str, Any]:
    result: dict[str, Any] = {
        "summaryByResult": [],
        "devices": [],
        "markdown": "",
        "reportFile": "report.xlsx",
    }
    if not report_bytes:
        return result

    columns, records = _read_sheet(report_bytes)
    summary = _summary(columns, records)
    devices = _devices_from_records(records)

    artifact_devices = getattr(artifact, "devices", None) if artifact is not None else None
    if isinstance(artifact_devices, list) and artifact_devices:
        devices = artifact_devices

    result["summaryByResult"] = summary
    result["devices"] = devices
    result["markdown"] = _markdown(summary)
    return result
