# -*- coding: utf-8 -*-
"""灵衢健康检查 · 解析 report.zip → Sheet「灵衢交换机健康检查统计」→ result.json + Markdown。"""
from __future__ import annotations

import io
import math
import zipfile
from pathlib import PurePosixPath
from typing import Any

SHEET_NAME = "灵衢交换机健康检查统计"
REPORT_DIR = "healthCheck/statistics_reports/"


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    return str(v).strip()


def _find_report_entry(names: list[str]) -> str:
    candidates: list[str] = []
    for name in names:
        norm = name.replace("\\", "/")
        if not norm.startswith(REPORT_DIR):
            continue
        base = PurePosixPath(norm).name.lower()
        if base.endswith(".xlsx") or base.endswith(".xls"):
            candidates.append(name)
    if not candidates:
        return ""
    return sorted(candidates)[-1]


def _records_to_markdown(columns: list[str], records: list[dict[str, Any]]) -> str:
    if not columns or not records:
        return ""
    lines = [
        "| " + " | ".join(columns) + " |",
        "| " + " | ".join("---" for _ in columns) + " |",
    ]
    for row in records:
        lines.append("| " + " | ".join(_cell_str(row.get(c)) for c in columns) + " |")
    return "\n".join(lines)


def _read_statistics_sheet(xlsx_bytes: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return [], []
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        columns = [_cell_str(c) for c in rows[0]]
        columns = [c for c in columns if c]
        if not columns:
            return [], []
        records: list[dict[str, Any]] = []
        for row in rows[1:]:
            vals = list(row)
            if not any(_cell_str(v) for v in vals):
                continue
            item = {
                columns[i]: _cell_str(vals[i]) if i < len(vals) else ""
                for i in range(len(columns))
            }
            records.append(item)
        return columns, records
    finally:
        wb.close()


def parse_report(
    report_bytes: bytes,
    artifact: Any = None,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "records": [],
        "markdown": "",
        "reportFile": "",
        "devices": [],
    }

    if not report_bytes:
        return result

    report_name = ""
    xlsx_bytes = b""
    with zipfile.ZipFile(io.BytesIO(report_bytes)) as zf:
        report_name = _find_report_entry(zf.namelist())
        if report_name:
            xlsx_bytes = zf.read(report_name)

    if not xlsx_bytes:
        return result

    result["reportFile"] = report_name.replace("\\", "/")

    columns, records = _read_statistics_sheet(xlsx_bytes)
    result["records"] = records
    result["markdown"] = _records_to_markdown(columns, records)
    return result
