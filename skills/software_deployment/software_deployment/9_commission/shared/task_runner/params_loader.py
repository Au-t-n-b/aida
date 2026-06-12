# -*- coding: utf-8 -*-
"""CloudOps 测试参数 xlsx：门禁 + templateId 合并（灵衢健康检查等 TOOLKIT_TASKS）。"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

TEMPLATE_ID_MAP: dict[str, str] = {
    "全量检测": "7e9fa60790634f1ba6cdc458aa08f7dc",
    "灵衢交换机健康检查_A3": "9d6bd871fc2946bd93aabfa5dfce7be7",
    "灵衢交换机健康检查_HPC": "8dbce99d525947d5bff2bdb5ad92a352",
}

_DEFAULT_TEMPLATE_ID = "7e9fa60790634f1ba6cdc458aa08f7dc"
_TEMPLATE_COLUMN = "检查项模板"

# 测试参数 xlsx 与「设备安装完工清单」同属 cloudops 目录，须排除后者。
_CHECKLIST_MARKERS = ("完工清单", "设备安装完工清单")
_PARAMS_TEMPLATE_MARKERS = ("task_params", "测试参数", "CloudOps_task_params")


def _is_checklist_xlsx(path: Path) -> bool:
    name = path.name
    return any(m in name for m in _CHECKLIST_MARKERS)


def _is_params_template_xlsx(path: Path) -> bool:
    low = path.name.lower()
    return any(m.lower() in low or m in path.name for m in _PARAMS_TEMPLATE_MARKERS)


def _pick_params_xlsx(candidates: list[Path]) -> Path | None:
    """优先 task_params 模板，其次非完工清单 xlsx，最后才兜底任意 xlsx。"""
    if not candidates:
        return None
    params = [p for p in candidates if _is_params_template_xlsx(p) and not _is_checklist_xlsx(p)]
    neutral = [p for p in candidates if not _is_checklist_xlsx(p) and p not in params]
    checklist = [p for p in candidates if _is_checklist_xlsx(p)]
    for group in (params, neutral, checklist):
        if group:
            return sorted(group, key=lambda p: p.stat().st_mtime, reverse=True)[0]
    return None


def _find_runtime_dir(start: Path) -> Path | None:
    for parent in start.resolve().parents:
        cand = parent / "runtime"
        if cand.is_dir() and (cand / "paths.py").is_file():
            return cand
    return None


_RUNTIME_DIR = _find_runtime_dir(Path(__file__))
if _RUNTIME_DIR and str(_RUNTIME_DIR) not in sys.path:
    sys.path.insert(0, str(_RUNTIME_DIR))


def _cloudops_params_path(skill_dir: str | Path) -> Path | None:
    root = Path(skill_dir).resolve()
    try:
        from paths import cloudops_input_dir, resolve_slot  # noqa: E402

        slot = resolve_slot("cloudops_params", skill_root=root)
        primary = str(slot.get("primary") or "").strip()
        if primary:
            p = Path(primary)
            if p.is_file():
                return p
        resolved = slot.get("resolved") or []
        if isinstance(resolved, list) and resolved:
            p = Path(str(resolved[0]))
            if p.is_file():
                return p
        inbox = cloudops_input_dir(root)
    except Exception:
        inbox = root / "ProjectData" / "input" / "cloudops"

    if not inbox.is_dir():
        return None
    candidates = [
        p for p in inbox.iterdir()
        if p.is_file() and p.suffix.lower() in {".xlsx", ".xls"}
    ]
    return _pick_params_xlsx(candidates)


def _read_template_name(path: Path, *, sheet: str) -> tuple[str | None, str | None]:
    """返回 (template_name, error_message)。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "缺少 openpyxl，无法读取 CloudOps 测试参数 xlsx。"

    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as e:
        return None, f"无法打开测试参数文件：{path.name}（{e}）"

    try:
        if sheet not in wb.sheetnames:
            return None, f"测试参数文件缺少 Sheet「{sheet}」（文件：{path.name}）。"
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return None, f"Sheet「{sheet}」为空（文件：{path.name}）。"
        header = [str(c or "").strip() for c in rows[0]]
        if _TEMPLATE_COLUMN not in header:
            return None, f"Sheet「{sheet}」缺少列「{_TEMPLATE_COLUMN}」（文件：{path.name}）。"
        col_i = header.index(_TEMPLATE_COLUMN)
        for row in rows[1:]:
            vals = list(row)
            if col_i >= len(vals):
                continue
            name = str(vals[col_i] or "").strip()
            if name:
                return name, None
        return None, None
    finally:
        wb.close()


def preflight_cloudops_params(skill_dir: str | Path, *, sheet: str) -> tuple[bool, str]:
    """检查测试参数 xlsx 存在且含指定 Sheet 与「检查项模板」列。"""
    sheet = str(sheet or "").strip()
    if not sheet:
        return False, "未配置 params_sheet。"

    path = _cloudops_params_path(skill_dir)
    if path is None:
        return (
            False,
            "请先上传 **CloudOps 测试参数 xlsx** 至 `ProjectData/input/cloudops`（槽位 cloudops_params）。",
        )

    name, err = _read_template_name(path, sheet=sheet)
    if err:
        return False, err
    if not name:
        return (
            True,
            "",
        )  # 无模板名时保留 execute.json 默认值，仍视为通过门禁
    if name not in TEMPLATE_ID_MAP and name != _DEFAULT_TEMPLATE_ID:
        known = "、".join(sorted(TEMPLATE_ID_MAP.keys()))
        return False, f"未知检查项模板「{name}」；已知模板：{known}。"
    return True, ""


def merge_template_id(
    body: dict[str, Any],
    skill_dir: str | Path,
    *,
    sheet: str,
) -> dict[str, Any]:
    """从测试参数 xlsx 合并 templateId；无文件或空模板名则保留 body 原值。"""
    out = dict(body)
    sheet = str(sheet or "").strip()
    if not sheet:
        return out

    path = _cloudops_params_path(skill_dir)
    if path is None:
        return out

    name, _ = _read_template_name(path, sheet=sheet)
    if not name:
        return out

    tid = TEMPLATE_ID_MAP.get(name)
    if tid:
        out["templateId"] = tid
        # 勿下发 templateName：CloudOps CreateTask 不识别该字段，会返回 9003。
    return out
