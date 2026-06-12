"""
CloudOps 手工补充 · 在线编辑保存（Phase 2b）。

基于第 4 步「CloudOps初始配置.xlsx」复制并 patch 指定 Sheet/列，输出「CloudOps配置_手工补充.xlsx」。
"""
from __future__ import annotations

import os
import shutil
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import os
import sys
from pathlib import Path


def _co_constants():
    root = Path(os.environ.get("SD_SKILL_ROOT", Path.cwd())).resolve()
    rt = root / "runtime"
    if str(rt) not in sys.path:
        sys.path.insert(0, str(rt))
    from sd_script_import import import_sd_script  # noqa: WPS433

    return import_sd_script(root, "4_cloudops_init", "constants")


def _supplement():
    root = Path(os.environ.get("SD_SKILL_ROOT", Path.cwd())).resolve()
    rt = root / "runtime"
    if str(rt) not in sys.path:
        sys.path.insert(0, str(rt))
    from sd_script_import import import_sd_script  # noqa: WPS433

    return import_sd_script(root, "5_cloudops_supplement", "supplement")


_c = _co_constants()
MANUAL_EDIT_SPEC = _c.MANUAL_EDIT_SPEC
CloudopsConfigSheets = _c.CloudopsConfigSheets
_sup = _supplement()
INIT_FILE_BASENAME = _sup.INIT_FILE_BASENAME
MANUAL_FILE_BASENAME = _sup.MANUAL_FILE_BASENAME
default_manual_output_path = _sup.default_manual_output_path


@dataclass
class ManualSaveResult:
    ok: bool
    message: str
    output_path: str = ""
    workspace_rel: str = ""
    rows_patched: dict[str, int] = field(default_factory=dict)
    logs: list[str] = field(default_factory=list)


def _skills_root(skill_dir: str) -> str:
    return os.path.abspath(os.path.join(skill_dir, os.pardir))


def _workspace_root(skill_dir: str) -> str:
    return os.path.abspath(os.path.join(_skills_root(skill_dir), os.pardir))


def default_init_path(skill_dir: str) -> str:
    return os.path.join(
        os.path.abspath(skill_dir), "ProjectData", "Output", INIT_FILE_BASENAME
    )


def _header_col_map(ws: Worksheet) -> dict[str, int]:
    out: dict[str, int] = {}
    for col in range(1, (ws.max_column or 0) + 1):
        val = ws.cell(row=1, column=col).value
        if val is None:
            continue
        key = str(val).strip()
        if key:
            out[key] = col
    return out


def _norm_cell(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _cell_value_for_write(v: Any) -> Any:
    """空字符串 → ``None``，用于把单元格真正清空（写入 Excel）。"""
    if v is None:
        return None
    if isinstance(v, str):
        return None if not v.strip() else v
    if isinstance(v, (int, float, bool)):
        return v
    s = str(v).strip()
    return None if not s else v


def _write_cell(ws: Worksheet, row: int, col: int, value: Any) -> None:
    """
    写入单元格。

    openpyxl 的 ``cell(row=, column=, value=None)`` 会把 ``None`` 当成「未传值」而**不修改**原内容；
    清空必须 ``cell = ws.cell(...); cell.value = None``。
    """
    ws.cell(row=row, column=col).value = _cell_value_for_write(value)


def _resolve_col(header_map: dict[str, int], name: str) -> int | None:
    if name in header_map:
        return header_map[name]
    for k, idx in header_map.items():
        if k.strip() == name.strip():
            return idx
    return None


def _patch_sheet_by_primary_key(
    ws: Worksheet,
    *,
    primary_key: str,
    editable_cols: list[str],
    rows: list[dict[str, Any]],
    logs: list[str],
) -> int:
    header_map = _header_col_map(ws)
    pk_col = _resolve_col(header_map, primary_key)
    if not pk_col:
        logs.append(f"未找到主键列「{primary_key}」")
        return 0

    edit_col_idx: dict[str, int] = {}
    for col_name in editable_cols:
        idx = _resolve_col(header_map, col_name)
        if idx:
            edit_col_idx[col_name] = idx
        else:
            logs.append(f"未找到可编辑列「{col_name}」")

    if not edit_col_idx:
        return 0

    pk_to_row: dict[str, int] = {}
    for r in range(2, (ws.max_row or 1) + 1):
        pk_val = _norm_cell(ws.cell(row=r, column=pk_col).value)
        if pk_val:
            pk_to_row[pk_val] = r

    patched = 0
    for item in rows:
        if not isinstance(item, dict):
            continue
        pk_val = _norm_cell(item.get(primary_key))
        if not pk_val or pk_val not in pk_to_row:
            continue
        row_idx = pk_to_row[pk_val]
        for col_name, col_idx in edit_col_idx.items():
            try:
                _write_cell(ws, row_idx, col_idx, item.get(col_name, ""))
            except Exception as exc:
                logs.append(f"行{row_idx}列「{col_name}」写入跳过 ({exc})")
        patched += 1
    return patched


def _os_effective_last_row(ws: Worksheet, *, tail_blank_rows: int = 5) -> int:
    """末行有数据的行号 + 预留空白行（避免遍历整张 3000+ 行模板）。"""
    header_map = _header_col_map(ws)
    if not header_map:
        return max(2, (ws.max_row or 2))
    cols = list(header_map.values())
    scan_to = min((ws.max_row or 2), 500)
    last_with_data = 1
    for r in range(2, scan_to + 1):
        for col in cols[: min(12, len(cols))]:
            if _norm_cell(ws.cell(row=r, column=col).value):
                last_with_data = r
                break
    return min((ws.max_row or 2), last_with_data + tail_blank_rows)


def _clear_os_excel_rows(
    ws: Worksheet,
    row_indices: list[int],
    header_map: dict[str, int],
    logs: list[str],
) -> int:
    """清空已从编辑页删除的 Excel 行（避免仅 patch 现存行导致删除不同步）。"""
    cleared = 0
    for raw in row_indices:
        try:
            row_idx = int(raw)
        except (TypeError, ValueError):
            continue
        if row_idx < 2:
            continue
        for col_idx in header_map.values():
            try:
                _write_cell(ws, row_idx, col_idx, "")
            except Exception as exc:
                logs.append(f"OS配置信息：清空行{row_idx}列{col_idx}跳过 ({exc})")
        cleared += 1
    if cleared:
        logs.append(f"OS配置信息：已清空删除的 {cleared} 行")
    return cleared


def _patch_os_sheet(ws: Worksheet, rows: list[dict[str, Any]], logs: list[str]) -> int:
    """
    按原表列位置原位写入（不重写表头、不清空 3000+ 空行）。

    行号优先使用 payload 中的 ``_excelRow``（与 Excel 行号一致）。
    """
    header_map = _header_col_map(ws)
    if not header_map:
        logs.append("OS配置信息：无表头")
        return 0

    allowed = set(header_map.keys())
    patched = 0
    next_fallback_row = 2

    for item in rows:
        if not isinstance(item, dict):
            continue
        row_idx_raw = item.get("_excelRow", item.get("_row"))
        if row_idx_raw is not None:
            try:
                row_idx = int(row_idx_raw)
            except (TypeError, ValueError):
                row_idx = next_fallback_row
        else:
            row_idx = next_fallback_row

        if row_idx < 2:
            row_idx = next_fallback_row
        next_fallback_row = max(next_fallback_row, row_idx) + 1

        wrote = False
        for key, val in item.items():
            if not isinstance(key, str) or key.startswith("_"):
                continue
            col_name = key.strip()
            if col_name not in allowed:
                continue
            col_idx = header_map[col_name]
            try:
                _write_cell(ws, row_idx, col_idx, val)
                wrote = True
            except Exception as exc:
                logs.append(f"OS配置信息：行{row_idx}列「{col_name}」写入跳过 ({exc})")
        if wrote:
            patched += 1

    return patched


def save_manual_from_sheets(
    *,
    skill_dir: str,
    init_path: str,
    sheets_payload: dict[str, Any],
    output_path: str | None = None,
    deleted_os_excel_rows: list[int] | None = None,
) -> ManualSaveResult:
    """
    ``sheets_payload`` 形如::

        {
          "服务器信息": [{"设备ID*（主键）": "...", "IBMC用户名": "...", ...}, ...],
          "交换机信息": [...],
          "OS配置信息": [{...}, ...]
        }
    """
    logs: list[str] = []
    out = output_path or default_manual_output_path(skill_dir)
    init_abs = os.path.abspath(init_path) if init_path else default_init_path(skill_dir)

    if not os.path.isfile(init_abs):
        return ManualSaveResult(
            ok=False,
            message=f"未找到 CloudOps 初始配置：{init_abs}（请先完成第 4 步）",
            logs=logs,
        )

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    shutil.copy2(init_abs, out)
    logs.append(f"已复制初配 → {out}")

    wb = load_workbook(out)
    rows_patched: dict[str, int] = {}

    try:
        for sheet_name, spec in MANUAL_EDIT_SPEC.items():
            raw = sheets_payload.get(sheet_name)
            if raw is None:
                continue
            if not isinstance(raw, list):
                logs.append(f"{sheet_name}：payload 非数组，已跳过")
                continue
            if sheet_name not in wb.sheetnames:
                logs.append(f"工作簿缺少 Sheet「{sheet_name}」")
                continue

            ws = wb[sheet_name]
            editable = spec.get("editable")
            if editable is None:
                header_map = _header_col_map(ws)
                if deleted_os_excel_rows:
                    _clear_os_excel_rows(ws, deleted_os_excel_rows, header_map, logs)
                n = _patch_os_sheet(ws, raw, logs)
            else:
                pk = str(spec.get("primary_key") or "").strip()
                if not pk:
                    logs.append(f"{sheet_name}：缺少 primary_key")
                    continue
                n = _patch_sheet_by_primary_key(
                    ws,
                    primary_key=pk,
                    editable_cols=list(editable),
                    rows=raw,
                    logs=logs,
                )
            rows_patched[sheet_name] = n
            logs.append(f"{sheet_name}：已更新 {n} 行")

        wb.save(out)
    except Exception as exc:
        logs.append(f"保存失败：{exc}")
        return ManualSaveResult(ok=False, message=f"保存 xlsx 失败：{exc}", output_path=out, logs=logs)
    finally:
        wb.close()

    ws_root = _workspace_root(skill_dir)
    try:
        rel = os.path.relpath(os.path.abspath(out), ws_root).replace("\\", "/")
    except ValueError:
        rel = ""

    return ManualSaveResult(
        ok=True,
        message=f"已保存「{MANUAL_FILE_BASENAME}」",
        output_path=out,
        workspace_rel=rel,
        rows_patched=rows_patched,
        logs=logs,
    )


def register_manual_save_in_chain(skill_dir: str, *, workspace_rel: str = "") -> None:
    """在线保存后写入 deploy_chain。"""
    import chain_io  # noqa: WPS433

    rel = (workspace_rel or "").strip().replace("\\", "/")
    if not rel:
        rel = "skills/deploymentandtest/ProjectData/Output/" + MANUAL_FILE_BASENAME
    chain_io.merge_chain(
        skill_dir,
        step5_cloudops_manual_path=rel,
        step5_cloudops_manual_saved_at=chain_io.iso_now(),
    )
